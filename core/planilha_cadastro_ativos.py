"""Planilha "Cadastro de Ativos" no formato do CSC Lançamentos."""
from __future__ import annotations

import io
import re

ABA = "Placa Patrimonial"
TITULO = "Placa Patrimonial (N° do bem)"
CABECALHO = ("Item", "Descrição do item", "Plaqueta", "Número de série", "Nf")
CHAVES = ("item", "descricao", "plaqueta", "numero_serie", "nf")

COR_FAIXA = "FFC00000"
LARGURAS = {"A": 7.6, "B": 123.3, "C": 14.1, "D": 28.9, "E": 8.1}
ALTURA_TITULO_E_CABECALHO = 15.6
ALTURA_DADOS = 12.75
LINHA_TITULO = 5
LINHA_CABECALHO = 6
LINHA_PRIMEIRO_DADO = 7

_SO_DIGITOS = re.compile(r"[0-9]+")
_FORA_DO_NOME = re.compile(r"[^A-Za-z0-9_-]")

def _texto(valor) -> str:
    return "" if valor is None else str(valor).strip()

def _celula_texto(valor) -> str | None:
    return _texto(valor) or None

def _celula_numero_ou_texto(valor) -> int | str | None:
    texto = _texto(valor)
    if not texto:
        return None
    return int(texto) if _SO_DIGITOS.fullmatch(texto) else texto

def montar(linhas: list[dict]) -> bytes:
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = ABA

    faixa = PatternFill("solid", fgColor=COR_FAIXA)
    fonte_faixa = Font(name="Calibri", size=12, bold=True)
    fonte_dados = Font(name="Calibri", size=11)
    centro_total = Alignment(horizontal="center", vertical="center")
    centro_horizontal = Alignment(horizontal="center")
    fina = Side(style="thin")
    borda_completa = Border(left=fina, right=fina, top=fina, bottom=fina)
    borda_titulo = Border(left=fina, top=fina, bottom=fina)

    for coluna, largura in LARGURAS.items():
        ws.column_dimensions[coluna].width = largura

    titulo = ws.cell(row=LINHA_TITULO, column=2, value=TITULO)
    titulo.font = fonte_faixa
    titulo.fill = faixa
    titulo.alignment = centro_total
    titulo.border = borda_titulo
    ws.row_dimensions[LINHA_TITULO].height = ALTURA_TITULO_E_CABECALHO

    for indice, rotulo in enumerate(CABECALHO, start=1):
        celula = ws.cell(row=LINHA_CABECALHO, column=indice, value=rotulo)
        celula.font = fonte_faixa
        celula.fill = faixa
        celula.alignment = centro_total
        celula.border = borda_completa
    ws.cell(row=LINHA_CABECALHO, column=4).number_format = "0"
    ws.row_dimensions[LINHA_CABECALHO].height = ALTURA_TITULO_E_CABECALHO

    linha = LINHA_PRIMEIRO_DADO
    for dado in linhas:
        valores = (
            _celula_texto(dado.get("item")),
            _celula_texto(dado.get("descricao")),
            _celula_numero_ou_texto(dado.get("plaqueta")),
            _celula_texto(dado.get("numero_serie")),
            _celula_numero_ou_texto(dado.get("nf")),
        )
        for indice, valor in enumerate(valores, start=1):
            celula = ws.cell(row=linha, column=indice, value=valor)
            celula.font = fonte_dados
            celula.alignment = centro_horizontal
            celula.border = borda_completa
        ws.cell(row=linha, column=4).number_format = "0"
        ws.row_dimensions[linha].height = ALTURA_DADOS
        linha += 1

    ws.freeze_panes = f"A{LINHA_PRIMEIRO_DADO}"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()

def nome_arquivo(nf: str | list[str] | tuple[str, ...]) -> str:
    partes = list(nf) if isinstance(nf, (list, tuple)) else [nf]
    limpas = [p for p in (_FORA_DO_NOME.sub("", _texto(x)) for x in partes) if p]
    return f"Cadastro_de_Ativos_NF_{'_'.join(limpas) or 'sem_numero'}.xlsx"
