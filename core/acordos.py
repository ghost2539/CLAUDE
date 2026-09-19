from __future__ import annotations

import io
import re
import unicodedata


def _norm(s) -> str:
    s = unicodedata.normalize("NFD", str(s or "").strip().lower())
    return "".join(c for c in s if unicodedata.category(c) != "Mn")


# Legenda do cabeçalho → chave interna. Casa por "contém".
_MAPA = [
    ("item_ebs", ("item ebs", "item do ebs", "cod item", "codigo ebs")),
    ("bu", ("bandeira", "bu", "unidade")),
    ("acordo_numero", ("acordo",)),
    ("vencimento", ("expira", "vencimento", "validade")),
    ("descricao_item", ("descricao", "descrição")),
    ("ncm", ("ncm",)),
    ("preco_acordo", ("preco", "preço", "valor")),
    ("fornecedor", ("fornecedor",)),
]


def _achar_colunas(cab: list) -> dict:
    """{chave_interna: indice_da_coluna} a partir da linha de cabeçalho."""
    col = {}
    for i, titulo in enumerate(cab):
        t = _norm(titulo)
        if not t:
            continue
        for chave, termos in _MAPA:
            if chave in col:
                continue
            if any(term in t for term in termos):
                col[chave] = i
                break
    return col


def _data_iso(v) -> str:
    if v is None or str(v).strip() == "":
        return ""
    try:
        return v.date().isoformat()          # datetime
    except Exception:  # noqa: BLE001
        pass
    try:
        return v.isoformat()                 # date
    except Exception:  # noqa: BLE001
        pass
    s = str(v).strip()
    m = re.match(r"^(\d{4})-(\d{2})-(\d{2})", s)
    if m:
        return m.group(0)
    m = re.match(r"^(\d{2})/(\d{2})/(\d{4})", s)
    if m:
        return f"{m.group(3)}-{m.group(2)}-{m.group(1)}"
    return s[:10]


def _preco(v) -> float:
    if v is None:
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace("R$", "").replace(" ", "")
    if "," in s and "." in s:      # 1.234,56
        s = s.replace(".", "").replace(",", ".")
    elif "," in s:                 # 1234,56
        s = s.replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return 0.0


def ler_acordos(dados: bytes) -> list[dict]:
    """Lê o XLSX (bytes) e devolve os acordos já deduplicados.

    Chave de deduplicação: (Item EBS, BU). Sem Item EBS, usa (descrição, NCM,
    preço, BU). Mantém a primeira ocorrência de cada chave.
    """
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(dados), read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]

    linhas = ws.iter_rows(values_only=True)
    col = {}
    for row in linhas:
        col = _achar_colunas(list(row))
        if "item_ebs" in col or ("descricao_item" in col and "preco_acordo" in col):
            break  # esta é a linha de cabeçalho
    if not col:
        return []

    vistos = set()
    out: list[dict] = []
    for row in linhas:
        def g(chave):
            i = col.get(chave)
            return row[i] if (i is not None and i < len(row)) else None

        item_ebs = str(g("item_ebs") or "").strip()
        descricao = str(g("descricao_item") or "").strip()
        bu = str(g("bu") or "").strip()
        ncm = str(g("ncm") or "").strip()
        preco = _preco(g("preco_acordo"))
        if not (item_ebs or descricao):
            continue
        chave = (item_ebs.lower(), bu.lower()) if item_ebs else (descricao.lower(), ncm, round(preco, 2), bu.lower())
        if chave in vistos:
            continue
        vistos.add(chave)
        out.append({
            "item_ebs": item_ebs,
            "descricao_item": descricao,
            "bu": bu,
            "ncm": ncm,
            "preco_acordo": preco,
            "acordo_numero": str(g("acordo_numero") or "").strip(),
            "fornecedor": str(g("fornecedor") or "").strip(),
            "vencimento": _data_iso(g("vencimento")),
        })
    return out
