#!/usr/bin/env python3
"""Verificação do Lançamento: planilha, chamado no ServiceNow como o usuário logado e reenvio."""
from __future__ import annotations

import http.server
import json
import os
import sys
import tempfile
import threading
import types
from pathlib import Path
from urllib.parse import parse_qs, urlsplit

RAIZ = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(RAIZ))
_TEMP = tempfile.mkdtemp(prefix="lanc-verif-")
os.environ["DATABASE_URL"] = f"sqlite:///{_TEMP}/portal.db"
os.environ.setdefault("PORTAL_SESSION_SECRET", "verificacao-local-com-64-caracteres-de-sobra-aqui")
os.environ["INITIAL_ADMIN_LOGIN"] = "admin.teste"
os.environ["AMBIENTE"] = "testes"
for _m in ("TRILHA", "OBSOLESCENCIA", "REVERSA", "AGENDAMENTOS_FORN", "INTERNALIZACAO"):
    os.environ[f"{_m}_DATABASE_URL"] = f"sqlite:///{_TEMP}/{_m.lower()}.db"
os.environ.setdefault("EBS_LOGIN_URL", "http://x")
os.environ.setdefault("EBS_SEARCH_URL", "http://x")
os.environ["PORTAL_COFRE_DIR"] = f"{_TEMP}/cofre"
os.environ["SN_PROXY"] = ""

import config as _config  # noqa: E402
_config.get_settings().DATA = Path(_TEMP) / "data"

POS = {"4512345": [
    {"po_numero": "4512345", "liberacao": None, "fornecedor": "ZEBRA DO BRASIL", "status_po": "APPROVED",
     "moeda": "BRL", "linha": 1, "item_ebs": "347191", "descricao": "ZEBRA IMPRESSORA INDUSTRIAL ZT231",
     "unidade": "UN", "quantidade_pedida": 2, "quantidade_recebida": 0, "quantidade_pendente": 2}]}
falso = types.ModuleType("integracoes.ebs_oracle")
falso.run_named = lambda nome, binds, max_rows=500: list(POS.get(binds["numero_po"], []))
sys.modules["integracoes.ebs_oracle"] = falso

G_CK = "a" * 72
ITEM = "66831e771b37b5901e870e9fe54bcb11"
REGISTRO: dict = {"pedidos": [], "anexos": [], "modo": "normal", "headers": []}


def _opc(*rotulos):
    return [{"label": r, "value": r.lower().replace(" ", "_")} for r in rotulos]


ITEM_DESCRITO = {"sys_id": ITEM, "name": "Lançamento de NF", "variables": [
    {"name": "requested_for", "type": "reference", "mandatory": True},
    {"name": "u_phone", "type": "string", "read_only": True},
    {"name": "u_impact_employee", "type": "reference"},
    {"name": "question_brand", "type": "select_box", "choices": _opc("Renner Brasil", "Camicado", "Youcom")},
    {"name": "question_document_type", "type": "select_box", "choices": _opc("Material", "Service")},
    {"name": "question_demand_type", "type": "select_box", "choices": _opc("Others", "Rent")},
    {"name": "question_has_a_purchase_order", "type": "select_box", "choices": _opc("Yes", "No")},
    {"name": "question_supplier", "type": "string"},
    {"name": "question_nature_of_the_transaction", "type": "select_box", "choices": _opc("Fixed assets (CAPEX)", "OPEX")},
    {"name": "question_the_document_is_being_sent_after_the_deadline", "type": "select_box", "choices": _opc("Yes", "No")},
    {"name": "question_number_of_documents_being_sent", "type": "string"},
    {"name": "documentos", "type": "multi_row", "children": [
        {"name": "question_order_number_po", "type": "string"},
        {"name": "question_document_number", "type": "string"},
        {"name": "question_due_date", "type": "date"}]},
    {"name": "question_description", "type": "multi_line_text"},
]}


class _SN(http.server.BaseHTTPRequestHandler):
    def _json(self, corpo, status=200):
        b = json.dumps(corpo).encode()
        self.send_response(status)
        self.send_header("Content-Type", "application/json")
        self.send_header("Content-Length", str(len(b)))
        self.end_headers()
        self.wfile.write(b)

    def _ok_headers(self):
        return self.headers.get("X-UserToken") == G_CK and self.headers.get("X-Requested-With") == "XMLHttpRequest"

    def do_GET(self):  # noqa: N802
        u = urlsplit(self.path)
        if u.path == "/login.do":
            b = b"<html>login</html>"
            self.send_response(200)
            self.send_header("Content-Type", "text/html")
            self.send_header("Content-Length", str(len(b)))
            self.end_headers()
            self.wfile.write(b)
            return
        if REGISTRO["modo"] == "expirado":
            self.send_response(302)
            self.send_header("Location", "/login.do")
            self.end_headers()
            return
        if u.path.startswith("/nav_to.do") or u.path == "/esc":
            b = f"<html><script>var g_ck = '{G_CK}';</script></html>".encode()
            self.send_response(200)
            self.send_header("Content-Type", "text/html")
            self.send_header("Content-Length", str(len(b)))
            self.end_headers()
            self.wfile.write(b)
            return
        if not self._ok_headers():
            return self._json({"error": "sem token"}, 401)
        q = parse_qs(u.query)
        if u.path == "/api/now/ui/user/current_user":
            return self._json({"result": {"user_sys_id": "1" * 32, "user_name": "rec.teste", "user_display_name": "Rec Teste", "email": "rec@x"}})
        if u.path == "/api/now/table/sys_user":
            consulta = q.get("sysparm_query", [""])[0]
            if "lucas.gonzaga" in consulta:
                return self._json({"result": [{"sys_id": "2" * 32, "name": "Lucas", "email": "lucas.gonzaga@lojasrenner.com.br", "user_name": "lucas"}]})
            return self._json({"result": []})
        if u.path == f"/api/sn_sc/servicecatalog/items/{ITEM}":
            return self._json({"result": ITEM_DESCRITO})
        if u.path == "/api/now/table/sc_req_item":
            return self._json({"result": [{"sys_id": "4" * 32, "number": "RITM0009999"}]})
        self._json({"error": "rota desconhecida " + u.path}, 404)

    def do_POST(self):  # noqa: N802
        u = urlsplit(self.path)
        n = int(self.headers.get("Content-Length") or 0)
        corpo = self.rfile.read(n) if n else b""
        if not self._ok_headers():
            return self._json({"error": "sem token"}, 401)
        if REGISTRO["modo"] == "fora":
            return self._json({"error": "manutenção"}, 503)
        if u.path == f"/api/sn_sc/servicecatalog/items/{ITEM}/order_now":
            REGISTRO["pedidos"].append(json.loads(corpo))
            return self._json({"result": {"request_id": "3" * 32, "request_number": "REQ0001234"}})
        if u.path == "/api/now/attachment/file":
            q = parse_qs(u.query)
            REGISTRO["anexos"].append({"tabela": q["table_name"][0], "sys_id": q["table_sys_id"][0],
                                       "nome": q["file_name"][0], "tipo": self.headers.get("Content-Type"),
                                       "tamanho": len(corpo)})
            return self._json({"result": {"sys_id": "5" * 32}}, 201)
        self._json({"error": "rota desconhecida " + u.path}, 404)

    def log_message(self, *a):
        pass


servidor = http.server.ThreadingHTTPServer(("127.0.0.1", 0), _SN)
threading.Thread(target=servidor.serve_forever, daemon=True).start()
SN_URL = f"http://127.0.0.1:{servidor.server_address[1]}"

from fastapi.testclient import TestClient  # noqa: E402
from sqlalchemy import select  # noqa: E402

from main import app  # noqa: E402
import core.security as sec  # noqa: E402
import db.internalizacao as int_db  # noqa: E402
import routers.servicenow as rsn  # noqa: E402

rsn.SERVICENOW_BASE = SN_URL
rsn.SN_PROXY = ""

falhas: list[str] = []
feitos = 0


def checar(cond, d) -> None:
    global feitos
    feitos += 1
    print(("  ok    " if cond else "  FALHA ") + d)
    if not cond:
        falhas.append(d)


c = TestClient(app)


def sessao(nome, modulos, admin=False, sn=True):
    pm = {m: {"can_view": True, "can_create": True, "can_edit": True, "can_admin": True} for m in modulos}
    dados = {"username": nome, "is_admin": admin, "permission_map": pm, "permissions": list(modulos)}
    if sn:
        dados["sn_cookies"] = {"glide_user_route": "abc", "JSESSIONID": "xyz"}
    _, cookie = sec.create_session(dados)
    return {"spare_session": cookie}


ADMIN = sessao("admin.teste", ("recebimento", "internalizacao", "agendamentos_forn"), admin=True)
LANC = sessao("lanc.teste", ("internalizacao",))
SEM_SN = sessao("semsn.teste", ("internalizacao",), sn=False)
B = "/api/internalizacao"
R = "/api/recebimento/fornecedores"

XML = """<?xml version="1.0" encoding="UTF-8"?>
<nfeProc xmlns="http://www.portalfiscal.inf.br/nfe" versao="4.00"><NFe><infNFe Id="NFe35250912345678000199550010000020010000002015" versao="4.00">
<ide><nNF>200153</nNF><serie>1</serie><dhEmi>2026-09-10T10:00:00-03:00</dhEmi></ide>
<emit><CNPJ>12345678000199</CNPJ><xNome>ZEBRA DO BRASIL</xNome></emit>
<det nItem="1"><prod><cProd>347191</cProd><xProd>ZEBRA IMPRESSORA INDUSTRIAL ZT231</xProd><qCom>2.0000</qCom></prod></det>
<cobr><dup><nDup>001</nDup><dVenc>2026-10-30</dVenc><vDup>15000.00</vDup></dup></cobr>
</infNFe></NFe></nfeProc>"""


def receber():
    r = c.post("/api/agendamentos-forn", cookies=ADMIN, json={
        "bu": "Renner", "fornecedor": "ZEBRA DO BRASIL", "estoque_destino": "REPOSICAO",
        "data_agendada": "2026-09-30", "pedidos": [{"po": "4512345", "nf": "200153"}],
        "equipamentos": [{"descricao": "ZEBRA ZT231", "quantidade": 2}]})
    aid = r.json()["id"]
    c.post(f"{R}/{aid}/nota/arquivo", cookies=ADMIN, data={"nf": "200153"},
           files={"arquivo": ("nf.xml", XML.encode(), "application/xml")})
    r = c.post(f"{R}/{aid}/confirmar", cookies=ADMIN, json={"itens": [
        {"po": "4512345", "nf": "200153", "linha": 1, "item_ebs": "347191", "descricao": "ZEBRA IMPRESSORA INDUSTRIAL ZT231",
         "unidade": "UN", "quantidade_pedida": 2, "quantidade_recebida": 2, "imobilizado": True,
         "seriais": [f"S{aid}A", f"S{aid}B"]}], "observacao": ""})
    assert r.status_code == 200, r.text
    return aid


c.post(f"{B}/cadastro/itens-imobilizados", cookies=ADMIN, json={"itens": [{"item_ebs": "347191", "descricao": "ZEBRA"}]})
c.post(f"{B}/cadastro/etiquetas", cookies=ADMIN, json={"prefixo": "PAT", "de": "001", "ate": "010", "local": "Armário 3"})
(Path(_TEMP) / "data" / "tmp" / "recebimento_forn").mkdir(parents=True, exist_ok=True)

print("[1] Conferir o formulário do catálogo como o usuário logado")
r = c.get(f"{B}/catalogo-sn/conferir", cookies=LANC)
checar(r.status_code == 200, f"conferir responde ({r.status_code}: {r.text[:120]})")
d = r.json()
checar(d["usuario"]["user_name"] == "rec.teste" and d["impacto"]["email"].startswith("lucas.gonzaga"), "usuário logado e impactado localizados")
nomes = [v["name"] for v in d["item"]["variables"]]
checar("documentos" in nomes and next(v for v in d["item"]["variables"] if v["name"] == "documentos")["type"] == "mrvs", "o bloco multilinha vem normalizado")
checar(d["exemplo"]["question_brand"] == "renner_brasil" and "u_phone" not in d["exemplo"] and d["faltantes"] == [],
       f"exemplo montado: opções por rótulo, telefone omitido por read_only ({d['exemplo'].get('question_brand')}, {d['faltantes']})")
checar(c.get(f"{B}/catalogo-sn/conferir", cookies=SEM_SN).status_code == 409, "sem sessão ServiceNow → 409")

print("\n[2] OK do lançamento: planilha + chamado + anexos")
a1 = receber()
r = c.get(f"{B}/{a1}", cookies=LANC).json()
checar(len(r["ativos"]) == 2 and r["ativos"][0]["plaqueta"] == "PAT001" and r["notas"][0]["tem_pdf"] in (True, False), "processo pronto com 2 ativos")
r = c.post(f"{B}/{a1}/lancar", cookies=LANC)
checar(r.status_code == 200, f"lancar responde ({r.status_code}: {r.text[:200]})")
d = r.json()
checar(d["status"] == "CONCLUIDA" and d["lancado_por"] == "lanc.teste" and d["planilha_arquivo"] == "Cadastro_de_Ativos_NF_200153.xlsx",
       f"processo concluído com a planilha nomeada pela NF ({d.get('planilha_arquivo')})")
checar(d["sn_request_number"] == "REQ0001234" and d["sn_ritm_number"] == "RITM0009999" and d["sn_erro"] == "" and not d["sn_pendente"],
       "chamado e RITM gravados")
ped = REGISTRO["pedidos"][-1]
v = ped["variables"]
checar(v["requested_for"] == "1" * 32 and v["u_impact_employee"] == "2" * 32, "requested_for é o usuário logado; impactado é o Lucas")
checar(v["question_brand"] == "renner_brasil" and v["question_document_type"] == "material" and v["question_demand_type"] == "others"
       and v["question_has_a_purchase_order"] == "yes" and v["question_nature_of_the_transaction"] == "fixed_assets_(capex)"
       and v["question_the_document_is_being_sent_after_the_deadline"] == "no", "opções fixas resolvidas por rótulo")
checar(v["question_supplier"] == "ZEBRA DO BRASIL" and v["question_number_of_documents_being_sent"] == "1", "fornecedor e quantidade de NFs")
checar(json.loads(v["documentos"]) == [{"question_order_number_po": "4512345", "question_document_number": "200153", "question_due_date": "2026-10-30"}],
       f"bloco multilinha com PO, NF e vencimento lido do XML ({v['documentos']})")
checar(v["question_description"] == "Prezados, gentileza enviar pra lançamento e cadastro de patrimonio a 200153 referente à PO 4512345.",
       "descrição com a frase exata")
checar("u_phone" not in v, "telefone não enviado (read_only)")
anexos = [a for a in REGISTRO["anexos"]]
checar(len(anexos) >= 1 and anexos[0]["tabela"] == "sc_req_item" and anexos[0]["sys_id"] == "4" * 32
       and anexos[0]["nome"] == "Cadastro_de_Ativos_NF_200153.xlsx" and anexos[0]["tamanho"] > 3000, f"planilha anexada no RITM ({anexos[:1]})")
pdf_anexado = any(a["nome"].endswith(".pdf") for a in anexos)
checar(pdf_anexado or "sem PDF" in d["aviso"], f"PDF da NF anexado quando existe, senão avisa ({d['aviso'][:80]!r})")
r = c.get(f"{B}/{a1}/planilha", cookies=LANC)
checar(r.status_code == 200 and r.content[:2] == b"PK", "a planilha baixa")
import openpyxl  # noqa: E402
import io  # noqa: E402
ws = openpyxl.load_workbook(io.BytesIO(r.content))["Placa Patrimonial"]
checar([ws.cell(row=6, column=i).value for i in range(1, 6)] == ["Item", "Descrição do item", "Plaqueta", "Número de série", "Nf"]
       and ws["A7"].value == "347191" and ws["C7"].value == "PAT001" and ws["D7"].value == f"S{a1}A" and ws["E7"].value == 200153,
       f"planilha no formato do CSC: {[ws.cell(row=7, column=i).value for i in range(1, 6)]}")
checar((Path(_TEMP) / "data" / "tmp" / "lancamentos" / str(a1) / "Cadastro_de_Ativos_NF_200153.xlsx").exists(), "arquivo em data/tmp/lancamentos/<id>/")
checar(c.post(f"{B}/{a1}/lancar", cookies=LANC).status_code == 409, "lançar de novo → 409")
checar(c.post(f"{B}/{a1}/lancar/reenviar", cookies=LANC).status_code == 409, "reenviar com chamado aberto → 409")
r = c.put(f"{B}/{a1}", cookies=LANC, json={"ativos": [], "concluir": False})
checar(r.status_code == 409, "depois do envio os ativos não mudam mais")
with int_db.SessionLocal() as s:
    at = s.scalars(select(int_db.Ativo)).all()
    checar(all(a.etapa == int_db.ETAPA_PATRIMONIO for a in at), "ativos seguem para a etapa Patrimônio")
lst = c.get(B, cookies=LANC).json()
checar(lst["itens"][0]["sn_request_number"] == "REQ0001234" and lst["itens"][0]["status"] == "CONCLUIDA", "a lista mostra o chamado")

print("\n[3] ServiceNow fora: planilha fica, chamado pendente, reenvio")
REGISTRO["modo"] = "fora"
a2 = receber()
r = c.post(f"{B}/{a2}/lancar", cookies=LANC)
checar(r.status_code == 200, f"lancar com o ServiceNow fora ainda responde 200 ({r.status_code})")
d = r.json()
checar(d["status"] == "CONCLUIDA" and d["planilha_arquivo"] and d["sn_pendente"] and "503" in d["sn_erro"], f"concluído, planilha gerada, chamado pendente com o erro ({d['sn_erro'][:60]})")
checar("não foi aberto" in d["aviso"], "aviso diz que o chamado não foi aberto")
REGISTRO["modo"] = "normal"
r = c.post(f"{B}/{a2}/lancar/reenviar", cookies=LANC)
checar(r.status_code == 200 and r.json()["sn_request_number"] == "REQ0001234" and not r.json()["sn_pendente"] and r.json()["sn_erro"] == "",
       f"reenvio abre o chamado e limpa o erro ({r.status_code}: {r.text[:100]})")

print("\n[4] Sessão ServiceNow expirada e sem sessão")
REGISTRO["modo"] = "expirado"
a3 = receber()
d = c.post(f"{B}/{a3}/lancar", cookies=LANC).json()
checar(d["sn_pendente"] and "sess" in d["sn_erro"].lower(), f"sessão expirada vira chamado pendente com o motivo ({d['sn_erro'][:60]})")
REGISTRO["modo"] = "normal"
a4 = receber()
d = c.post(f"{B}/{a4}/lancar", cookies=SEM_SN).json()
checar(d["sn_pendente"] and "Logon AD" in d["sn_erro"], "sem cookies do ServiceNow: pendente pedindo login")

print("\n[5] Vencimento da nota pelo Lançamento e validações")
a5 = receber()
r = c.patch(f"{B}/{a5}/nota/200153", cookies=LANC, json={"vencimento": "2026-12-01"})
checar(r.status_code == 200 and r.json()["vencimento"] == "2026-12-01", "vencimento gravado pela Internalização")
checar(c.patch(f"{B}/{a5}/nota/200153", cookies=LANC, json={"vencimento": "01/12/2026"}).status_code == 422, "data fora do formato → 422")
r = c.get(f"{B}/{a5}", cookies=LANC).json()
ids = [a["id"] for a in r["ativos"]]
c.put(f"{B}/{a5}", cookies=LANC, json={"ativos": [{"id": ids[0], "ebs_item": "347191", "descricao": "Z", "plaqueta": "PAT", "numero_serie": ""},
                                                  {"id": ids[1], "ebs_item": "347191", "descricao": "Z", "plaqueta": "PAT", "numero_serie": "X"}], "concluir": False})
r = c.post(f"{B}/{a5}/lancar", cookies=LANC)
checar(r.status_code == 422 and "série" in r.json()["detail"], "linha sem serial não lança (422)")
checar(c.post(f"{B}/999/lancar", cookies=LANC).status_code in (404, 422), "agendamento inexistente")
checar(c.get(f"{B}/nfe/certificados", cookies=LANC).json()["certificados"][0]["configurado"] is False, "estado dos certificados NF-e responde (nenhum configurado)")
checar(c.get(f"{B}/{a1}/planilha").status_code == 401, "sem sessão 401")

print(f"\n{feitos - len(falhas)} de {feitos} verificações passaram.")
if falhas:
    print("Falhou:\n  - " + "\n  - ".join(falhas))
    sys.exit(1)
print("Lançamento: planilha do CSC, chamado como o usuário logado, reenvio quando falha.")
