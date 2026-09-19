#!/usr/bin/env python3
"""A planilha "Cadastro de Ativos" sai igual ao modelo do CSC, e os arquivos"""
from __future__ import annotations

import io
import logging
import os
import re
import stat
import sys
import tempfile
import time
from pathlib import Path

RAIZ = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(RAIZ))

_TMP = Path(tempfile.mkdtemp())
for chave, valor in {
    "DATABASE_URL": f"sqlite:///{_TMP/'portal.db'}",
    "PORTAL_SESSION_SECRET": "verificacao-local-sem-valor",
    "PORTAL_COFRE_DIR": str(_TMP / "cofre"),
    "INTERNALIZACAO_DATABASE_URL": f"sqlite:///{_TMP/'int.db'}",
    "AGENDAMENTOS_FORN_DATABASE_URL": f"sqlite:///{_TMP/'agf.db'}",
}.items():
    os.environ[chave] = valor

falhas: list[str] = []
feitos = 0

def checar(cond, descricao):
    global feitos
    feitos += 1
    print(("  ok    " if cond else "  FALHA ") + descricao)
    if not cond:
        falhas.append(descricao)

_DATA_REPO = RAIZ / "data"
_antes = sorted(p.name for p in _DATA_REPO.iterdir()) if _DATA_REPO.is_dir() else []

import config  # noqa: E402

config.Settings.DATA = _TMP / "data"

import openpyxl  # noqa: E402

from core import arquivos_temporarios as at  # noqa: E402
from core import planilha_cadastro_ativos as pca  # noqa: E402

def _lado(borda, nome):
    lado = getattr(borda, nome)
    return None if lado is None else lado.style

def _bordas(celula) -> tuple:
    return tuple(_lado(celula.border, n) for n in ("left", "right", "top", "bottom"))

def _fonte_ok(celula, tamanho, negrito) -> bool:
    f = celula.font
    return (f.name == "Calibri" and float(f.size or 0) == tamanho
            and bool(f.bold) == negrito and f.color is None)

def _fill_faixa(celula) -> bool:
    return (celula.fill.fill_type == "solid"
            and celula.fill.fgColor.rgb == "FFC00000")

def _perto(a, b) -> bool:
    return a is not None and abs(float(a) - b) < 0.001

LINHAS = [
    {"item": "347191", "descricao": "ZEBRA IMPRESSORA INDUSTRIAL ZT231",
     "plaqueta": "2759960", "numero_serie": "T3N1ABC", "nf": "200153"},
    {"item": "347191", "descricao": "ZEBRA IMPRESSORA INDUSTRIAL ZT231",
     "plaqueta": " 2759961 ", "numero_serie": "0012345678", "nf": "200153"},
    {"item": "", "plaqueta": None},
]

conteudo = pca.montar(LINHAS)
checar(conteudo[:2] == b"PK", f"montar() devolve um .xlsx ({len(conteudo)} bytes)")
wb = openpyxl.load_workbook(io.BytesIO(conteudo))
ws = wb.active
MAX_LINHA, MAX_COLUNA = ws.max_row, ws.max_column

print("[1] Aba, linhas vazias e título em B5")
checar(wb.sheetnames == ["Placa Patrimonial"], "aba única 'Placa Patrimonial'")
vazias = all(c.value is None for linha in ws.iter_rows(min_row=1, max_row=4, max_col=5)
             for c in linha)
checar(vazias, "linhas 1-4 sem valor")
checar(ws["A5"].value is None and ws["A5"].fill.fill_type is None,
       "A5 vazia e sem preenchimento (o título é só B5)")
checar(all(ws[f"{col}5"].value is None for col in "CDE"), "C5..E5 vazias")
b5 = ws["B5"]
checar(b5.value == "Placa Patrimonial (N° do bem)", "B5 = 'Placa Patrimonial (N° do bem)'")
checar(_fonte_ok(b5, 12, True), "B5 Calibri 12 negrito, cor da fonte não definida")
checar(_fill_faixa(b5), "B5 preenchimento sólido FFC00000")
checar(b5.alignment.horizontal == "center" and b5.alignment.vertical == "center",
       "B5 alinhamento center/center")
checar(_bordas(b5) == ("thin", None, "thin", "thin"),
       f"B5 bordas finas left/top/bottom e SEM right {_bordas(b5)}")
checar(not list(ws.merged_cells.ranges), "nenhuma célula mesclada")

print("\n[2] Cabeçalho (linha 6)")
esperado = ["Item", "Descrição do item", "Plaqueta", "Número de série", "Nf"]
valores6 = [ws.cell(row=6, column=i).value for i in range(1, 6)]
checar(valores6 == esperado, f"cabeçalho {valores6}")
checar(ws.cell(row=6, column=6).value is None, "nada além da coluna E na linha 6")
for col in "ABCDE":
    c = ws[f"{col}6"]
    checar(_fonte_ok(c, 12, True) and _fill_faixa(c)
           and c.alignment.horizontal == "center" and c.alignment.vertical == "center"
           and _bordas(c) == ("thin",) * 4,
           f"{col}6 Calibri 12 negrito sem cor, fill FFC00000, center/center, 4 bordas finas")
checar(ws["D6"].number_format == "0", "D6 number_format '0'")
checar(ws["C6"].number_format == "General", "C6 sem formato especial (o '0' é só da coluna D)")

print("\n[3] Dados (linhas 7-9): tipos e estilo")
checar(MAX_LINHA == 9 and MAX_COLUNA == 5,
       f"três linhas de dados, cinco colunas (max_row={MAX_LINHA}, max_column={MAX_COLUNA})")
for r in (7, 8, 9):
    for col in "ABCDE":
        c = ws[f"{col}{r}"]
        checar(_fonte_ok(c, 11, False) and c.alignment.horizontal == "center"
               and c.alignment.vertical is None and _bordas(c) == ("thin",) * 4
               and c.fill.fill_type is None,
               f"{col}{r} Calibri 11 sem negrito, horizontal center (vertical não definido), 4 bordas, sem fill")
    checar(ws[f"D{r}"].number_format == "0", f"D{r} number_format '0'")
    checar(ws[f"C{r}"].number_format == "General", f"C{r} sem formato especial")
checar(ws["A7"].value == "347191" and isinstance(ws["A7"].value, str),
       "A7 item '347191' como TEXTO")
checar(ws["B7"].value == "ZEBRA IMPRESSORA INDUSTRIAL ZT231", "B7 descrição")
checar(ws["C7"].value == 2759960 and isinstance(ws["C7"].value, int),
       "C7 plaqueta 2759960 como INTEIRO")
checar(ws["D7"].value == "T3N1ABC", "D7 número de série")
checar(ws["E7"].value == 200153 and isinstance(ws["E7"].value, int),
       "E7 NF 200153 como INTEIRO")
checar(ws["C8"].value == 2759961 and isinstance(ws["C8"].value, int),
       "C8 plaqueta com espaços em volta vira inteiro")
checar(ws["D8"].value == "0012345678" and isinstance(ws["D8"].value, str),
       "D8 serial só de dígitos continua TEXTO (zero à esquerda preservado)")
checar(all(ws[f"{col}9"].value is None for col in "ABCDE"),
       "linha 9: vazio, None e chave faltante viram célula vazia")
checar(ws.row_dimensions[10].height is None, "linha 10 não existe (sem altura definida)")

print("\n[4] Larguras, alturas, congelamento e grade")
for col, larg in {"A": 7.6, "B": 123.3, "C": 14.1, "D": 28.9, "E": 8.1}.items():
    checar(_perto(ws.column_dimensions[col].width, larg),
           f"largura {col} = {larg} ({ws.column_dimensions[col].width})")
for r in (5, 6):
    checar(_perto(ws.row_dimensions[r].height, 15.6),
           f"altura linha {r} = 15.6 ({ws.row_dimensions[r].height})")
for r in (7, 8, 9):
    checar(_perto(ws.row_dimensions[r].height, 12.75),
           f"altura linha {r} = 12.75 ({ws.row_dimensions[r].height})")
checar(ws.row_dimensions[4].height is None, "linha 4 sem altura definida")
checar(ws.freeze_panes == "A7", f"freeze_panes = A7 ({ws.freeze_panes})")
checar(ws.sheet_view.showGridLines is not False, "linhas de grade no padrão (não desligadas)")

ws_vazia = openpyxl.load_workbook(io.BytesIO(pca.montar([]))).active
checar(ws_vazia.max_row == 6 and ws_vazia["A6"].value == "Item"
       and ws_vazia.freeze_panes == "A7",
       "sem linhas: título e cabeçalho ficam, congelamento também")

print("\n[5] nome_arquivo")
checar(pca.nome_arquivo("200153") == "Cadastro_de_Ativos_NF_200153.xlsx",
       "uma NF → Cadastro_de_Ativos_NF_200153.xlsx")
checar(pca.nome_arquivo(["200153", "200154"]) == "Cadastro_de_Ativos_NF_200153_200154.xlsx",
       "várias NFs unidas por '_'")
PADRAO = re.compile(r"^Cadastro_de_Ativos_NF_[A-Za-z0-9_-]+\.xlsx$")
for bruto in ("200.153/A b", "../../x", "", ["", " 7 "], "NF ~ 9"):
    nome = pca.nome_arquivo(bruto)
    checar(bool(PADRAO.fullmatch(nome)), f"{bruto!r} → {nome} (só [A-Za-z0-9_-])")
checar(pca.nome_arquivo("../../x") == "Cadastro_de_Ativos_NF_x.xlsx",
       "'../../x' vira só 'x'")

print("\n[6] Contraprova: a exportação antiga da Internalização não mudou")
import routers.internalizacao as ri  # noqa: E402

antigo = openpyxl.load_workbook(io.BytesIO(ri._montar_xlsx("200153", [
    {"descricao": "X", "plaqueta": "1", "numero_serie": "S1", "ebs_item": "347191"},
]))).active
cab_antigo = [antigo.cell(row=6, column=i).value for i in range(1, 6)]
checar(cab_antigo[0] == "NF", f"cabeçalho antigo continua começando por NF {cab_antigo}")
checar("A5:D5" in {str(r) for r in antigo.merged_cells.ranges},
       "título antigo continua mesclado em A5:D5")
checar(antigo["A6"].font.color is not None and antigo["A6"].font.color.rgb == "00FFFFFF",
       "fonte branca do cabeçalho antigo continua")
checar(antigo.sheet_view.showGridLines is False, "grade desligada no antigo continua")

print("\n[7] Arquivos temporários: pasta sob data/tmp, segmentos validados")
RAIZ_TMP = _TMP / "data" / "tmp"
p7 = at.pasta("lancamentos", 7)
checar(p7 == RAIZ_TMP / "lancamentos" / "7" and p7.is_dir(),
       f"pasta('lancamentos', 7) = {p7.relative_to(_TMP)}")
checar(stat.S_IMODE(p7.stat().st_mode) == 0o750, "pasta criada com modo 0o750")
checar(stat.S_IMODE(p7.parent.stat().st_mode) == 0o750, "pasta da área também 0o750")
checar(at.RETENCAO_DIAS == 5, "RETENCAO_DIAS = 5")
for area, chave in (("..", 7), ("lancamentos", "../x"), ("lanc/amentos", 1),
                    ("lancamentos", ""), ("lancamentos", "7 "), ("notas", "a.b")):
    try:
        at.pasta(area, chave)
        recusou = False
    except ValueError:
        recusou = True
    checar(recusou, f"pasta({area!r}, {chave!r}) recusada com ValueError")
checar(not (_TMP / "x").exists() and not (_TMP / "data" / "x").exists(),
       "nada criado fora de data/tmp pelas tentativas recusadas")

print("\n[8] gravar / caminho / listar")
dados = b"conteudo da planilha"
arq = at.gravar("lancamentos", 7, "Cadastro_de_Ativos_NF_200153.xlsx", dados)
checar(arq == p7 / "Cadastro_de_Ativos_NF_200153.xlsx" and arq.read_bytes() == dados,
       "gravar() escreve o conteúdo no nome pedido")
checar(stat.S_IMODE(arq.stat().st_mode) == 0o640, "arquivo com modo 0o640")
checar(not any(p.name.startswith(".") for p in p7.iterdir()),
       "nenhum arquivo parcial sobrou depois da escrita atômica")
mal = at.gravar("lancamentos", 7, "../../fora.xlsx", b"x")
checar(mal.parent == p7 and "/" not in mal.name and not mal.name.startswith("."),
       f"nome com '../' fica dentro da pasta e sem ponto inicial ({mal.name})")
checar(not (_TMP / "fora.xlsx").exists() and not (_TMP / "data" / "fora.xlsx").exists(),
       "nada gravado fora da pasta")
mal.unlink()
try:
    at.gravar("lancamentos", 7, "..", b"x")
    recusou = False
except ValueError:
    recusou = True
checar(recusou, "gravar() com nome '..' recusado")

checar(at.caminho("lancamentos", 7, "Cadastro_de_Ativos_NF_200153.xlsx") == arq.resolve(),
       "caminho() devolve o arquivo gravado")
checar(at.caminho("lancamentos", 7, "../../../../etc/passwd") is None,
       "caminho() com '../' devolve None")
checar(at.caminho("lancamentos", 7, "inexistente.xlsx") is None,
       "caminho() de arquivo inexistente devolve None")
checar(at.caminho("..", 7, "Cadastro_de_Ativos_NF_200153.xlsx") is None,
       "caminho() com área inválida devolve None (sem exceção)")
checar(at.caminho("lancamentos", 99, "x.xlsx") is None
       and not (RAIZ_TMP / "lancamentos" / "99").exists(),
       "caminho() não cria pasta")
fora = _TMP / "fora.txt"
fora.write_bytes(b"segredo")
os.symlink(fora, p7 / "link.txt")
checar(at.caminho("lancamentos", 7, "link.txt") is None,
       "link simbólico para fora da pasta é negado (resolve e compara)")
checar(at.caminho("lancamentos", 7, ".parcial-abc") is None,
       "nome com ponto inicial (arquivo parcial) é negado")

lista = at.listar("lancamentos", 7)
checar(len(lista) == 1 and lista[0]["nome"] == "Cadastro_de_Ativos_NF_200153.xlsx",
       f"listar() traz só o arquivo gravado (link simbólico fora) {[i['nome'] for i in lista]}")
checar(set(lista[0]) == {"nome", "tamanho", "modificado_em"}
       and lista[0]["tamanho"] == len(dados)
       and re.fullmatch(r"\d{4}-\d{2}-\d{2}T\d{2}:\d{2}:\d{2}", lista[0]["modificado_em"]),
       "cada item tem nome, tamanho e modificado_em (ISO, segundos)")
checar(at.listar("lancamentos", 123) == [], "listar() de pasta inexistente é lista vazia")
(p7 / "link.txt").unlink()

print("\n[9] limpar_expirados: só o que passou do prazo, e a pasta vazia")
SEIS_DIAS = time.time() - 6 * 86400
QUATRO_DIAS = time.time() - 4 * 86400
velho = at.gravar("lancamentos", 8, "velho.xlsx", b"v")
os.utime(velho, (SEIS_DIAS, SEIS_DIAS))
quase = at.gravar("lancamentos", 9, "quase.xlsx", b"q")
os.utime(quase, (QUATRO_DIAS, QUATRO_DIAS))
nota_velha = at.gravar("notas", 3, "nf.xml", b"<nfe/>")
os.utime(nota_velha, (SEIS_DIAS, SEIS_DIAS))

n = at.limpar_expirados(5, area="lancamentos")
checar(n == 1, f"área 'lancamentos': apagou 1 ({n})")
checar(not velho.exists(), "o de 6 dias foi apagado")
checar(not velho.parent.exists(), "a pasta que ficou vazia foi removida")
checar(quase.exists() and arq.exists(), "o de 4 dias e o recente ficaram")
checar(nota_velha.exists(), "outra área não foi tocada quando area= é dada")
checar((RAIZ_TMP / "lancamentos").is_dir(), "a pasta da área com conteúdo fica")

n = at.limpar_expirados(5)
checar(n == 1, f"sem área: apagou a nota velha ({n})")
checar(not nota_velha.exists() and not (RAIZ_TMP / "notas").exists(),
       "nf.xml apagado; notas/3 e notas/ vazias removidas")
checar(RAIZ_TMP.is_dir(), "data/tmp em si fica")
checar(at.limpar_expirados(5, area="nada") == 0, "área inexistente: 0, sem erro")
try:
    at.limpar_expirados(5, area="../..")
    recusou = False
except ValueError:
    recusou = True
checar(recusou, "área '../..' recusada (nunca sobe acima de data/tmp)")

class _Coletor(logging.Handler):
    def __init__(self):
        super().__init__()
        self.mensagens: list[str] = []

    def emit(self, registro):
        self.mensagens.append(registro.getMessage())

coletor = _Coletor()
logging.getLogger("arquivos_temporarios").addHandler(coletor)
travado = at.gravar("lancamentos", 10, "travado.xlsx", b"t")
solto = at.gravar("lancamentos", 11, "solto.xlsx", b"s")
for p in (travado, solto):
    os.utime(p, (SEIS_DIAS, SEIS_DIAS))
_unlink_real = Path.unlink

def _unlink_falho(self, *a, **k):
    if self.name == "travado.xlsx":
        raise OSError("simulado: arquivo em uso")
    return _unlink_real(self, *a, **k)

Path.unlink = _unlink_falho
try:
    n = at.limpar_expirados(5)
finally:
    Path.unlink = _unlink_real
    logging.getLogger("arquivos_temporarios").removeHandler(coletor)
checar(n == 1 and not solto.exists() and not solto.parent.exists(),
       "erro num arquivo não interrompe: o outro foi apagado e sua pasta removida")
checar(travado.exists() and travado.parent.exists(),
       "o arquivo que falhou fica, e a pasta dele também")
checar(any("travado.xlsx" in m for m in coletor.mensagens),
       f"a falha foi para o log ({len(coletor.mensagens)} mensagem)")

print("\n[10] Contraprova: nada gravado no data/ do repositório")
_depois = sorted(p.name for p in _DATA_REPO.iterdir()) if _DATA_REPO.is_dir() else []
checar(_antes == _depois, f"data/ do repositório inalterado {_depois}")
checar(RAIZ_TMP.is_relative_to(_TMP), "tudo ficou sob o Settings.DATA temporário")

print(f"\n{feitos - len(falhas)} de {feitos} verificações passaram.")
if falhas:
    print("Falhou:\n  - " + "\n  - ".join(falhas))
    sys.exit(1)
print("Planilha igual ao modelo do CSC; temporários presos em data/tmp e com prazo.")
