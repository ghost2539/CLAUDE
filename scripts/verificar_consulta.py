#!/usr/bin/env python3
"""Verificação da Consulta de ativos direto na base do EBS (portal e Times).

    python3 scripts/verificar_consulta.py
"""
from __future__ import annotations
import os, re, sys, tempfile, types
from pathlib import Path
RAIZ = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(RAIZ))
_T = tempfile.mkdtemp(prefix="consulta-verif-")
for v in sorted(set(re.findall(r"([A-Z_]*DATABASE_URL)", (RAIZ / "config.py").read_text()))):
    os.environ[v] = f"sqlite:///{_T}/{v.lower()}.db"
os.environ["PORTAL_SESSION_SECRET"] = "verificacao-local"
os.environ["PORTAL_COFRE_DIR"] = f"{_T}/cofre"
CRED = {"ORACLE_EBS_DSN": "ebsdb.interno.local:1521/EBSPRD",
        "ORACLE_EBS_USER": "leitor_spare", "ORACLE_EBS_PASS": "segredo-teste"}
os.environ.update(CRED)
import logging; logging.disable(logging.CRITICAL)

BASE = [
    {"asset_id": 1001, "imobilizado": "100200", "etiqueta": "RN000123", "numero_serie": "SN-ABC-1",
     "descricao": "COLETOR TC21", "fabricante": "ZEBRA", "modelo_ebs": "TC21", "livro": "RENNER CORP",
     "custo": 3500.5, "dpis": "2024-03-01", "baixado": "N", "data_baixa": None,
     "local_atribuido": "RS.POA.LOJA101", "po": "2570313-32", "nf": "NF445566"},
    {"asset_id": 1002, "imobilizado": "300400", "etiqueta": "CM000777", "numero_serie": "SN-XYZ-9",
     "descricao": "PDV DELL", "fabricante": "DELL", "modelo_ebs": "OPTIPLEX", "livro": "CAMICADO CORP",
     "custo": 5000, "dpis": "2021-01-15", "baixado": "S", "data_baixa": "2025-06-30",
     "local_atribuido": "SP.SAO.CD324", "po": None, "nf": None},
]
COLS = list(BASE[0].keys())
CLASSICO = {
    "FA_ADDITIONS_B": {"ASSET_ID", "ASSET_NUMBER", "TAG_NUMBER", "SERIAL_NUMBER", "DESCRIPTION",
                       "MANUFACTURER_NAME", "MODEL_NUMBER"},
    "FA_BOOKS": {"ASSET_ID", "BOOK_TYPE_CODE", "COST", "DATE_PLACED_IN_SERVICE", "DATE_INEFFECTIVE",
                 "PERIOD_COUNTER_FULLY_RETIRED"},
    "FA_BOOK_CONTROLS": {"BOOK_TYPE_CODE", "BOOK_CLASS"},
    "FA_RETIREMENTS": {"ASSET_ID", "BOOK_TYPE_CODE", "DATE_RETIRED", "STATUS"},
    "FA_DISTRIBUTION_HISTORY": {"ASSET_ID", "LOCATION_ID", "DATE_INEFFECTIVE"},
    "FA_LOCATIONS": {"LOCATION_ID", "SEGMENT1", "SEGMENT2", "SEGMENT3", "SEGMENT4", "SEGMENT5",
                     "SEGMENT6", "SEGMENT7"},
    "FA_ASSET_INVOICES": {"ASSET_ID", "PO_NUMBER", "INVOICE_NUMBER", "DATE_INEFFECTIVE",
                          "INVOICE_DISTRIBUTION_ID"},
    "AP_INVOICE_DISTRIBUTIONS_ALL": {"INVOICE_DISTRIBUTION_ID", "PO_DISTRIBUTION_ID"},
    "PO_DISTRIBUTIONS_ALL": {"PO_DISTRIBUTION_ID", "PO_RELEASE_ID"},
    "PO_RELEASES_ALL": {"PO_RELEASE_ID", "RELEASE_NUM"},
}
# A instalação de verdade do cliente: R12, descrição na tabela traduzida e
# FA_BOOKS sem date_retired.
R12 = {t: set(c) for t, c in CLASSICO.items()}
R12["FA_ADDITIONS_B"].discard("DESCRIPTION")
R12["FA_ADDITIONS_TL"] = {"ASSET_ID", "DESCRIPTION", "LANGUAGE"}
CATALOGO = {t: set(c) for t, c in CLASSICO.items()}
REG = {"conexoes": 0, "binds": [], "sql": [], "catalogo": 0, "connect_kwargs": None, "falha": None}


class _LOB:
    pass


class _Cursor:
    arraysize = 100

    def __init__(self):
        self._rows = []
        self.description = [(c.upper(),) for c in COLS]

    def __enter__(self):
        return self

    def __exit__(self, *a):
        return False

    def execute(self, sql, binds=None):
        if sql.strip().upper().startswith("SET TRANSACTION"):
            return
        # O catálogo responde só o que ESTA conta enxerga: é assim que o portal
        # descobre coluna e tabela que faltam nesta instalação.
        if "ALL_TAB_COLUMNS" in sql.upper():
            REG["catalogo"] += 1
            self.description = [("TABELA",), ("COLUNA",)]
            self._rows = [(tab, col) for tab, cols in sorted(CATALOGO.items()) for col in sorted(cols)]
            return
        REG["sql"].append(sql)
        REG["binds"].append(dict(binds or {}))
        self.description = [(c.upper(),) for c in COLS]
        termos = {str(v) for v in (binds or {}).values()}
        linhas = []
        for r in BASE:
            if not (r["imobilizado"] in termos or r["etiqueta"] in termos or r["numero_serie"] in termos):
                continue
            d = dict(r)
            # Sem a cadeia da liberação no SQL, a base só teria a raiz da PO.
            if "PO_RELEASES_ALL" not in sql and "-" in str(d.get("po") or ""):
                d["po"] = str(d["po"]).split("-")[0]
            for campo in COLS:
                if f"NULL AS {campo}" in sql:
                    d[campo] = None
            linhas.append(tuple(d[c] for c in COLS))
        self._rows = linhas

    def fetchmany(self, n):
        return self._rows[:n]

    def fetchall(self):
        return self._rows


class _Conn:
    def cursor(self):
        return _Cursor()

    def rollback(self):
        pass

    def close(self):
        pass


def _connect(**kw):
    if REG["falha"]:
        raise REG["falha"]
    REG["conexoes"] += 1
    REG["connect_kwargs"] = kw
    return _Conn()


fake = types.ModuleType("oracledb")
fake.connect = _connect
fake.init_oracle_client = lambda **kw: None
fake.LOB = _LOB
sys.modules["oracledb"] = fake

import main  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402
from core.security import create_session  # noqa: E402
from config import get_settings  # noqa: E402
import db.portal as dbp  # noqa: E402
from db.portal import SessionLocal, Classification, Asset, ReceiptCycle, LocalAsset  # noqa: E402
import routers.consulta_times as rct  # noqa: E402

falhas, feitos = [], 0
def checar(c, d):
    global feitos; feitos += 1
    print(("  ok   " if c else "  FALHA ") + d)
    if not c: falhas.append(d)

dbp.init_db()
c = TestClient(main.app)
sid, cookie = create_session({"username": "sso.user", "display_name": "SSO", "is_admin": True, "auth_source": "SSO",
                              "permissions": ["admin"], "permission_map": {}, "user_id": 1, "ebs_auth": None})
c.cookies.set("spare_session", cookie)
COLUNAS = ["empresa", "imobilizado", "etiqueta", "numero_serie", "descricao", "categoria",
           "local_atribuido", "baixado", "po", "nf", "erro"]

print("\n[1] Busca por número de série, etiqueta ou imobilizado, numa única ida à base por lote")
c.post("/api/consulta", json={"identificadores": ["aquece-o-catalogo"]})
REG.update(conexoes=0, binds=[], sql=[], catalogo=0)
r = c.post("/api/consulta", json={"identificadores": ["sn-abc-1", "CM000777", "300400", "NAOEXISTE1"]})
checar(r.status_code == 200, f"POST /api/consulta → 200 (veio {r.status_code})")
d = r.json()
checar(d["encontrados"] == 3 and d["nao_encontrados"] == 1, "três achados (série, etiqueta, imobilizado) e um não")
por = {x["pesquisado"]: x for x in d["resultados"]}
a = por["sn-abc-1"]
checar(a["encontrado"] and a["imobilizado"] == "100200" and a["etiqueta"] == "RN000123"
       and a["numero_serie"] == "SN-ABC-1", "série digitada em minúsculas acha o ativo")
checar(a["empresa"] == "RENNER" and a["descricao"] == "COLETOR TC21"
       and a["local_atribuido"] == "RS.POA.LOJA101", "empresa pelo livro, descrição e local atribuído")
checar(a["baixado"] == "Não" and a["po"] == "2570313-32" and a["nf"] == "NF445566" and a["erro"] == "",
       "Baixado?, PO (com a liberação) e NF vindos da base")
b = por["CM000777"]
checar(b["encontrado"] and b["empresa"] == "CAMICADO" and b["baixado"] == "Sim"
       and b["data_baixa"] == "2025-06-30", "busca pela etiqueta; ativo baixado vem como Sim, com a data")
checar(por["300400"]["encontrado"] and por["300400"]["etiqueta"] == "CM000777", "busca pelo imobilizado")
checar(not por["NAOEXISTE1"]["encontrado"] and por["NAOEXISTE1"]["erro"], "não encontrado traz o motivo na coluna Erro")
checar(all(k in a for k in COLUNAS), "a linha traz as onze colunas da tela")
checar(REG["conexoes"] == 1 and len(REG["sql"]) == 1,
       f"o lote inteiro foi numa conexão e numa consulta (conexões={REG['conexoes']}, consultas={len(REG['sql'])})")
sql = REG["sql"][0]
checar("/*TERMOS*/" not in sql and ":t0" in sql and sql.count("IN (:t0") == 3,
       "SQL do arquivo com os binds :tN nos três campos (asset_number, tag_number, serial_number)")
checar(set(REG["binds"][0].values()) == {"sn-abc-1", "SN-ABC-1", "CM000777", "300400", "NAOEXISTE1"},
       "binds: cada termo como digitado e em maiúsculas, sem repetir")
checar(REG["connect_kwargs"] == {"user": "leitor_spare", "password": "segredo-teste",
                                 "dsn": "ebsdb.interno.local:1521/EBSPRD"},
       "conexão com a credencial do ambiente do serviço (os.environ)")
checar(REG["catalogo"] == 0, "o catálogo da base é lido uma vez por processo, não a cada consulta")
REG.update(conexoes=0, binds=[], sql=[])
muitos = [f"SN-{n:04d}" for n in range(120)] + ["SN-ABC-1"]
d = c.post("/api/consulta", json={"identificadores": muitos}).json()
checar(len(d["resultados"]) == 121 and d["encontrados"] == 1 and REG["conexoes"] == 2,
       f"121 termos → dois lotes de até 100, nada descartado (conexões={REG['conexoes']})")
r = c.get("/api/consulta/single", params={"identificador": "RN000123"})
checar(r.status_code == 200 and r.json()["imobilizado"] == "100200" and "ciclo" in r.json(),
       "GET /api/consulta/single acha pela etiqueta")

print("\n[1b] A consulta é montada com as colunas que a conta enxerga")
import integracoes.ebs_ativos as ea  # noqa: E402
sql_classico = ea.montar_sql(":t0", CLASSICO)
checar("fb.date_retired" not in sql_classico and "period_counter_fully_retired" in sql_classico,
       "FA_BOOKS não tem date_retired: a baixa sai de period_counter_fully_retired")
checar("fa.description" in sql_classico and "FA_ADDITIONS_TL" not in sql_classico,
       "instalação com descrição na tabela base continua lendo de lá")
busca = sql_classico[sql_classico.index("WITH alvos"):sql_classico.index(")\nSELECT")]
checar(busca.count("SELECT fa0.asset_id") == 3 and busca.count("UNION") == 2 and " OR " not in busca,
       "busca por UNION de três colunas indexadas, sem OR (era o que deixava a consulta lenta)")
sql_r12 = ea.montar_sql(":t0", R12)
checar("fa.description" not in sql_r12 and "tl.description" in sql_r12
       and "LEFT JOIN APPS.FA_ADDITIONS_TL" in sql_r12 and "USERENV('LANG')" in sql_r12,
       "R12 sem DESCRIPTION na tabela base: a descrição vem da tabela traduzida")
sem_nota = {t_: set(c) for t_, c in R12.items() if t_ != "FA_ASSET_INVOICES"}
sql_curto = ea.montar_sql(":t0", sem_nota)
checar("APPS.FA_ASSET_INVOICES" not in sql_curto and "NULL AS po" in sql_curto and "NULL AS nf" in sql_curto,
       "tabela que a conta não enxerga sai do SQL, e o campo vira NULL")
sem_loc = {t_: set(c) for t_, c in R12.items() if t_ != "FA_LOCATIONS"}
checar("NULL AS local_atribuido" in ea.montar_sql(":t0", sem_loc),
       "bloco que precisa de duas tabelas sai quando falta qualquer uma delas")
sql_cego = ea.montar_sql(":t0", {})
checar("fa.description" in sql_cego and "APPS.FA_ASSET_INVOICES" in sql_cego
       and "NULL AS" not in sql_cego and "FA_ADDITIONS_TL" not in sql_cego,
       "catálogo que não diz nada não descarta nada: vale a consulta inteira")
checar(ea.mapa(R12)["DESCRICAO"] == "tl.description" and ea.mapa(CLASSICO)["DESCRICAO"] == "fa.description",
       "o mapa de campos diz de qual coluna cada campo veio")

# PO de contrato: a raiz é a mesma e o que muda é a liberação (2570313-32).
import re as _re  # noqa: E402
def _quantos_po(sql):
    return len(_re.findall(r"AS po,?\s*$", sql, _re.M))
checar("PO_RELEASES_ALL" in sql_classico and "'-' || pr.release_num" in sql_classico
       and _quantos_po(sql_classico) == 1,
       "PO sai com a liberação (2570313-32), pela distribuição da fatura")
sem_release = {t_: set(c) for t_, c in CLASSICO.items() if t_ != "PO_RELEASES_ALL"}
sql_sem_rel = ea.montar_sql(":t0", sem_release)
checar("PO_RELEASES_ALL" not in sql_sem_rel and "MAX(ai.po_number)" in sql_sem_rel
       and _quantos_po(sql_sem_rel) == 1,
       "sem as tabelas da liberação, a PO ainda vem, só sem o sufixo")
checar("MAX(ai.po_number)" in sql_cego and "PO_RELEASES_ALL" not in sql_cego
       and _quantos_po(sql_cego) == 1,
       "catálogo que não diz nada não arrisca a cadeia da liberação")
sql_sem_nota = ea.montar_sql(":t0", {t_: set(c) for t_, c in CLASSICO.items() if t_ != "FA_ASSET_INVOICES"})
checar(sql_sem_nota.count("NULL AS po") == 1 and _quantos_po(sql_sem_nota) == 1,
       "sem a tabela das faturas, a coluna PO entra uma vez só, vazia")

CATALOGO.clear(); CATALOGO.update({t_: set(c) for t_, c in R12.items()})
CATALOGO.pop("FA_ASSET_INVOICES")
ea._catalogo = None
REG.update(conexoes=0, sql=[], catalogo=0)
d = c.post("/api/consulta", json={"identificadores": ["SN-ABC-1"]}).json()
linha = d["resultados"][0]
checar(REG["catalogo"] == 1 and linha["encontrado"], "na instalação do cliente, a consulta responde")
checar(linha["po"] == "" and linha["nf"] == "" and linha["descricao"] == "COLETOR TC21",
       "o que a conta não enxerga vem em branco; o resto continua vindo da base")
CATALOGO.update({t_: set(c) for t_, c in CLASSICO.items()})
ea._catalogo = None
linha = c.post("/api/consulta", json={"identificadores": ["SN-ABC-1"]}).json()["resultados"][0]
checar(linha["po"] == "2570313-32", "na tela, a PO de contrato aparece com a liberação")
CATALOGO.pop("PO_RELEASES_ALL")
ea._catalogo = None
linha = c.post("/api/consulta", json={"identificadores": ["SN-ABC-1"]}).json()["resultados"][0]
checar(linha["po"] == "2570313", "sem as tabelas da liberação, ainda aparece a raiz da PO")
CATALOGO.clear()
CATALOGO.update({t_: set(c) for t_, c in R12.items() if t_ != "FA_ASSET_INVOICES"})
ea._catalogo = None
r = c.get("/api/ebs-oracle/consulta-de-ativos")
diag = r.json()
checar(r.status_code == 200 and diag["campos"]["DESCRICAO"] == "tl.description"
       and "FA_ADDITIONS_TL" in diag["tabelas"],
       "a tela de diagnóstico mostra a coluna escolhida para cada campo")
checar("NULL AS po" in diag["sql"], "o diagnóstico mostra o SQL realmente montado")
CATALOGO.clear(); CATALOGO.update({t_: set(c) for t_, c in CLASSICO.items()})
ea._catalogo = None
c.post("/api/consulta", json={"identificadores": ["aquece"]})

print("\n[2] Categoria do cadastro e PO/NF do recebimento quando a base não traz")
with SessionLocal.begin() as s:
    s.add(Classification(description_pattern="PDV DELL", company="CAMICADO", category="PDV", model="OptiPlex 3000"))
    ativo = Asset(company="CAMICADO", asset_number="300400", tag_number="CM000777",
                  serial_number="SN-XYZ-9", description="PDV DELL")
    s.add(ativo); s.flush()
    s.add(ReceiptCycle(asset_id=ativo.id, cycle_number=1, origem_entrada="fornecedor",
                       po="PO112233", nf="NF998877", created_by="teste"))
d = c.post("/api/consulta", json={"identificadores": ["SN-XYZ-9", "SN-ABC-1"]}).json()
por = {x["pesquisado"]: x for x in d["resultados"]}
checar(por["SN-XYZ-9"]["categoria"] == "PDV" and por["SN-XYZ-9"]["modelo"] == "OptiPlex 3000",
       "categoria e modelo do cadastro da tela de recebimento")
checar(por["SN-XYZ-9"]["po"] == "PO112233" and por["SN-XYZ-9"]["nf"] == "NF998877",
       "PO e NF do ciclo de recebimento quando o EBS não tem")
checar(por["SN-ABC-1"]["po"] == "2570313-32" and por["SN-ABC-1"]["nf"] == "NF445566", "PO e NF do EBS prevalecem")
checar(por["SN-ABC-1"]["categoria"] == "NÃO CLASSIFICADA", "sem cadastro → NÃO CLASSIFICADA")
r = c.post("/api/consulta/export", json={"identificadores": ["SN-XYZ-9"]})
checar(r.status_code == 200 and "spreadsheetml" in r.headers.get("content-type", ""), "exportação em Excel")
import io  # noqa: E402
from openpyxl import load_workbook  # noqa: E402
ws = load_workbook(io.BytesIO(r.content)).active
cab = [x.value for x in ws[1]]
linha = dict(zip(cab, [x.value for x in ws[2]]))
checar(all(x in cab for x in ("Baixado", "Po", "Nf", "Local Atribuido")), "planilha traz Baixado, PO, NF e Local")
checar(linha["Baixado"] == "Sim" and linha["Po"] == "PO112233" and linha["Categoria"] == "PDV", "valores na planilha")

print("\n[3] Base do EBS fora: erro na coluna, sem dado de acesso; base local como apoio")
with SessionLocal.begin() as s:
    s.add(LocalAsset(company="YOUCOM", asset_number="555666", tag_number="YC000555",
                     serial_number="SN-LOCAL-5", description="IMPRESSORA ZEBRA", active=True))
REG["falha"] = RuntimeError("ORA-12541: TNS:no listener em ebsdb.interno.local:1521/EBSPRD para leitor_spare")
r = c.post("/api/consulta", json={"identificadores": ["SN-ABC-1", "SN-LOCAL-5"]})
checar(r.status_code == 200, f"base fora não derruba a tela: 200 (veio {r.status_code})")
por = {x["pesquisado"]: x for x in r.json()["resultados"]}
e = por["SN-ABC-1"]
checar(not e["encontrado"] and e["erro"].startswith("A base do EBS recusou"), "linha não achada traz o erro na coluna Erro")
checar("ebsdb.interno.local" not in e["erro"] and "EBSPRD" not in e["erro"] and "leitor_spare" not in e["erro"],
       "erro sem endereço, instância ou usuário da base")
loc = por["SN-LOCAL-5"]
checar(loc["encontrado"] and loc["fonte"] == "BASE LOCAL" and loc["etiqueta"] == "YC000555" and loc["erro"],
       "ativo da base local aparece, com o aviso na coluna Erro")
REG["falha"] = None
for k in CRED:
    os.environ.pop(k, None)
d = c.post("/api/consulta", json={"identificadores": ["SN-ABC-1"]}).json()
erro = d["resultados"][0]["erro"]
checar("Credencial da base do EBS ausente" in erro and "ORACLE_EBS_DSN" in erro,
       "sem credencial no ambiente do serviço, a linha diz qual chave falta")
os.environ.update(CRED)
d = c.post("/api/consulta", json={"identificadores": ["SN-LOCAL-5"]}).json()
checar(d["resultados"][0]["encontrado"] and d["resultados"][0]["fonte"] == "BASE LOCAL"
       and d["resultados"][0]["erro"] == "Não está no EBS; dados da base local.",
       "com a base no ar, ativo só da base local ainda aparece, avisando a origem")

print("\n[4] Consulta Times sem login; ServiceNow do espaço e o portal continuam exigindo login")
anon = TestClient(main.app)
r = anon.get("/consulta-times")
checar(r.status_code == 200 and "Consulta de Ativos" in r.text, "página /consulta-times abre sem sessão")
r = anon.post("/api/consulta-times/consulta", json={"identificadores": ["RN000123", "NAOEXISTE2"]})
checar(r.status_code == 200 and r.json()["encontrados"] == 1, "API da consulta Times responde sem sessão")
l0 = r.json()["resultados"][0]
checar(all(k in l0 for k in COLUNAS) and l0["empresa"] == "RENNER" and l0["local_atribuido"] == "RS.POA.LOJA101",
       "Times traz as mesmas onze colunas")
r = anon.post("/api/consulta-times/consulta/export", json={"identificadores": ["RN000123"]})
checar(r.status_code == 200 and "spreadsheetml" in r.headers.get("content-type", ""), "exportação Times sem sessão")
for rota in ("/api/consulta-times/gestao-ativos", "/api/consulta-times/stockrooms",
             "/api/consulta-times/liberacoes", "/api/consulta-times/eu"):
    r = anon.get(rota)
    checar(r.status_code == 401, f"{rota} sem sessão → 401 (veio {r.status_code})")
r = anon.post("/api/consulta", json={"identificadores": ["RN000123"]})
checar(r.status_code == 401, f"POST /api/consulta do portal sem sessão → 401 (veio {r.status_code})")
esp = TestClient(rct.criar_app_espelho())
r = esp.get("/consulta-times")
checar(r.status_code == 200, "espelho na porta antiga abre a página sem sessão")
r = esp.post("/api/consulta-times/consulta", json={"identificadores": ["300400"]})
checar(r.status_code == 200 and r.json()["encontrados"] == 1, "espelho consulta sem sessão")
idx = (RAIZ / "static/consulta-times/index.html").read_text(encoding="utf-8")
item_sn = idx.split('data-route="servicenow"')[1].split("</a>")[0]
checar("disabled" not in item_sn and "login" in item_sn.lower(), "menu Times: item ServiceNow ativo, avisando que exige login")
app_js = (RAIZ / "static/consulta-times/app.js").read_text(encoding="utf-8")
checar("'/#gestao_ativos/entrada'" in app_js and "exige login" in app_js,
       "ServiceNow do espaço leva ao portal, que aplica login e permissão")
checar("setProgresso" in app_js and "LOTE = 50" in app_js, "Times: barra de progresso e lotes de 50")
checar(all(f"key: '{k}'" in app_js for k in COLUNAS), "Times: as onze colunas na tela")

print("\n[5] Nada da API REST antiga sobrou")
padrao = re.compile(r"EBS_LOGIN_URL|EBS_SEARCH_URL|ebs_service\b|ebs_logged|ebs_public_username")
sobras = []
for p in RAIZ.rglob("*"):
    if not p.is_file() or p.suffix not in (".py", ".js", ".html", ".sh", ".md", ".service", ".example", ".modelo", ".json"):
        continue
    if any(parte.startswith(".") or parte in ("node_modules", "__pycache__", "data") for parte in p.relative_to(RAIZ).parts[:-1]):
        continue
    if p.resolve() == Path(__file__).resolve():
        continue
    try:
        texto = p.read_text(encoding="utf-8")
    except (UnicodeDecodeError, OSError):
        continue
    if padrao.search(texto):
        sobras.append(str(p.relative_to(RAIZ)))
checar(not sobras, "nenhuma referência a EBS_LOGIN_URL, EBS_SEARCH_URL, ebs_service ou ebs_logged"
       + (" (sobrou: " + ", ".join(sobras[:6]) + ")" if sobras else ""))
checar(not (RAIZ / "integracoes/ebs_service.py").exists() and not (RAIZ / "integracoes/ebs_logged.py").exists(),
       "ebs_service.py e ebs_logged.py removidos")
cfg = get_settings()
checar(not hasattr(cfg, "EBS_LOGIN_URL") and not hasattr(cfg, "EBS_SEARCH_URL") and not hasattr(cfg, "CREDENTIALS_DIRECTORY"),
       "config sem EBS_LOGIN_URL, EBS_SEARCH_URL e CREDENTIALS_DIRECTORY")
r = c.get("/api/parametros/ebs")
checar(r.status_code in (404, 405), f"/api/parametros/ebs não existe mais (veio {r.status_code})")
cols = c.get("/api/consulta/colunas").json()
checar([x["chave"] for x in cols["disponiveis"]] == COLUNAS, "/api/consulta/colunas oferece as onze colunas, na ordem")
for arq in ("modulos/consulta.js", "modulos-times/consulta.js"):
    t = (RAIZ / arq).read_text(encoding="utf-8")
    checar(all(f"key: '{k}'" in t for k in COLUNAS) and "barraProgresso" in t and "LOTE = 50" in t,
           f"{arq}: onze colunas, barra de progresso e lotes de 50")
    checar("etiqueta ou imobilizado" in t, f"{arq}: campo diz que aceita série, etiqueta ou imobilizado")
pa = (RAIZ / "modulos/parametros_admin.js").read_text(encoding="utf-8")
checar("_renderEbs" not in pa and "API de consulta" not in pa, "Parâmetros sem o cartão da API do EBS")

print(f"\n{feitos - len(falhas)} de {feitos} verificações passaram.")
if falhas:
    print("Falhas:\n  - " + "\n  - ".join(falhas)); sys.exit(1)
print("Consulta íntegra.")
