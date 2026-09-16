#!/usr/bin/env python3
"""TMA dos chamados do Spare a partir de uma planilha de INCs, consultando o SN.

Script AUTÔNOMO: não depende de nada do portal. Só precisa de Python 3,
`requests` e `openpyxl`. A autenticação é HTTP Basic com uma CONTA DE SERVIÇO
do ServiceNow (não usa SSO) — o usuário e a senha são informados na hora.

A planilha (coluna 'Tarefa') dá os números dos chamados; TODO o resto vem do
ServiceNow, consultado em LOTES:

- incident:            sys_id, subcategory, opened_at, resolved_at, assignment_group
- sys_journal_field:   work notes/comments — confirma apontamento do time
                       (contém "Modalidade de envio" / "Modalidade Envio")
- sys_audit:           histórico de assignment_group (com data) — entrada/saída do Spare

Janela de atendimento (do mais novo para o mais antigo, tratando "bouncing"):
- Entrada  = última vez que o chamado foi atribuído à fila do Spare
             (ou a abertura, se nasceu na nossa fila).
- Saída    = a transferência para outra fila logo após essa entrada; ou, se foi
             resolvido dentro do Spare, a data de RESOLVIDO (resolved_at).
- TMA      = Saída − Entrada, em DIAS: corrido e útil (sem sábado/domingo).

Uso:
    python3 tma_chamados.py ENTRADA.xlsx [SAIDA.xlsx] [opções]

Opções:
    --limit N        processa só os N primeiros chamados (teste)
    --usuario NOME   conta de serviço (senão, pergunta na hora)
    --instancia URL  base do ServiceNow (padrão: https://renner.service-now.com)
    --proxy URL      proxy https, se necessário (ex.: http://10.115.35.45:8888)
    --sem-proxy      força NÃO usar proxy
    --selftest       roda só o autoteste da lógica (sem SN)

O usuário pode vir por --usuario ou pela variável SN_USER; a senha, pela
variável SN_PASS ou digitada na hora (getpass, não aparece na tela).
"""
from __future__ import annotations

import getpass
import os
import re
import sys
import time
import unicodedata
from datetime import datetime, timedelta

SPARE_GRUPO = "TI_N2_FLD_RNR_LOJAS_SPARE"
SN_BASE_PADRAO = "https://renner.service-now.com"
APONTAMENTOS = ("modalidade de envio", "modalidade envio")

# Lotes: gentis com a conta de serviço.
LOTE_INCIDENTS = 100
LOTE_SYSIDS = 60
PAUSA_S = 0.2


# ── Utilidades puras (testáveis sem SN) ─────────────────────────────────────
def _norm(s) -> str:
    s = unicodedata.normalize("NFD", str(s or "").lower())
    return "".join(c for c in s if unicodedata.category(c) != "Mn")


def _parse_dt(v):
    if not v:
        return None
    if isinstance(v, datetime):
        return v
    s = str(v).strip()
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%d",
                "%d/%m/%Y %H:%M:%S", "%d/%m/%Y %H:%M", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, fmt)
        except ValueError:
            continue
    return None


def tem_apontamento(textos) -> bool:
    """True se algum work note/comentário fala de envio de equipamento."""
    for t in textos:
        n = _norm(t)
        if any(marca in n for marca in APONTAMENTOS):
            return True
    return False


def _eh_spare(valor, tokens) -> bool:
    v = _norm(valor)
    return bool(v) and v in tokens


def resolver_janela(criacao, audits, resolved_at, grupo_atual, tokens_spare):
    """Devolve (entrada, saida, tipo_saida) do último ciclo no Spare.

    `audits`: lista de dicts {ts: datetime, old: str, new: str} (mudanças de
    assignment_group). `tokens_spare`: set com nome e sys_id do grupo Spare,
    normalizados. Trata bouncing pegando a ENTRADA mais recente no Spare.
    """
    linhas = sorted([a for a in audits if a.get("ts")], key=lambda a: a["ts"])
    entradas = [a["ts"] for a in linhas if _eh_spare(a.get("new"), tokens_spare)
                and not _eh_spare(a.get("old"), tokens_spare)]
    saidas = [a["ts"] for a in linhas if _eh_spare(a.get("old"), tokens_spare)
              and not _eh_spare(a.get("new"), tokens_spare)]

    nasceu_no_spare = False
    if linhas:
        # Se a 1ª mudança registrada saiu do Spare, é porque nasceu no Spare.
        nasceu_no_spare = _eh_spare(linhas[0].get("old"), tokens_spare)
    elif _eh_spare(grupo_atual, tokens_spare):
        nasceu_no_spare = True

    if entradas:
        entrada = max(entradas)
    elif nasceu_no_spare or _eh_spare(grupo_atual, tokens_spare):
        entrada = criacao
    else:
        return None, None, "sem_passagem_spare"

    if entrada is None:
        return None, None, "sem_data_entrada"

    pos = [s for s in saidas if s >= entrada]
    if pos:
        return entrada, min(pos), "transferencia"
    if resolved_at and resolved_at >= entrada:
        return entrada, resolved_at, "resolvido"
    return entrada, None, "em_aberto"


def dias_corridos(a, b) -> float:
    return round((b - a).total_seconds() / 86400.0, 2) if (a and b) else None


def dias_uteis(a, b) -> float:
    """Dias entre a e b contando só segunda a sexta (24h por dia útil)."""
    if not (a and b) or b <= a:
        return 0.0 if (a and b) else None
    seg = 0.0
    cur = a
    while cur < b:
        fim_dia = (cur + timedelta(days=1)).replace(hour=0, minute=0, second=0, microsecond=0)
        trecho = min(fim_dia, b)
        if cur.weekday() < 5:                      # 0-4 = seg-sex
            seg += (trecho - cur).total_seconds()
        cur = trecho
    return round(seg / 86400.0, 2)


# ── Carga da planilha ───────────────────────────────────────────────────────
_RE_INC = re.compile(r"\b(INC\d{5,})\b", re.IGNORECASE)


def _linhas_da_planilha(caminho: str):
    """Retorna uma lista de linhas (cada linha = lista de células em texto),
    detectando o formato REAL do arquivo — não confia na extensão.

    O ServiceNow costuma exportar como .xls/.xlsx que na verdade é HTML-tabela
    ou CSV, ou como .xls binário antigo. Cobrimos todos.
    """
    with open(caminho, "rb") as fh:
        cabeca = fh.read(8)

    # 1) .xlsx de verdade (zip: assinatura PK\x03\x04)
    if cabeca[:4] == b"PK\x03\x04":
        import openpyxl
        wb = openpyxl.load_workbook(caminho, read_only=True, data_only=True)
        ws = wb[wb.sheetnames[0]]
        return [[c for c in row] for row in ws.iter_rows(values_only=True)]

    # 2) .xls binário antigo (OLE: assinatura D0 CF 11 E0)
    if cabeca[:4] == b"\xd0\xcf\x11\xe0":
        try:
            import xlrd  # type: ignore
        except ImportError:
            raise SystemExit(
                "Arquivo é .xls binário antigo. Instale xlrd (pip install xlrd) "
                "ou reexporte como CSV / .xlsx real."
            )
        wb = xlrd.open_workbook(caminho)
        sh = wb.sheet_by_index(0)
        return [[sh.cell_value(r, c) for c in range(sh.ncols)]
                for r in range(sh.nrows)]

    # 3) Texto: pode ser HTML-tabela (SN "Excel") ou CSV/TSV
    raw = open(caminho, "rb").read()
    for enc in ("utf-8-sig", "utf-16", "latin-1"):
        try:
            texto = raw.decode(enc)
            break
        except (UnicodeDecodeError, UnicodeError):
            continue
    else:
        texto = raw.decode("latin-1", "replace")

    baixo = texto.lower()
    if "<table" in baixo or "<tr" in baixo or "<td" in baixo:
        return _linhas_html(texto)

    # CSV/TSV — deixa o csv.Sniffer achar o delimitador
    import csv
    import io
    amostra = texto[:4096]
    try:
        dialeto = csv.Sniffer().sniff(amostra, delimiters=",;\t|")
    except csv.Error:
        class _D(csv.Dialect):
            delimiter = ";" if amostra.count(";") > amostra.count(",") else ","
            quotechar = '"'
            doublequote = True
            skipinitialspace = True
            lineterminator = "\n"
            quoting = csv.QUOTE_MINIMAL
        dialeto = _D
    return [list(row) for row in csv.reader(io.StringIO(texto), dialeto)]


def _linhas_html(texto: str):
    """Extrai as linhas da primeira <table> de um export HTML do ServiceNow."""
    from html.parser import HTMLParser
    from html import unescape

    class _P(HTMLParser):
        def __init__(self):
            super().__init__()
            self.linhas, self.cel, self.buf = [], None, []
            self.in_td = False
            self.parou = False

        def handle_starttag(self, tag, attrs):
            if self.parou:
                return
            if tag == "tr":
                self.cel = []
            elif tag in ("td", "th"):
                self.in_td, self.buf = True, []

        def handle_endtag(self, tag):
            if self.parou:
                return
            if tag in ("td", "th") and self.in_td:
                self.cel.append(unescape("".join(self.buf)).strip())
                self.in_td = False
            elif tag == "tr" and self.cel is not None:
                self.linhas.append(self.cel)
                self.cel = None
            elif tag == "table" and self.linhas:
                self.parou = True  # só a primeira tabela

        def handle_data(self, data):
            if self.in_td:
                self.buf.append(data)

    p = _P()
    p.feed(texto)
    return p.linhas


def carregar_incs(caminho: str):
    linhas = _linhas_da_planilha(caminho)
    linhas = [lin for lin in linhas if any(str(c or "").strip() for c in lin)]
    if not linhas:
        raise SystemExit("Planilha vazia ou ilegível.")

    hdr = [str(c or "").strip().lower() for c in linhas[0]]

    def idx(*nomes):
        for n in nomes:
            if n in hdr:
                return hdr.index(n)
        return None

    i_num = idx("tarefa", "number", "chamado", "incident", "número", "numero")
    i_grp = idx("assignment group", "assignment_group", "grupo")

    out, vistos = [], set()

    def _add(num, grp=""):
        num = str(num or "").strip()
        m = _RE_INC.search(num)
        if m:
            num = m.group(1).upper()
        if not num or num in vistos:
            return
        vistos.add(num)
        out.append({"number": num, "grupo_planilha": str(grp or "").strip()})

    if i_num is not None:
        for r in linhas[1:]:
            grp = r[i_grp] if (i_grp is not None and i_grp < len(r)) else ""
            _add(r[i_num] if i_num < len(r) else "", grp)
    else:
        # Sem cabeçalho reconhecível: varre tudo procurando padrão INC#####
        for r in linhas:
            for c in r:
                m = _RE_INC.search(str(c or ""))
                if m:
                    _add(m.group(1))
    if not out:
        raise SystemExit(
            "Nenhum número de chamado (INC…) encontrado. Confira a coluna 'Tarefa'."
        )
    return out


# ── Credenciais e sessão SN (Basic Auth — conta de serviço) ─────────────────
def _get_requests():
    try:
        import requests
        import urllib3
        urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)
        return requests
    except ImportError:
        raise SystemExit("Pacote 'requests' não instalado. Rode: pip install requests")


def obter_credenciais(usuario_cli=None):
    """Usuário/senha da conta de serviço, informados na hora (ou por env/CLI)."""
    usuario = usuario_cli or os.environ.get("SN_USER") or ""
    if not usuario:
        try:
            usuario = input("Usuário da conta de serviço do ServiceNow: ").strip()
        except EOFError:
            usuario = ""
    if not usuario:
        raise SystemExit("Usuário não informado.")
    senha = os.environ.get("SN_PASS") or ""
    if not senha:
        senha = getpass.getpass(f"Senha de {usuario} (não aparece na tela): ")
    if not senha:
        raise SystemExit("Senha não informada.")
    return usuario, senha


def abrir_sessao_sn(base, usuario, senha, proxy=None):
    """Sessão requests com HTTP Basic Auth (conta de serviço, sem SSO).

    Faz um GET de teste no incident (1 registro) para validar o login antes
    de começar os lotes. Retorna a sessão pronta.
    """
    requests = _get_requests()
    session = requests.Session()
    session.verify = False
    session.auth = (usuario, senha)
    session.headers.update({
        "Accept": "application/json",
        "X-Requested-With": "XMLHttpRequest",
    })
    if proxy:
        session.proxies = {"https": proxy, "http": proxy}
    session._sn_base = base.rstrip("/")  # guardado para os helpers

    print(f"[SN] Testando acesso como {usuario} em {session._sn_base} …")
    url = (f"{session._sn_base}/incident.do?JSONv2"
           "&sysparm_action=getRecords&sysparm_record_count=1&sysparm_fields=number")
    try:
        r = session.get(url, timeout=30)
    except Exception as exc:  # noqa: BLE001
        raise SystemExit(f"Falha de conexão ao ServiceNow: {exc}")
    if r.status_code in (401, 403) or "login" in r.url.lower():
        raise SystemExit(
            f"Autenticação recusada (HTTP {r.status_code}). Confira usuário/senha "
            "da conta de serviço e se ela tem acesso à API (JSONv2)."
        )
    if "json" not in r.headers.get("Content-Type", "").lower():
        raise SystemExit(
            f"Resposta inesperada (HTTP {r.status_code}). Talvez precise de proxy "
            "(--proxy) ou a conta caiu em tela de SSO. Confira o acesso."
        )
    print("[SN] Acesso OK.")
    return session


def _chunks(lst, n):
    for i in range(0, len(lst), n):
        yield lst[i:i + n]


def _plain(v):
    return v.get("value", "") if isinstance(v, dict) else (v if v is not None else "")


def _sn_query(session, table, query="", fields="", limit=500, offset=0):
    """Uma página da tabela via JSONv2 + Basic Auth. Retorna lista de dicts."""
    params = ["sysparm_action=getRecords", f"sysparm_record_count={limit}"]
    if query:
        from urllib.parse import quote
        params.append("sysparm_query=" + quote(query, safe="=^,<>!@."))
    if fields:
        params.append(f"sysparm_fields={fields}")
    if offset:
        params.append(f"sysparm_first_row={offset}")
    url = f"{session._sn_base}/{table}.do?JSONv2&{'&'.join(params)}"
    r = session.get(url, timeout=60)
    if r.status_code in (401, 403) or "login" in r.url.lower():
        raise SystemExit("Sessão ServiceNow recusada/expirou (HTTP "
                         f"{r.status_code}). Rode de novo e confira as credenciais.")
    if r.status_code != 200 or "json" not in r.headers.get("Content-Type", "").lower():
        raise SystemExit(f"ServiceNow retornou HTTP {r.status_code} em {table}.")
    return r.json().get("records", [])


def _sn_query_all(session, table, query="", fields="", page_size=500, max_records=100000):
    """Pagina via sysparm_first_row até esgotar (ou max_records)."""
    todos, offset = [], 0
    while offset < max_records:
        lote = _sn_query(session, table, query, fields, page_size, offset)
        if not lote:
            break
        todos.extend(lote)
        if len(lote) < page_size:
            break
        offset += page_size
        time.sleep(PAUSA_S)
    return todos


def coletar_sn(session, numeros):
    """Consulta incident + work notes + histórico de fila em lotes.
    Devolve dict number -> {sys_id, subcategoria, resolved_at, grupo_atual,
    textos:[...], audits:[{ts,old,new}]}."""
    # 1) Grupo Spare: nome + sys_id (para casar no audit).
    tokens = {_norm(SPARE_GRUPO)}
    try:
        g = _sn_query_all(session, "sys_user_group", f"name={SPARE_GRUPO}",
                          "sys_id,name", page_size=1, max_records=1)
        if g:
            tokens.add(_norm(_plain(g[0].get("sys_id"))))
    except SystemExit:
        raise
    except Exception as exc:  # noqa: BLE001
        print(f"[SN] aviso: não resolvi o sys_id do grupo Spare ({exc})")

    dados = {}
    # 2) Incidents por número (lotes).
    total = len(numeros)
    for k, chunk in enumerate(_chunks(numeros, LOTE_INCIDENTS), 1):
        q = "numberIN" + ",".join(chunk)
        recs = _sn_query_all(session, "incident", q,
                             "number,sys_id,subcategory,opened_at,resolved_at,assignment_group",
                             page_size=LOTE_INCIDENTS, max_records=LOTE_INCIDENTS)
        for r in recs:
            num = _plain(r.get("number"))
            dados[num] = {
                "sys_id": _plain(r.get("sys_id")),
                "subcategoria": _plain(r.get("subcategory")) or "(sem subcategoria)",
                "opened_at": _parse_dt(_plain(r.get("opened_at"))),
                "resolved_at": _parse_dt(_plain(r.get("resolved_at"))),
                "grupo_atual": _plain(r.get("assignment_group")),
                "textos": [], "audits": [],
            }
        print(f"[SN] incidents {min(k*LOTE_INCIDENTS, total)}/{total}")
        time.sleep(PAUSA_S)

    sysid_para_num = {d["sys_id"]: n for n, d in dados.items() if d["sys_id"]}
    sysids = list(sysid_para_num)

    # 3) Work notes / comentários (lotes por element_id).
    for chunk in _chunks(sysids, LOTE_SYSIDS):
        q = "element_idIN" + ",".join(chunk) + "^elementINwork_notes,comments"
        recs = _sn_query_all(session, "sys_journal_field", q,
                             "element_id,value,sys_created_on",
                             page_size=1000, max_records=5000)
        for r in recs:
            num = sysid_para_num.get(_plain(r.get("element_id")))
            if num:
                dados[num]["textos"].append(_plain(r.get("value")))
        time.sleep(PAUSA_S)
    print("[SN] work notes coletadas.")

    # 4) Histórico de assignment_group (sys_audit, lotes por documentkey).
    for chunk in _chunks(sysids, LOTE_SYSIDS):
        q = ("tablename=incident^fieldname=assignment_group^documentkeyIN"
             + ",".join(chunk))
        recs = _sn_query_all(session, "sys_audit", q,
                             "documentkey,oldvalue,newvalue,sys_created_on",
                             page_size=1000, max_records=20000)
        for r in recs:
            num = sysid_para_num.get(_plain(r.get("documentkey")))
            if num:
                dados[num]["audits"].append({
                    "ts": _parse_dt(_plain(r.get("sys_created_on"))),
                    "old": _plain(r.get("oldvalue")),
                    "new": _plain(r.get("newvalue")),
                })
        time.sleep(PAUSA_S)
    print("[SN] histórico de filas coletado.")
    return dados, tokens


# ── Relatório ───────────────────────────────────────────────────────────────
def montar(numeros_planilha, dados, tokens, saida_xlsx):
    import openpyxl
    from statistics import mean, median

    linhas = []
    for item in numeros_planilha:
        num = item["number"]
        d = dados.get(num)
        if not d:
            linhas.append({"number": num, "encontrado": False})
            continue
        apont = tem_apontamento(d["textos"])
        entrada, saida, tipo = resolver_janela(
            d["opened_at"], d["audits"], d["resolved_at"], d["grupo_atual"], tokens)
        corrido = dias_corridos(entrada, saida)
        util = dias_uteis(entrada, saida)
        linhas.append({
            "number": num, "encontrado": True,
            "subcategoria": d["subcategoria"], "grupo_atual": d["grupo_atual"],
            "apontamento": apont, "entrada": entrada, "saida": saida,
            "tipo_saida": tipo, "tma_corrido": corrido, "tma_util": util,
        })

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Chamados"
    cab = ["Chamado", "Encontrado no SN", "Subcategoria", "Fila atual",
           "Apontamento de envio", "Entrada no Spare", "Saída/Resolvido",
           "Tipo de saída", "TMA corrido (dias)", "TMA útil (dias)"]
    ws.append(cab)
    for L in linhas:
        if not L.get("encontrado"):
            ws.append([L["number"], "NÃO", "", "", "", "", "", "", "", ""])
            continue
        ws.append([
            L["number"], "Sim", L["subcategoria"], L["grupo_atual"],
            "Sim" if L["apontamento"] else "Não",
            L["entrada"].strftime("%d/%m/%Y %H:%M") if L["entrada"] else "",
            L["saida"].strftime("%d/%m/%Y %H:%M") if L["saida"] else "",
            L["tipo_saida"], L["tma_corrido"], L["tma_util"],
        ])

    # Resumo por subcategoria: só chamados com apontamento e TMA calculado.
    val = [L for L in linhas if L.get("encontrado") and L.get("apontamento")
           and L.get("tma_corrido") is not None]
    grupos = {}
    for L in val:
        grupos.setdefault(L["subcategoria"], []).append(L)
    ws2 = wb.create_sheet("Resumo por subcategoria")
    ws2.append(["Subcategoria", "Qtd chamados", "TMA corrido médio (dias)",
                "TMA corrido mediana", "TMA útil médio (dias)", "TMA útil mediana"])
    for sub in sorted(grupos):
        g = grupos[sub]
        c = [x["tma_corrido"] for x in g]
        u = [x["tma_util"] for x in g]
        ws2.append([sub, len(g), round(mean(c), 2), round(median(c), 2),
                    round(mean(u), 2), round(median(u), 2)])
    if val:
        c = [x["tma_corrido"] for x in val]
        u = [x["tma_util"] for x in val]
        ws2.append(["TOTAL", len(val), round(mean(c), 2), round(median(c), 2),
                    round(mean(u), 2), round(median(u), 2)])

    for w in (ws, ws2):
        for col in w.columns:
            larg = max((len(str(c.value)) for c in col if c.value is not None), default=10)
            w.column_dimensions[col[0].column_letter].width = min(max(larg + 2, 12), 46)

    wb.save(saida_xlsx)
    total = len(linhas)
    achados = sum(1 for L in linhas if L.get("encontrado"))
    com_ap = sum(1 for L in linhas if L.get("apontamento"))
    print(f"\nResumo: {total} chamados | {achados} no SN | {com_ap} com apontamento "
          f"| {len(val)} no TMA | {len(grupos)} subcategorias")
    print(f"Arquivo gerado: {saida_xlsx}")


# ── Autoteste da lógica pura (sem SN) ───────────────────────────────────────
def _selftest():
    T = _norm(SPARE_GRUPO)
    d = datetime
    # Parcial/bouncing: entra, sai, volta, resolvido no Spare.
    audits = [
        {"ts": d(2026, 1, 1, 9), "old": "FILA_A", "new": SPARE_GRUPO},   # entra
        {"ts": d(2026, 1, 3, 9), "old": SPARE_GRUPO, "new": "FILA_B"},   # sai
        {"ts": d(2026, 1, 6, 9), "old": "FILA_B", "new": SPARE_GRUPO},   # volta (mais novo)
    ]
    ent, sai, tipo = resolver_janela(d(2026, 1, 1, 8), audits, d(2026, 1, 8, 9),
                                     SPARE_GRUPO, {T})
    assert ent == d(2026, 1, 6, 9) and sai == d(2026, 1, 8, 9) and tipo == "resolvido", (ent, sai, tipo)
    # Transferência após a entrada:
    audits2 = [
        {"ts": d(2026, 2, 2, 9), "old": "FILA_A", "new": SPARE_GRUPO},
        {"ts": d(2026, 2, 5, 9), "old": SPARE_GRUPO, "new": "FILA_C"},
    ]
    ent, sai, tipo = resolver_janela(d(2026, 2, 1), audits2, None, "FILA_C", {T})
    assert tipo == "transferencia" and sai == d(2026, 2, 5, 9), (ent, sai, tipo)
    # Nasceu no Spare, resolvido lá (sem audits de entrada):
    ent, sai, tipo = resolver_janela(d(2026, 3, 2, 10), [], d(2026, 3, 4, 10),
                                     SPARE_GRUPO, {T})
    assert ent == d(2026, 3, 2, 10) and tipo == "resolvido", (ent, sai, tipo)
    # Dias corrido x útil: sex 10h -> seg 10h = 3 corridos, 1 útil.
    assert dias_corridos(d(2026, 1, 2, 10), d(2026, 1, 5, 10)) == 3.0
    assert dias_uteis(d(2026, 1, 2, 10), d(2026, 1, 5, 10)) == 1.0
    # Apontamento (com/sem acento):
    assert tem_apontamento(["Modalidade de Envio: Correios"]) is True
    assert tem_apontamento(["modalidade envio sedex"]) is True
    assert tem_apontamento(["texto qualquer"]) is False
    print("selftest OK")


def _opt(argv, nome, default=None):
    """Lê --nome VALOR ou --nome=VALOR."""
    for i, a in enumerate(argv):
        if a == nome and i + 1 < len(argv):
            return argv[i + 1]
        if a.startswith(nome + "="):
            return a.split("=", 1)[1]
    return default


def main(argv):
    if "--selftest" in argv:
        _selftest()
        return

    # posicionais = tudo que não é flag nem valor de flag
    flags_com_valor = {"--limit", "--usuario", "--instancia", "--proxy"}
    args, pular = [], False
    for i, a in enumerate(argv):
        if pular:
            pular = False
            continue
        if a.startswith("--"):
            if a in flags_com_valor and "=" not in a:
                pular = True
            continue
        args.append(a)

    if not args:
        print(__doc__)
        raise SystemExit(2)

    entrada = args[0]
    saida = args[1] if len(args) > 1 else "TMA_chamados_resultado.xlsx"

    limite_s = _opt(argv, "--limit")
    limite = int(limite_s) if limite_s else None
    usuario_cli = _opt(argv, "--usuario")
    base = _opt(argv, "--instancia", os.environ.get("SN_BASE", SN_BASE_PADRAO))
    if "--sem-proxy" in argv:
        proxy = None
    else:
        proxy = _opt(argv, "--proxy", os.environ.get("SN_PROXY"))

    # 1) Lê a planilha ANTES de pedir senha (falha cedo se o arquivo estiver ruim).
    itens = carregar_incs(entrada)
    if limite:
        itens = itens[:limite]
    print(f"{len(itens)} chamados para consultar.")

    # 2) Credenciais (na hora) e sessão.
    usuario, senha = obter_credenciais(usuario_cli)
    session = abrir_sessao_sn(base, usuario, senha, proxy)

    # 3) Consulta e relatório.
    numeros = [i["number"] for i in itens]
    dados, tokens = coletar_sn(session, numeros)
    montar(itens, dados, tokens, saida)


if __name__ == "__main__":
    main(sys.argv[1:])
