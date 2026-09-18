#!/usr/bin/env python3
"""Verificação do Controle de Orçamento: "Em andamento" e o disponível.

    python3 scripts/verificar_controle_orcamento.py

"Em andamento" é o que ainda NÃO está comprometido no EBS mas já está em
curso — uma PO aguardando aprovação, por exemplo. É digitado na tela, e
sincronizar com o EBS não pode apagá-lo.

O saldo da tela é um só: disponível = orçamento − comprometido − em
andamento − realizado. As três parcelas descontam dele.

Roda contra um SQLite temporário.
"""
from __future__ import annotations

import os
import sqlite3
import sys
import tempfile
from decimal import Decimal
from pathlib import Path

RAIZ = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(RAIZ))

from fastapi import HTTPException  # noqa: E402

_TMP = Path(tempfile.mkdtemp())
os.environ.setdefault("DATABASE_URL", f"sqlite:///{_TMP/'portal.db'}")
os.environ.setdefault("PORTAL_SESSION_SECRET", "verificacao-local-sem-valor")
os.environ["ORCAMENTO_EXEC_DATABASE_URL"] = f"sqlite:///{_TMP/'orcamento_exec.db'}"

import db.orcamento_exec as dbo  # noqa: E402
import routers.controle_orcamento_exec as co  # noqa: E402
from sqlalchemy import select  # noqa: E402

falhas: list[str] = []
feitos = 0


def checar(cond, descricao):
    global feitos
    feitos += 1
    print(("  ok   " if cond else "  FALHA ") + descricao)
    if not cond:
        falhas.append(descricao)


dbo.init_db()
P = dbo.BudgetProject

print("\n[1] A coluna existe e sai na API")
with dbo.SessionLocal.begin() as s:
    s.add(P(code="PRJ-1", name="Troca de coletores", kind="CAPEX",
            approved_budget=Decimal("100000"), committed=Decimal("30000"),
            realized=Decimal("20000"), a_realizar=Decimal("50000"),
            em_andamento=Decimal("15000")))
with dbo.SessionLocal() as s:
    linha = s.scalar(select(P).where(P.code == "PRJ-1"))
    saida = co._dict(linha)
checar(saida["em_andamento"] == 15000.0,
       f"em_andamento vai para a tela ({saida.get('em_andamento')})")
checar(saida["a_realizar"] == 50000.0, "a_realizar continua vindo do EBS")
checar("em_andamento" in co._CAMPOS and co._CAMPOS["em_andamento"] == "em_andamento",
       "o campo é gravável pela tela")

print("\n[2] Em andamento é editável; a realizar não")
corpo = co.ProjetoIn(em_andamento=12500.9)
checar(corpo.em_andamento == Decimal("12500.90"),
       f"aceita o valor da tela ({corpo.em_andamento})")
# Mesma régua dos outros campos de dinheiro do módulo: nada de negativo.
try:
    co.ProjetoIn(em_andamento=-1)
    checar(False, "valor negativo é recusado")
except Exception:
    checar(True, "valor negativo é recusado")
try:
    co.ProjetoIn(a_realizar=10)
    checar(False, "a_realizar NÃO pode ser enviado pela tela")
except Exception:
    checar(True, "a_realizar NÃO pode ser enviado pela tela")
checar(co._PADRAO["em_andamento"] == Decimal("0"),
       "projeto novo nasce com em andamento zerado")

print("\n[3] Sincronizar com o EBS não apaga o que foi digitado")
with dbo.SessionLocal.begin() as s:
    alvo = s.scalar(select(P).where(P.code == "PRJ-1"))
    co._aplicar_ebs(alvo, {"saldo_inicial": "120000", "comprometido": "40000",
                           "reservados": "5000", "realizado": "25000",
                           "saldo_dia": "50000", "empresa": "RENNER"},
                    rates={})
with dbo.SessionLocal() as s:
    depois = s.scalar(select(P).where(P.code == "PRJ-1"))
checar(float(depois.committed) == 45000.0,
       f"comprometido = comprometido + reservados ({depois.committed})")
checar(float(depois.a_realizar) == 50000.0, "a_realizar veio do saldo do dia")
checar(float(depois.em_andamento) == 15000.0,
       f"em andamento sobreviveu à sincronização ({depois.em_andamento})")

print("\n[4] Banco que já existe ganha a coluna sem perder dado")
alvo_db = Path(str(dbo.DATABASE_URL).replace("sqlite:///", ""))
con = sqlite3.connect(alvo_db)
con.execute("ALTER TABLE budget_projects DROP COLUMN em_andamento")
con.commit()
cols = {c[1] for c in con.execute("PRAGMA table_info(budget_projects)")}
checar("em_andamento" not in cols, "banco antigo montado, sem a coluna")
antes = con.execute("SELECT COUNT(*) FROM budget_projects").fetchone()[0]
con.close()

dbo._ready = False
dbo._ensure_coluna_a_realizar()
con = sqlite3.connect(alvo_db)
cols = {c[1] for c in con.execute("PRAGMA table_info(budget_projects)")}
checar("em_andamento" in cols, "a migração acrescentou a coluna")
checar(con.execute("SELECT COUNT(*) FROM budget_projects").fetchone()[0] == antes,
       f"nenhuma linha se perdeu ({antes})")
zerados = con.execute("SELECT COUNT(*) FROM budget_projects "
                      "WHERE em_andamento IS NULL OR em_andamento = 0").fetchone()[0]
checar(zerados == antes, "as linhas antigas ficam com em andamento zerado")
con.close()
dbo._ensure_coluna_a_realizar()
checar(True, "rodar a migração de novo não quebra")

print("\n[5] A conta da tela fecha")
# Regra da área: comprometido, em andamento e realizado descontam do
# disponível. "A realizar" não é mais mostrado. O em andamento foi zerado
# pela migração da seção [4], então é digitado de novo — que é exatamente
# o que o usuário faria depois de subir a versão.
dbo._engine = None
dbo._factory = None
with dbo.SessionLocal.begin() as s:
    s.scalar(select(P).where(P.code == "PRJ-1")).em_andamento = Decimal("15000")
with dbo.SessionLocal() as s:
    d = co._dict(s.scalar(select(P).where(P.code == "PRJ-1")))
# Depois da sincronização da seção [3]: orçamento 120.000, comprometido
# 45.000 (40.000 + 5.000 reservados), realizado 25.000, em andamento 15.000.
disponivel = (d["orcamento"] - d["comprometido"] - d["em_andamento"]
              - d["realizado"])
checar(d["orcamento"] == 120000.0 and d["comprometido"] == 45000.0
       and d["realizado"] == 25000.0,
       f"parcelas conferem ({d['orcamento']}, {d['comprometido']}, {d['realizado']})")
checar(disponivel == 35000.0,
       f"disponível = 120.000 − 45.000 − 15.000 − 25.000 = {disponivel}")
# Cada parcela desconta: subir qualquer uma derruba o disponível na mesma medida.
for campo, quanto in (("comprometido", 1000.0), ("em_andamento", 2000.0),
                      ("realizado", 3000.0)):
    alterado = dict(d)
    alterado[campo] += quanto
    novo = (alterado["orcamento"] - alterado["comprometido"]
            - alterado["em_andamento"] - alterado["realizado"])
    checar(abs((disponivel - novo) - quanto) < 0.001,
           f"+{quanto:.0f} em {campo} tira {quanto:.0f} do disponível")

print("\n[Câmbio] A taxa informada na tela converte ARS/UYU para BRL")
# Sem isto, projeto da Argentina ficava com o valor no peso: a variável de
# ambiente está zerada no servidor e a URL de cotação ao vivo não é
# alcançável da rede interna. A tela é a única fonte que quem opera controla.
import asyncio as _asyncio  # noqa: E402
import io as _io  # noqa: E402

import routers.controle_orcamento_exec as _coe  # noqa: E402
from db.orcamento_exec import (  # noqa: E402
    BudgetProject as _BP, OpexItem as _OI, SessionLocal as _SL,
    ensure_db as _ensure, gravar_cambio as _grava, ler_cambio as _le,
)


class _Req:
    headers: dict = {}
    client = None


_coe._exigir = lambda req, acao, registro="", detalhe="": {"username": "verificacao"}
_coe.check_rate_limit = lambda req, nome=None: None
_coe.registrar_acesso = lambda *a, **k: None

_LINHA_AR = {"nro_projeto": "AR-1", "empresa": "RENNER ARGENTINA S.A.",
             "saldo_inicial": "1000000", "comprometido": "0", "reservados": "0"}

_p = _BP(code="AR-1", name="x")
_aviso, _ = _coe._aplicar_ebs(_p, dict(_LINHA_AR))
checar(float(_p.approved_budget) == 1000000.0,
       f"sem taxa, o valor fica no peso ({_p.approved_budget})")
checar("sem cotação ARS" in _aviso, f"e a tela é avisada do porquê ({_aviso[:40]}…)")

_grava({"ARS": 0.0055}, "verificacao")
_p2 = _BP(code="AR-1", name="x")
_aviso2, _ = _coe._aplicar_ebs(_p2, dict(_LINHA_AR))
checar(float(_p2.approved_budget) == 5500.0,
       f"com a taxa informada, converte (1.000.000 x 0,0055 = {_p2.approved_budget})")
checar(not _aviso2, "e não sobra aviso")

_p3 = _BP(code="BR-1", name="x")
_coe._aplicar_ebs(_p3, {"nro_projeto": "BR-1", "empresa": "LOJAS RENNER S.A.",
                        "saldo_inicial": "1000000", "comprometido": "0", "reservados": "0"})
checar(float(_p3.approved_budget) == 1000000.0, "empresa do Brasil não é convertida")

_lido = _coe.cambio_ler(_Req())
checar(_lido["cambio"]["ARS"]["fonte"] == "tela",
       f"a tela diz de onde a taxa veio ({_lido['cambio']['ARS']['fonte']})")
checar(_lido["cambio"]["UYU"]["fonte"] == "nenhuma", "e diz quando não há taxa")

try:
    _coe.cambio_gravar(_coe.CambioIn(UYU=5500), _Req())
    checar(False, "taxa absurda deveria ser recusada")
except HTTPException as _e:
    checar(_e.status_code == 422, f"1 peso valendo 5.500 reais é recusado ({_e.status_code})")

_coe.cambio_gravar(_coe.CambioIn(ARS=0), _Req())
checar(_le()["ARS"]["valor"] == 0.0, "zero apaga a taxa e volta ao ambiente/ao vivo")

print("\n[OPEX] Planilha modelo e importação")


async def _corpo(resp):
    partes = []
    async for pedaco in resp.body_iterator:
        partes.append(pedaco if isinstance(pedaco, bytes) else pedaco.encode())
    return b"".join(partes)


from openpyxl import load_workbook as _load  # noqa: E402

_bytes = _asyncio.run(_corpo(_coe.opex_modelo(_Req())))
_wb = _load(_io.BytesIO(_bytes))
checar("OPEX" in _wb.sheetnames and "Instruções" in _wb.sheetnames,
       f"o modelo tem as duas abas ({_wb.sheetnames})")
_cab = [c.value for c in _wb["OPEX"][1]]
checar(len(_cab) == 31, f"7 colunas fixas + 12 orçado + 12 realizado ({len(_cab)})")
checar(_cab[0].startswith("País") and _cab[1].startswith("Ano"),
       f"país e ano vêm primeiro e marcados ({_cab[:2]})")


class _Arq:
    def __init__(self, dados: bytes, nome: str = "opex.xlsx"):
        self.filename = nome
        self._dados = dados

    async def read(self) -> bytes:
        return self._dados


def _enviar(wb, **kw):
    buf = _io.BytesIO()
    wb.save(buf)
    return _asyncio.run(_coe.opex_importar(_Req(), _Arq(buf.getvalue()), **kw))


# O modelo que o portal gera tem de voltar pela importação sem retoque: o
# cabeçalho escreve "País *", e o asterisco já derrubou essa volta uma vez.
_previa = _enviar(_wb, dry_run=True)
checar(_previa["lidas"] == 2,
       f"o próprio modelo volta pela importação ({_previa['lidas']} linha(s))")
checar(_previa["dry_run"] is True and not _previa["avisos"],
       f"prévia sem aviso ({_previa['avisos']})")
_ensure()
with _SL() as _s:
    checar(_s.query(_OI).count() == 0, "a prévia NÃO grava nada")

_gravou = _enviar(_wb, dry_run=False)
with _SL() as _s:
    checar(_gravou["incluidas"] == 2 and _s.query(_OI).count() == 2,
           f"a confirmação grava ({_gravou['incluidas']})")

_de_novo = _enviar(_wb, dry_run=False, substituir=True)
with _SL() as _s:
    checar(_s.query(_OI).count() == 2,
           f"reenviar com substituir não soma em dobro ({_s.query(_OI).count()})")
checar(_de_novo["apagadas"] == 2, f"e diz quantas trocou ({_de_novo['apagadas']})")

# Cabeçalho digitado à mão, número em pt-BR e linha inválida.
from openpyxl import Workbook as _WB  # noqa: E402

_mao = _WB()
_ws = _mao.active
_ws.title = "OPEX"
_MES = ("jan", "fev", "mar", "abr", "mai", "jun", "jul", "ago", "set", "out", "nov", "dez")
_ws.append(["PAIS", "ano", "bu", "Fornecedor", "conta contabil", "descricao da conta",
            "tipo de despesa"] + [f"orçado {m}" for m in _MES]
           + [f"REALIZADO {m.upper()}" for m in _MES])
_ws.append(["br", 2026, "Infra", "Um Ltda", "", "", "Manutenção", "1.234,56"] + [""] * 23)
_ws.append([])
_ws.append(["XX", 2026, "", "País inválido"] + [""] * 27)
_ws.append(["uy", 2026, "Infra", "Dos S.A.", "6110", "Servicios", "", "R$ 900,10"] + [""] * 23)

_livre = _enviar(_mao, dry_run=True)
checar(_livre["lidas"] == 2,
       f"cabeçalho sem acento e em caixa diferente é aceito ({_livre['lidas']})")
_por_linha = {x["linha"]: x for x in _livre["previa"]}
checar(_por_linha[2]["total_orcado"] == 1234.56,
       f"'1.234,56' vira 1234.56 ({_por_linha[2]['total_orcado']})")
checar(_por_linha[5]["total_orcado"] == 900.10,
       f"'R$ 900,10' vira 900.10 ({_por_linha[5]['total_orcado']})")
checar(any("XX" in a for a in _livre["avisos"]),
       f"país inválido vira aviso, não erro ({_livre['avisos']})")
checar(_por_linha[2]["moeda"] == "BRL" and _por_linha[5]["moeda"] == "UYU",
       "cada país fica na sua moeda, sem conversão no OPEX")

_sem_coluna = _WB()
_sem_coluna.active.title = "OPEX"
_sem_coluna.active.append(["Fornecedor", "BU"])
try:
    _enviar(_sem_coluna, dry_run=True)
    checar(False, "planilha sem País/Ano deveria ser recusada")
except HTTPException as _e:
    checar(_e.status_code == 422 and "País" in _e.detail,
           f"planilha fora do modelo é recusada dizendo o que falta ({_e.detail[:48]}…)")


print(f"\n{feitos - len(falhas)} de {feitos} verificações passaram.")
if falhas:
    print("Falhas:\n  - " + "\n  - ".join(falhas))
    sys.exit(1)
print("Controle de Orçamento íntegro.")
