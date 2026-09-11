"""Orçamento de Manutenção — reparo de coletores e SLEDs do SPARE.

Substitui a planilha `Manutenção.xlsx`. O contrato completo (modelo, regras
de normalização, regra dos 60 %, agregados do painel e endpoints) está em
`docs/ORCAMENTO_MANUTENCAO.md` — este arquivo só o implementa.

Banco PRÓPRIO (`db/orcamento_manutencao.py`). Exige sessão do portal e o
módulo `orcamento_manutencao` liberado: `view` lê, `create` inclui, `edit`
altera, `export` exporta, `admin` importa planilha, exclui e configura.
"""
from __future__ import annotations

import io
import json
import logging
import numbers
import re
import unicodedata
from collections import Counter
from datetime import date, datetime
from pathlib import Path
from types import SimpleNamespace
from typing import Any, Optional

from fastapi import APIRouter, File, Form, HTTPException, Request, UploadFile
from fastapi.responses import JSONResponse, StreamingResponse
from pydantic import BaseModel, ConfigDict
from sqlalchemy import case, delete, func, insert, or_, select, update
from sqlalchemy.exc import IntegrityError

import db.orcamento_manutencao as db
from core.security import check_rate_limit, require_permission

_log = logging.getLogger("orcamento_manutencao")

router = APIRouter(prefix="/api/orcamento-manutencao", tags=["Orçamento Manutenção"],
                   include_in_schema=False)

MODULO = "orcamento_manutencao"
R = db.Reparo

# Categorias com rótulo canônico (3.1); qualquer outra vai por família.
CATEGORIAS_CANONICAS = {
    "coletor": ("Coletor", "COLETOR"),
    # A planilha do fornecedor chama o Coletor pelo fabricante (8.2).
    "coletor - bluebird": ("Coletor", "COLETOR"),
    "coletor bluebird": ("Coletor", "COLETOR"),
    "coletor hf550x": ("Coletor HF550X", "COLETOR"),
    "coletor s70": ("Coletor S70", "COLETOR"),
    "sled rfid": ("Sled RFID", "SLED"),
    "sled rfr901": ("Sled RFR901", "SLED"),
}


def _exigir(req: Request, acao: str) -> dict:
    db.ensure_db()
    return require_permission(req, MODULO, acao)


def _usuario(sd: dict) -> str:
    return str(sd.get("login") or sd.get("username") or "")[:80]


# ── Utilitários de texto/número ─────────────────────────────────────────
def _vazio(v) -> bool:
    if v is None:
        return True
    try:
        if v != v:  # NaN / NaT
            return True
    except Exception:  # noqa: BLE001
        pass
    return isinstance(v, str) and not v.strip()


def _numero_puro(v) -> bool:
    return isinstance(v, numbers.Number) and not isinstance(v, bool)


def _texto(v, tam: Optional[int] = None) -> str:
    if _vazio(v):
        return ""
    if _numero_puro(v) and float(v).is_integer():
        v = int(v)
    t = " ".join(str(v).replace("\xa0", " ").split())
    return t[:tam] if tam else t


def _sem_acento(v) -> str:
    """Minúsculas, sem acento, espaços únicos — base de toda comparação."""
    t = unicodedata.normalize("NFKD", _texto(v))
    return "".join(c for c in t if not unicodedata.combining(c)).lower()


def _inteiro(v) -> Optional[int]:
    if _vazio(v):
        return None
    if _numero_puro(v):
        return int(v)
    try:
        return int(float(_texto(v).replace(",", ".")))
    except ValueError:
        return None


def _numero(v) -> Optional[float]:
    """Número em formato livre ("1.234,56", "10,4%"); None quando não é número."""
    if _vazio(v):
        return None
    if _numero_puro(v):
        return float(v)
    t = _texto(v).replace("r$", "").replace("R$", "").replace(" ", "")
    pct = t.endswith("%")
    t = t.rstrip("%")
    if "," in t and "." in t:
        t = t.replace(".", "").replace(",", ".")
    elif "," in t:
        t = t.replace(",", ".")
    try:
        n = float(t)
    except ValueError:
        return None
    return n / 100 if pct else n


def _texto_rma(v) -> str:
    """RMA sem nenhum espaço (a planilha traz `\\xa0`)."""
    if _vazio(v):
        return ""
    if _numero_puro(v) and float(v).is_integer():
        return str(int(v))
    return "".join(str(v).replace("\xa0", " ").split())[:40]


# ── Normalização (seção 3) — funções puras ──────────────────────────────
def normalizar_categoria(texto) -> tuple[str, str]:
    original = _texto(texto, 40)
    chave = _sem_acento(original)
    if chave in CATEGORIAS_CANONICAS:
        return CATEGORIAS_CANONICAS[chave]
    if "coletor" in chave:
        return original.title(), "COLETOR"
    if "sled" in chave:
        return original.title(), "SLED"
    return original, "OUTRO"


def normalizar_status(texto) -> tuple[str, bool]:
    """Devolve (status canônico, reconhecido). Aceita o canônico e o texto humano."""
    chave = _sem_acento(texto).replace("_", " ")
    if chave.startswith("aprovado"):
        return "APROVADO", True
    if chave.startswith("reprovado"):
        return "REPROVADO", True
    if chave.startswith("aguardando aprovacao"):
        return "AGUARDANDO_APROVACAO", True
    if chave.startswith("aguardando orcamento"):
        return "AGUARDANDO_ORCAMENTO", True
    if chave.startswith("validando"):
        return "VALIDANDO_ORCAMENTO", True
    # 8.2: o fornecedor fecha o reparo sem cobrar — decidido, e sem custo.
    if chave.startswith("bonificado") or chave.startswith("garantia"):
        return "APROVADO", True
    # Faturado só acontece depois de aprovado: a nota já saiu.
    if chave.startswith("faturado"):
        return "APROVADO", True
    return "AGUARDANDO_ORCAMENTO", False


def normalizar_tipo(texto) -> tuple[str, bool]:
    chave = _sem_acento(texto)
    if "contrato" in chave:
        return "CONTRATO", True
    if "avulsa" in chave or "po extra" in chave:
        return "AVULSA", True
    return "CONTRATO", False


def normalizar_status_retorno(texto) -> str:
    return "DEVOLVIDO" if "devolvid" in _sem_acento(texto) else "EM_MANUTENCAO"


def normalizar_empresa(texto) -> str:
    chave = _sem_acento(texto).upper()
    for nome in db.EMPRESAS:
        if nome in chave:
            return nome
    return ""


def interpretar_orcamento(valor) -> tuple[float, bool]:
    """(orcamento, garantia). "Garantia" → (0, True); "R$ -" → (0, False);
    outro texto não numérico levanta ValueError com o motivo."""
    if _vazio(valor):
        return 0.0, False
    if _numero_puro(valor):
        return round(float(valor), 2), False
    chave = _sem_acento(valor)
    if "garantia" in chave:
        return 0.0, True
    compacto = chave.replace(" ", "")
    if compacto.replace("r$", "").strip("-.,") == "":
        return 0.0, False
    n = _numero(valor)
    if n is None:
        raise ValueError(f"orçamento não numérico: {_texto(valor)!r}")
    return round(n, 2), False


_RE_MES = re.compile(r"^(\d{4})-(\d{2})")


def mes_referencia(valor) -> Optional[str]:
    """datetime/date/Timestamp → "AAAA-MM"; texto "AAAA-MM" (ou ISO) mantém; outro → None."""
    if _vazio(valor):
        return None
    if hasattr(valor, "year") and hasattr(valor, "month"):
        try:
            return f"{int(valor.year):04d}-{int(valor.month):02d}"
        except (TypeError, ValueError):
            return None
    m = _RE_MES.match(_texto(valor))
    if m and 1 <= int(m.group(2)) <= 12:
        return f"{m.group(1)}-{m.group(2)}"
    return None


def data_celula(valor) -> Optional[date]:
    """Célula de data → `date`; texto "AAAA-MM-DD"/"DD/MM/AAAA" também; outro → None."""
    if _vazio(valor):
        return None
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    if hasattr(valor, "to_pydatetime"):  # pandas.Timestamp
        try:
            return valor.to_pydatetime().date()
        except (TypeError, ValueError):
            return None
    texto = _texto(valor)[:10]
    for formato in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(texto, formato).date()
        except ValueError:
            continue
    return None


# ── 8.2 — o status da planilha do fornecedor carrega mais de uma coisa ──
MESES_PT = {
    "janeiro": 1, "fevereiro": 2, "marco": 3, "abril": 4, "maio": 5, "junho": 6,
    "julho": 7, "agosto": 8, "setembro": 9, "outubro": 10, "novembro": 11, "dezembro": 12,
}
_RE_MES_PT = re.compile(r"\b(" + "|".join(MESES_PT) + r")\b")
_RE_ANO = re.compile(r"\b(20\d{2})\b")


def meses_por_extenso(texto) -> list[int]:
    """Meses citados por extenso, na ordem em que aparecem. A comparação é sem
    acento, então "MARÇO" e "marco" valem o mesmo."""
    return [MESES_PT[m] for m in _RE_MES_PT.findall(_sem_acento(texto))]


def mes_por_extenso(texto) -> Optional[str]:
    """Primeiro "<mês por extenso> ... <ano>" do texto → "AAAA-MM".
    "OUTUBRO / NOVEMBRO 2025" vale o primeiro mês: "2025-10"."""
    meses = meses_por_extenso(texto)
    anos = _RE_ANO.findall(_sem_acento(texto))
    return f"{anos[0]}-{meses[0]:02d}" if meses and anos else None


def interpretar_status_fornecedor(texto) -> tuple:
    """Lê o texto de status do fornecedor (8.2), que junta até quatro coisas:
    "APROVADO - PO EXTRA CAMICADO - SETEMBRO 2026" é status APROVADO, tipo
    AVULSA, empresa CAMICADO e mês 2026-09.

    Devolve `(status, tipo|None, mes|None, empresa|None, garantia, zerar_orcamento)`.
    `None` quer dizer "o texto não fala disso" — quem chama mantém o que já tem.
    Sem mês, sem tipo e sem empresa o resultado é o de `normalizar_status`
    sozinho, então `Aprovado Via Contrato` da base antiga segue APROVADO com
    tipo CONTRATO e sem mês, como sempre foi.
    """
    chave = _sem_acento(texto)
    status, _ = normalizar_status(texto)
    tipo, tipo_ok = normalizar_tipo(texto)
    garantia = chave.startswith("garantia")
    zerar = garantia or chave.startswith("bonificado")
    return (status, tipo if tipo_ok else None, mes_por_extenso(texto),
            normalizar_empresa(texto) or None, garantia, zerar)


def calcular(r, limiar: float) -> bool:
    """Recalcula `percentual`/`avaliacao` e aplica a reprovação automática (3.5).
    Serve para ORM e para o SimpleNamespace da importação. Devolve se algo mudou."""
    antes = (r.percentual, r.avaliacao, r.status, r.status_original)
    orc = float(r.orcamento or 0)
    vc = float(r.valor_compra) if r.valor_compra is not None else 0.0
    perc = None
    if vc > 0:
        perc = orc / vc
        r.percentual = round(perc, 4)
        r.avaliacao = "DENTRO" if perc <= limiar else "FORA"
    else:
        r.percentual = None
        r.avaliacao = ""
    # Garantia é reparo sem custo e já resolvido: entra como aprovado, e a
    # regra dos 60 % nem chega a olhar para ele.
    if r.garantia:
        if r.status != "APROVADO":
            r.status = "APROVADO"
            r.status_original = "Aprovado — garantia"
        return antes != (r.percentual, r.avaliacao, r.status, r.status_original)
    # RFR901, HF550 e S70 podem ser aprovados mesmo acima dos 60 %: ficam de
    # fora da reprovação automática e seguem pendentes para decisão manual.
    if r.avaliacao == "FORA" and r.status in db.STATUS_PENDENTES and not _isento_60(r.categoria):
        r.status = "REPROVADO"
        pct = f"{perc * 100:.1f}".replace(".", ",")
        lim = f"{limiar * 100:g}".replace(".", ",")
        r.status_original = f"Reprovado automaticamente ({pct} % > {lim} %)"[:60]
    return antes != (r.percentual, r.avaliacao, r.status, r.status_original)


# Modelos que podem ser aprovados acima dos 60 %: a regra automática não os
# reprova; a aprovação é decidida manualmente.
_ISENTOS_60 = ("RFR901", "HF550", "S70")


def _isento_60(categoria) -> bool:
    chave = _sem_acento(categoria).replace(" ", "")
    return any(m.lower() in chave for m in _ISENTOS_60)


def _padrao_categoria(config: dict, categoria: str) -> Optional[float]:
    chave = _sem_acento(categoria)
    for nome, valor in (config.get("valor_compra_padrao") or {}).items():
        if _sem_acento(nome) == chave:
            v = _numero(valor)
            if v and v > 0:
                return round(v, 2)
    return None


def _valor_compra_para(r, config: dict, ebs_custo: Optional[float] = None,
                       razao: Optional[float] = None,
                       sobrescrever_manual: bool = False) -> tuple[Optional[float], str]:
    """Ordem 3.4: EBS → planilha → padrão. Uma fonte mais fraca nunca rebaixa
    um valor já gravado por fonte mais forte; MANUAL só cede se pedido."""
    atual = float(r.valor_compra) if r.valor_compra is not None else None
    fonte = r.valor_compra_fonte or ""
    if atual and fonte == "MANUAL" and not sobrescrever_manual:
        return atual, "MANUAL"
    if ebs_custo is not None and ebs_custo > 0:
        return round(float(ebs_custo), 2), "EBS"
    if atual and fonte == "EBS":
        return atual, "EBS"
    orc = float(r.orcamento or 0)
    if razao is not None and orc > 0 and razao > 0:
        return round(orc / razao, 2), "PLANILHA"
    if atual and fonte == "PLANILHA":
        return atual, "PLANILHA"
    padrao = _padrao_categoria(config, r.categoria)
    if padrao:
        return padrao, "PADRAO"
    if atual and fonte == "MANUAL":
        return atual, "MANUAL"
    return None, ""


# ── EBS ─────────────────────────────────────────────────────────────────
def _ebs_buscar(serie: str) -> dict:
    """Chamada crua ao EBS (isolada para poder ser substituída em testes)."""
    from integracoes.ebs_service import search_one
    from routers.public_assets import _auth
    try:
        res = search_one(_auth(), serie)
    except PermissionError:
        res = {"encontrado": False, "erro": "Sessão EBS expirada"}
    if "expirad" in _sem_acento(res.get("erro") if isinstance(res, dict) else ""):
        res = search_one(_auth(True), serie)
    return res if isinstance(res, dict) else {"encontrado": False, "erro": "resposta inválida do EBS"}


def _consultar_ebs(serie: str) -> dict:
    """Nunca levanta: {"encontrado", "custo", "empresa", "erro"}."""
    saida = {"encontrado": False, "custo": None, "empresa": "", "erro": ""}
    try:
        res = _ebs_buscar(serie)
        if not res.get("encontrado"):
            saida["erro"] = _texto(res.get("erro") or "Ativo não encontrado no EBS", 200)
            return saida
        empresa = res.get("empresa") or ""
        try:
            from routers.public_assets import _company
            empresa = _company(empresa)
        except Exception:  # noqa: BLE001
            pass
        saida.update(encontrado=True, custo=_numero(res.get("custo_asset")),
                     empresa=normalizar_empresa(empresa))
        if not saida["custo"]:
            saida["erro"] = "EBS sem custo do ativo"
    except Exception as exc:  # noqa: BLE001 — falha do EBS nunca derruba a inclusão
        _log.warning("EBS falhou para série %s: %s", serie, exc)
        saida["erro"] = _texto(f"Falha na consulta ao EBS: {exc}", 200)
    return saida


def _aplicar_ebs(r, res: dict, config: dict, sobrescrever_manual: bool) -> bool:
    """Grava o resultado do EBS na linha. Devolve se valor/empresa mudaram."""
    r.ebs_consultado_em = db.utcnow()
    r.ebs_erro = (res.get("erro") or "")[:200]
    mudou = False
    if res.get("encontrado"):
        if res.get("empresa") and r.empresa != res["empresa"]:
            r.empresa = res["empresa"]
            mudou = True
        valor, fonte = _valor_compra_para(r, config, ebs_custo=res.get("custo"),
                                          sobrescrever_manual=sobrescrever_manual)
        if (valor, fonte) != (r.valor_compra, r.valor_compra_fonte):
            r.valor_compra, r.valor_compra_fonte = valor, fonte
            mudou = True
    calcular(r, config["limiar_percentual"])
    return mudou


# ── Entrada ─────────────────────────────────────────────────────────────
class ReparoIn(BaseModel):
    """Corpo do POST. Números aceitam texto ("389,29", "Garantia") — a tela
    pode mandar campos a mais, que são ignorados."""
    model_config = ConfigDict(extra="ignore")
    rma: str
    serie: str
    categoria: str
    loja: Any = None
    empresa: str = ""
    orcamento: Any = None
    garantia: Optional[bool] = None
    valor_compra: Any = None
    status: str = ""
    tipo_manutencao: str = ""
    status_retorno: str = ""
    ano: Any = None
    mes_referencia: Any = None
    ano_devolucao: Any = None
    lote_prime: str = ""
    qtde: Any = None
    observacao: str = ""


class ReparoPatch(BaseModel):
    """Corpo do PUT: só o que vier é alterado."""
    model_config = ConfigDict(extra="ignore")
    rma: Optional[str] = None
    serie: Optional[str] = None
    categoria: Optional[str] = None
    loja: Any = None
    empresa: Optional[str] = None
    orcamento: Any = None
    garantia: Optional[bool] = None
    valor_compra: Any = None
    status: Optional[str] = None
    tipo_manutencao: Optional[str] = None
    status_retorno: Optional[str] = None
    ano: Any = None
    mes_referencia: Any = None
    ano_devolucao: Any = None
    lote_prime: Optional[str] = None
    qtde: Any = None
    observacao: Optional[str] = None


class ConfigIn(BaseModel):
    model_config = ConfigDict(extra="ignore")
    cota_mensal: Optional[dict[str, Any]] = None
    limiar_percentual: Any = None
    valor_compra_padrao: Optional[dict[str, Any]] = None


def _status_entrada(texto: str) -> tuple[str, str]:
    """(canônico, status_original) para o corpo; texto não reconhecido → 422."""
    status, ok = normalizar_status(texto)
    if texto and not ok:
        raise HTTPException(422, f"Status não reconhecido: {texto!r}")
    return status, (_texto(texto, 60) or db.STATUS_ROTULOS[status])


def _valor_manual(v) -> Optional[float]:
    if _vazio(v):
        return None
    n = _numero(v)
    if n is None or n < 0:
        raise HTTPException(422, f"Valor de compra inválido: {_texto(v)!r}")
    return round(n, 2)


def _mes_entrada(v) -> Optional[str]:
    if _vazio(v):
        return None
    mes = mes_referencia(v)
    if not mes:
        raise HTTPException(422, f"Mês de referência inválido (use AAAA-MM): {_texto(v)!r}")
    return mes


# ── Filtros e serialização ──────────────────────────────────────────────
def _filtrar(stmt, ano=None, mes=None, familia=None, categoria=None, status=None,
             status_retorno=None, empresa=None, tipo_manutencao=None, q=None,
             min_reparos=None, lote=None):
    if min_reparos and int(min_reparos) > 1:
        # Reincidência (8.4): só as séries com N atendimentos ou mais. Uma
        # subconsulta agrupada, para valer igual na lista e na contagem.
        reincidentes = (select(R.serie).where(R.serie != "").group_by(R.serie)
                        .having(func.count(R.id) >= int(min_reparos)).scalar_subquery())
        stmt = stmt.where(R.serie.in_(reincidentes))
    if ano:
        stmt = stmt.where(R.ano == ano)
    if mes:
        stmt = stmt.where(R.mes_referencia == mes)
    if familia:
        stmt = stmt.where(R.familia == familia.strip().upper())
    if categoria:
        stmt = stmt.where(R.categoria == normalizar_categoria(categoria)[0])
    if status:
        stmt = stmt.where(R.status == normalizar_status(status)[0])
    if status_retorno:
        stmt = stmt.where(R.status_retorno == normalizar_status_retorno(status_retorno))
    if empresa:
        stmt = stmt.where(R.empresa == normalizar_empresa(empresa))
    if lote and lote.strip():
        # Casa por conteúdo e sem diferenciar maiúsculas (vale em SQLite e
        # Postgres): digitar "lote" traz todos que contêm "lote".
        stmt = stmt.where(func.lower(R.lote_prime).like(f"%{lote.strip().lower()}%"))
    if tipo_manutencao:
        stmt = stmt.where(R.tipo_manutencao == normalizar_tipo(tipo_manutencao)[0])
    if q and q.strip():
        like = f"%{q.strip()}%"
        stmt = stmt.where(or_(R.rma.like(like), R.serie.like(like.upper()),
                              R.lote_prime.like(like)))
    return stmt


def _ordenado(stmt):
    return stmt.order_by(R.mes_referencia.desc().nullslast(), R.id.desc())


def _mes_do_filtro(mes: Optional[str], ano: int) -> Optional[str]:
    """Aceita "AAAA-MM", "MM" ou "3"; devolve "AAAA-MM" ou None (ano inteiro)."""
    texto = str(mes or "").strip()
    if not texto or texto.lower() in ("todos", "ano"):
        return None
    if len(texto) >= 7 and texto[4] == "-" and texto[:4].isdigit() and texto[5:7].isdigit():
        numero = int(texto[5:7])
        return f"{int(texto[:4])}-{numero:02d}" if 1 <= numero <= 12 else None
    if texto.isdigit() and 1 <= int(texto) <= 12:
        return f"{ano}-{int(texto):02d}"
    return None


def _dinheiro(v) -> float:
    return round(float(v or 0), 2)


# ── 8.4 Reincidência — a série é o equipamento, o RMA é o atendimento ────
_CUSTO_APROVADO = func.coalesce(
    func.sum(case((R.status == "APROVADO", R.orcamento), else_=0)), 0)


def _agregado_series(s, series=None) -> dict:
    """`{série: (atendimentos, custo aprovado)}` numa consulta agregada só —
    nunca uma consulta por linha. `series=None` agrega a base inteira; uma
    lista agrega só aquelas, em blocos que cabem no limite de parâmetros."""
    saida: dict[str, tuple[int, float]] = {}
    base = select(R.serie, func.count(R.id), _CUSTO_APROVADO).where(R.serie != "").group_by(R.serie)
    if series is None:
        blocos = [base]
    else:
        lista = [x for x in dict.fromkeys(series) if x]
        blocos = [base.where(R.serie.in_(lista[i:i + 400])) for i in range(0, len(lista), 400)]
    for stmt in blocos:
        for serie, qtde, custo in s.execute(stmt).all():
            saida[serie] = (int(qtde), _dinheiro(custo))
    return saida


def _com_reincidencia(itens: list[dict], agregado: dict) -> list[dict]:
    for item in itens:
        qtde, custo = agregado.get(item.get("serie") or "", (1, 0.0))
        item["serie_reparos"], item["serie_custo"] = qtde, custo
    return itens


# ── 5.1 Painel ──────────────────────────────────────────────────────────
@router.get("/resumo")
def resumo(req: Request, ano: Optional[int] = None, mes: Optional[str] = None):
    """Painel. `mes` ("AAAA-MM" ou "01".."12") estreita cartões e quebras ao mês;
    as séries mensais continuam mostrando o ano inteiro, para dar contexto."""
    _exigir(req, "view")
    with db.SessionLocal() as s:
        config = db.ler_config(s)
        # Anos: da coluna `ano` e também dos meses de referência (na planilha
        # o mês de contrato pode cair no ano seguinte ao do registro).
        anos = {int(a) for (a,) in s.execute(select(R.ano).distinct()).all() if a}
        anos |= {int(m) for (m,) in s.execute(
            select(func.substr(R.mes_referencia, 1, 4)).distinct()
            .where(R.mes_referencia.isnot(None))).all() if m and str(m).isdigit()}
        if not anos:
            anos = {date.today().year}
        ano = ano or max(anos)

        mes_sel = _mes_do_filtro(mes, ano)
        escopo = (R.mes_referencia == mes_sel) if mes_sel else R.mes_referencia.like(f"{ano}-%")

        # ── Séries mensais (sempre o ano inteiro) ───────────────────────
        vazio = lambda: {"COLETOR": 0.0, "SLED": 0.0, "TOTAL": 0.0}  # noqa: E731
        vazio_tipo = lambda: {"CONTRATO": 0.0, "AVULSA": 0.0, "TOTAL": 0.0}  # noqa: E731
        meses = {}
        for m in range(1, 13):
            meses[f"{ano}-{m:02d}"] = {
                "consumo": vazio(), "reparados": vazio(),
                "reprovados_valor": vazio(), "reprovados_qtde": vazio(),
                "consumo_tipo": vazio_tipo(), "reparados_tipo": vazio_tipo(),
                "consumo_contrato": vazio(), "consumo_avulso": vazio(),
            }
        rows = s.execute(
            select(R.mes_referencia, R.familia, R.tipo_manutencao, R.status, func.count(R.id),
                   func.coalesce(func.sum(R.orcamento), 0),
                   func.coalesce(func.sum(R.valor_compra), 0))
            .where(R.mes_referencia.like(f"{ano}-%"), R.status.in_(("APROVADO", "REPROVADO")))
            .group_by(R.mes_referencia, R.familia, R.tipo_manutencao, R.status)
        ).all()
        for mes_ref, familia, tipo, status, qtde, soma_orc, soma_vc in rows:
            bloco = meses.get(mes_ref)
            if bloco is None:
                continue
            qtde = int(qtde)
            if status == "APROVADO":
                alvo_v, alvo_q, valor = bloco["consumo"], bloco["reparados"], float(soma_orc)
                # Aprovados separam contrato de avulso: a cota mensal é do contrato.
                if tipo in bloco["consumo_tipo"]:
                    bloco["consumo_tipo"][tipo] += valor
                    bloco["reparados_tipo"][tipo] += qtde
                # A mesma quebra por família das demais visões, por tipo.
                por_fam = bloco["consumo_avulso" if tipo == "AVULSA" else "consumo_contrato"]
                if familia in por_fam:
                    por_fam[familia] += valor
                por_fam["TOTAL"] += valor
                bloco["consumo_tipo"]["TOTAL"] += valor
                bloco["reparados_tipo"]["TOTAL"] += qtde
            else:
                alvo_v, alvo_q, valor = bloco["reprovados_valor"], bloco["reprovados_qtde"], float(soma_vc)
            if familia in alvo_v:
                alvo_v[familia] += valor
                alvo_q[familia] += qtde
            alvo_v["TOTAL"] += valor
            alvo_q["TOTAL"] += qtde

        lista_meses = []
        for mes_ref, bloco in meses.items():
            for chave in ("consumo", "reprovados_valor", "consumo_tipo",
                          "consumo_contrato", "consumo_avulso"):
                bloco[chave] = {k: _dinheiro(v) for k, v in bloco[chave].items()}
            for chave in ("reparados", "reprovados_qtde", "reparados_tipo"):
                bloco[chave] = {k: int(v) for k, v in bloco[chave].items()}
            lista_meses.append({"mes": mes_ref, **bloco})

        com_consumo = [b for b in lista_meses if b["consumo"]["TOTAL"] > 0]

        # ── Quebras do escopo (mês selecionado ou ano inteiro) ──────────
        categorias: dict[tuple, dict] = {}
        gerais = {f: {"consumo": 0.0, "reparados": 0} for f in db.FAMILIAS}
        por_tipo = {t: {"consumo": 0.0, "reparados": 0} for t in db.TIPOS_ROTULOS}
        totais = {"aprovados": 0, "aprovados_contrato": 0, "aprovados_avulsa": 0,
                  "reprovados": 0, "pendentes": 0, "garantia": 0,
                  "consumo": 0.0, "consumo_contrato": 0.0, "consumo_avulsa": 0.0,
                  "reprovados_valor": 0.0}
        for cat, fam, tipo, status, garantia, q, soma_orc, soma_vc in s.execute(
            select(R.categoria, R.familia, R.tipo_manutencao, R.status, R.garantia,
                   func.count(R.id), func.coalesce(func.sum(R.orcamento), 0),
                   func.coalesce(func.sum(R.valor_compra), 0))
            .where(escopo)
            .group_by(R.categoria, R.familia, R.tipo_manutencao, R.status, R.garantia)
        ).all():
            c = categorias.setdefault((cat, fam), {
                "categoria": cat, "familia": fam, "consumo": 0.0,
                "consumo_contrato": 0.0, "consumo_avulsa": 0.0, "reparados": 0,
                "reparados_contrato": 0, "reparados_avulsa": 0,
                "reprovados_qtde": 0, "reprovados_valor": 0.0,
            })
            q, valor_orc, valor_vc = int(q), float(soma_orc), float(soma_vc)
            sufixo = "contrato" if tipo == "CONTRATO" else "avulsa"
            if status == "APROVADO":
                c["consumo"] += valor_orc
                c["reparados"] += q
                c[f"consumo_{sufixo}"] += valor_orc
                c[f"reparados_{sufixo}"] += q
                gerais.setdefault(fam, {"consumo": 0.0, "reparados": 0})
                gerais[fam]["consumo"] += valor_orc
                gerais[fam]["reparados"] += q
                por_tipo.setdefault(tipo, {"consumo": 0.0, "reparados": 0})
                por_tipo[tipo]["consumo"] += valor_orc
                por_tipo[tipo]["reparados"] += q
                totais["aprovados"] += q
                totais[f"aprovados_{sufixo}"] += q
                totais["consumo"] += valor_orc
                totais[f"consumo_{sufixo}"] += valor_orc
                if garantia:
                    totais["garantia"] += q
            elif status == "REPROVADO":
                c["reprovados_valor"] += valor_vc
                c["reprovados_qtde"] += q
                totais["reprovados"] += q
                totais["reprovados_valor"] += valor_vc
            else:
                totais["pendentes"] += q

        lista_categorias = sorted(categorias.values(), key=lambda x: -x["consumo"])
        for c in lista_categorias:
            for chave in ("consumo", "consumo_contrato", "consumo_avulsa", "reprovados_valor"):
                c[chave] = _dinheiro(c[chave])
        for bloco in (gerais, por_tipo):
            for dados in bloco.values():
                dados["consumo"] = _dinheiro(dados["consumo"])
                dados["media"] = _dinheiro(dados["consumo"] / dados["reparados"]) if dados["reparados"] else 0.0
        for chave in ("consumo", "consumo_contrato", "consumo_avulsa", "reprovados_valor"):
            totais[chave] = _dinheiro(totais[chave])

        # ── Cartões financeiros ────────────────────────────────────────
        cota = _dinheiro((config.get("cota_mensal") or {}).get(str(ano)))
        if mes_sel:
            foco = next((b for b in lista_meses if b["mes"] == mes_sel), None)
        else:
            foco = com_consumo[-1] if com_consumo else None
        cota_em_uso = foco["mes"] if foco else mes_sel
        consumo_atual = foco["consumo"]["TOTAL"] if foco else 0.0
        consumo_contrato = foco["consumo_tipo"]["CONTRATO"] if foco else 0.0
        consumo_avulsa = foco["consumo_tipo"]["AVULSA"] if foco else 0.0

        totais["meses_com_consumo"] = len(com_consumo)
        total_ano = _dinheiro(sum(b["consumo"]["TOTAL"] for b in lista_meses))
        totais["total_ano"] = total_ano
        totais["media_mensal"] = _dinheiro(total_ano / len(com_consumo)) if com_consumo else 0.0
        totais["cota_anual"] = _dinheiro(cota * 12)
        # A cota mensal acompanha o contrato: o avulso é medido à parte.
        totais["percentual_cota_mes"] = round(consumo_contrato / cota, 4) if cota else None

        # Cards: retrato de agora, sem filtro de ano.
        ag_aprovacao = [
            {"categoria": cat, "familia": fam, "qtde": int(q), "valor": _dinheiro(v)}
            for cat, fam, q, v in s.execute(
                select(R.categoria, R.familia, func.count(R.id),
                       func.coalesce(func.sum(R.orcamento), 0))
                .where(R.status == "AGUARDANDO_APROVACAO")
                .group_by(R.categoria, R.familia).order_by(R.categoria)
            ).all()
        ]
        devolucao: dict[tuple, dict] = {}
        for cat, fam, status, q in s.execute(
            select(R.categoria, R.familia, R.status, func.count(R.id))
            .where(R.status_retorno == "EM_MANUTENCAO")
            .group_by(R.categoria, R.familia, R.status).order_by(R.categoria)
        ).all():
            d = devolucao.setdefault((cat, fam), {
                "categoria": cat, "familia": fam, "total": 0,
                "ag_manutencao": 0, "ag_orcamento": 0, "ag_aprovacao": 0, "reprovado": 0,
            })
            d["total"] += int(q)
            chave = {"APROVADO": "ag_manutencao", "AGUARDANDO_ORCAMENTO": "ag_orcamento",
                     "VALIDANDO_ORCAMENTO": "ag_orcamento",
                     "AGUARDANDO_APROVACAO": "ag_aprovacao", "REPROVADO": "reprovado"}.get(status)
            if chave:
                d[chave] += int(q)

    return {
        "ano": ano,
        "mes": mes_sel,
        "anos_disponiveis": sorted(anos),
        "meses_disponiveis": [
            {"mes": b["mes"], "tem_dado": b["consumo"]["TOTAL"] > 0 or b["reprovados_qtde"]["TOTAL"] > 0}
            for b in lista_meses
        ],
        "cota_mensal": cota,
        "cota_em_uso": cota_em_uso,
        "consumo_atual": consumo_atual,
        "consumo_atual_contrato": consumo_contrato,
        "consumo_atual_avulsa": consumo_avulsa,
        "residual": _dinheiro(cota - consumo_contrato),
        "total_investido": totais["consumo"] if mes_sel else total_ano,
        "limiar_percentual": config["limiar_percentual"],
        "meses": lista_meses,
        "gerais": gerais,
        "por_tipo": por_tipo,
        "aguardando_aprovacao": ag_aprovacao,
        "aguardando_devolucao": list(devolucao.values()),
        "categorias": lista_categorias,
        "totais": totais,
    }


# ── 5.2 Lista ───────────────────────────────────────────────────────────
@router.get("/reparos")
def listar(req: Request, ano: Optional[int] = None, mes: Optional[str] = None,
           familia: Optional[str] = None, categoria: Optional[str] = None,
           status: Optional[str] = None, status_retorno: Optional[str] = None,
           empresa: Optional[str] = None, tipo_manutencao: Optional[str] = None,
           q: Optional[str] = None, min_reparos: Optional[int] = None,
           lote: Optional[str] = None, limit: int = 100, offset: int = 0):
    """`min_reparos=2` deixa só o que é reincidente (8.4). Cada item leva
    `serie_reparos` e `serie_custo` — o histórico da série na base inteira."""
    _exigir(req, "view")
    limit = max(1, min(int(limit or 100), 1000))
    offset = max(0, int(offset or 0))
    filtros = dict(ano=ano, mes=mes, familia=familia, categoria=categoria, status=status,
                   status_retorno=status_retorno, empresa=empresa,
                   tipo_manutencao=tipo_manutencao, q=q, min_reparos=min_reparos, lote=lote)
    with db.SessionLocal() as s:
        total = s.scalar(_filtrar(select(func.count(R.id)), **filtros)) or 0
        itens = s.scalars(_ordenado(_filtrar(select(R), **filtros))
                          .limit(limit).offset(offset)).all()
        saida = [r.to_dict() for r in itens]
        _com_reincidencia(saida, _agregado_series(s, [r.serie for r in itens]))
        return {"total": int(total), "itens": saida}


# ── 5.3 Inclusão ────────────────────────────────────────────────────────
def _conflito(reparo_id: int) -> JSONResponse:
    return JSONResponse(status_code=409, content={"detail": "RMA já cadastrado", "id": reparo_id})


@router.post("/reparos", status_code=201)
def criar(body: ReparoIn, req: Request):
    sd = _exigir(req, "create")
    check_rate_limit(req, "api")
    usuario = _usuario(sd)

    rma = _texto_rma(body.rma)
    serie = _texto(body.serie, 60).upper()
    categoria, familia = normalizar_categoria(body.categoria)
    if not rma or not serie or not categoria:
        raise HTTPException(422, "RMA, série e categoria são obrigatórios.")
    try:
        orcamento, garantia = interpretar_orcamento(body.orcamento)
    except ValueError as exc:
        raise HTTPException(422, str(exc))
    status, status_original = _status_entrada(body.status or "")
    mes = _mes_entrada(body.mes_referencia)

    with db.SessionLocal() as s:
        existente = s.scalar(select(R.id).where(R.rma == rma))
        if existente:
            return _conflito(existente)
        config = db.ler_config(s)

    r = R(
        rma=rma, serie=serie, categoria=categoria, familia=familia,
        loja=_inteiro(body.loja), empresa=normalizar_empresa(body.empresa),
        orcamento=orcamento, garantia=bool(garantia or body.garantia),
        status=status, status_original=status_original,
        tipo_manutencao=normalizar_tipo(body.tipo_manutencao)[0],
        tipo_original=_texto(body.tipo_manutencao, 60),
        status_retorno=normalizar_status_retorno(body.status_retorno),
        mes_referencia=mes,
        ano=_inteiro(body.ano) or (int(mes[:4]) if mes else date.today().year),
        ano_devolucao=_inteiro(body.ano_devolucao),
        lote_prime=_texto(body.lote_prime, 200), qtde=_inteiro(body.qtde) or 1,
        observacao=(body.observacao or "").strip(),
        criado_por=usuario, atualizado_por=usuario, origem="PORTAL",
    )

    ebs = {"consultado": False, "erro": ""}
    manual = _valor_manual(body.valor_compra)
    if manual is not None:
        r.valor_compra, r.valor_compra_fonte = manual, "MANUAL"
    else:
        # EBS fora de transação: chamada de rede não segura o banco.
        res = _consultar_ebs(serie)
        ebs = {"consultado": True, "erro": res["erro"]}
        r.ebs_consultado_em = db.utcnow()
        r.ebs_erro = res["erro"][:200]
        if res["empresa"]:
            r.empresa = res["empresa"]
        r.valor_compra, r.valor_compra_fonte = _valor_compra_para(r, config, ebs_custo=res["custo"])
    calcular(r, config["limiar_percentual"])

    try:
        with db.SessionLocal.begin() as s:
            s.add(r)
            s.flush()
            saida = r.to_dict()
    except IntegrityError:
        with db.SessionLocal() as s:
            return _conflito(s.scalar(select(R.id).where(R.rma == rma)) or 0)
    return {"reparo": saida, "ebs": ebs}


# ── 5.4 Edição ──────────────────────────────────────────────────────────
@router.put("/reparos/{reparo_id}")
def atualizar(reparo_id: int, body: ReparoPatch, req: Request):
    sd = _exigir(req, "edit")
    check_rate_limit(req, "api")
    dados = body.model_dump(exclude_unset=True)
    with db.SessionLocal.begin() as s:
        r = s.get(R, reparo_id)
        if not r:
            raise HTTPException(404, "Reparo não encontrado.")
        config = db.ler_config(s)

        if "rma" in dados:
            novo = _texto_rma(dados["rma"])
            if not novo:
                raise HTTPException(422, "RMA não pode ficar vazio.")
            if novo != r.rma and s.scalar(select(R.id).where(R.rma == novo)):
                raise HTTPException(409, "RMA já cadastrado")
            r.rma = novo
        if "serie" in dados:
            r.serie = _texto(dados["serie"], 60).upper()
            if not r.serie:
                raise HTTPException(422, "Série não pode ficar vazia.")
        if "categoria" in dados:
            r.categoria, r.familia = normalizar_categoria(dados["categoria"])
            if not r.categoria:
                raise HTTPException(422, "Categoria não pode ficar vazia.")
        if "loja" in dados:
            r.loja = _inteiro(dados["loja"])
        if "empresa" in dados:
            r.empresa = normalizar_empresa(dados["empresa"])
        if "orcamento" in dados:
            try:
                r.orcamento, garantia = interpretar_orcamento(dados["orcamento"])
            except ValueError as exc:
                raise HTTPException(422, str(exc))
            if garantia:
                r.garantia = True
        if "garantia" in dados and dados["garantia"] is not None:
            r.garantia = bool(dados["garantia"])
        if "status" in dados and dados["status"] is not None:
            r.status, r.status_original = _status_entrada(dados["status"])
        if "tipo_manutencao" in dados and dados["tipo_manutencao"] is not None:
            r.tipo_manutencao = normalizar_tipo(dados["tipo_manutencao"])[0]
            r.tipo_original = _texto(dados["tipo_manutencao"], 60)
        if "status_retorno" in dados and dados["status_retorno"] is not None:
            r.status_retorno = normalizar_status_retorno(dados["status_retorno"])
        if "mes_referencia" in dados:
            r.mes_referencia = _mes_entrada(dados["mes_referencia"])
        if "ano" in dados and _inteiro(dados["ano"]):
            r.ano = _inteiro(dados["ano"])
        if "ano_devolucao" in dados:
            r.ano_devolucao = _inteiro(dados["ano_devolucao"])
        if "lote_prime" in dados and dados["lote_prime"] is not None:
            r.lote_prime = _texto(dados["lote_prime"], 200)
        if "qtde" in dados and _inteiro(dados["qtde"]):
            r.qtde = _inteiro(dados["qtde"])
        if "observacao" in dados and dados["observacao"] is not None:
            r.observacao = str(dados["observacao"]).strip()
        if "valor_compra" in dados:
            manual = _valor_manual(dados["valor_compra"])
            # Valor digitado vira MANUAL; nulo limpa e deixa a ordem 3.4 decidir.
            r.valor_compra, r.valor_compra_fonte = (manual, "MANUAL") if manual is not None else (None, "")

        r.valor_compra, r.valor_compra_fonte = _valor_compra_para(r, config)
        calcular(r, config["limiar_percentual"])
        r.atualizado_por = _usuario(sd)
        s.flush()
        return r.to_dict()


# ── 5.5 Exclusão ────────────────────────────────────────────────────────
@router.delete("/reparos/{reparo_id}")
def excluir(reparo_id: int, req: Request):
    _exigir(req, "admin")
    check_rate_limit(req, "api")
    with db.SessionLocal.begin() as s:
        r = s.get(R, reparo_id)
        if not r:
            raise HTTPException(404, "Reparo não encontrado.")
        s.delete(r)
    return {"ok": True}


# ── 5.6 EBS ─────────────────────────────────────────────────────────────
@router.post("/reparos/ebs-pendentes")
def ebs_pendentes(req: Request, limite: int = 200):
    """Completa valor de compra/empresa pelo EBS nas linhas sem um dos dois,
    as mais recentes primeiro. Não sobrescreve valor MANUAL."""
    sd = _exigir(req, "admin")
    check_rate_limit(req, "api")
    limite = max(1, min(int(limite or 200), 1000))
    with db.SessionLocal() as s:
        alvos = s.execute(
            select(R.id, R.serie)
            .where(or_(R.valor_compra.is_(None), R.empresa == ""))
            .order_by(R.id.desc()).limit(limite)
        ).all()
    # Rede primeiro, banco depois: nada de transação aberta durante as chamadas.
    resultados = [(rid, _consultar_ebs(serie)) for rid, serie in alvos]
    return _gravar_ebs(resultados, _usuario(sd), sobrescrever_manual=False)


@router.post("/reparos/{reparo_id}/ebs")
def ebs_um(reparo_id: int, req: Request):
    """Reconsulta explícita de uma linha — aqui o valor MANUAL cede (3.4)."""
    sd = _exigir(req, "edit")
    check_rate_limit(req, "api")
    with db.SessionLocal() as s:
        r = s.get(R, reparo_id)
        if not r:
            raise HTTPException(404, "Reparo não encontrado.")
        serie = r.serie
    saida = _gravar_ebs([(reparo_id, _consultar_ebs(serie))], _usuario(sd), sobrescrever_manual=True)
    with db.SessionLocal() as s:
        saida["reparo"] = s.get(R, reparo_id).to_dict()
    return saida


def _gravar_ebs(resultados: list, usuario: str, sobrescrever_manual: bool) -> dict:
    atualizados, falhas = 0, []
    with db.SessionLocal.begin() as s:
        config = db.ler_config(s)
        for rid, res in resultados:
            r = s.get(R, rid)
            if not r:
                continue
            if _aplicar_ebs(r, res, config, sobrescrever_manual):
                atualizados += 1
                r.atualizado_por = usuario
            if res.get("erro"):
                falhas.append({"id": rid, "erro": res["erro"]})
    return {"consultados": len(resultados), "atualizados": atualizados, "falhas": falhas}


# ── 5.7 Importação ──────────────────────────────────────────────────────
# Nomes de coluna que identificam a série em cada layout: a planilha histórica
# diz SÉRIE, a aba AVULSO do fornecedor diz S/N (8.2).
COLUNAS_SERIE = ("serie", "numero de serie", "n serie", "serial")
COLUNAS_SN = ("s/n", "sn")

ALIASES = {
    "serie": [*COLUNAS_SERIE, *COLUNAS_SN],
    "loja": ["loja"],
    "rma": ["rma"],
    "categoria": ["categoria"],
    "orcamento": ["orcamento", "valor orcamento", "valor", "total"],
    "razao": ["60% orcamento", "percentual", "60%"],
    "ano": ["ano"],
    "lote_prime": ["lote prime", "lote"],
    "status": ["status orcamento", "status", "aprovado via contrato"],
    "mes_contrato": ["mes contrato", "mes - contrato", "mes", "mes referencia", "mes_referencia"],
    "tipo": ["tipo de manutencao", "tipo manutencao", "tipo"],
    "status_retorno": ["status de retorno", "status retorno"],
    "ano_devolucao": ["ano devolucao"],
    "avaliacao": ["avaliacao orcamento", "avaliacao"],  # recalculada; só mapeada
    "qtde": ["qtde", "quantidade"],
    "empresa": ["empresa"],
    # 8.2 — planilha do fornecedor
    "disponibilizacao": ["disponibilizacao", "disponibilização", "data disponibilizacao"],
    "origem_equipamento": ["origem"],
    "po": ["po", "pedido"],
}
OBRIGATORIAS = ("rma", "serie", "categoria")
# Colunas que só existem se o arquivo as trouxer. A planilha do fornecedor tem
# meia dúzia de colunas (8.2): gravá-las mesmo assim apagaria loja, lote e —
# pior — o DEVOLVIDO que a tela de retorno acabou de marcar. Sem a coluna, o
# upsert não mexe no que está gravado.
CAMPOS_OPCIONAIS = ("loja", "lote_prime", "qtde", "status_retorno", "ano_devolucao",
                    "origem_equipamento", "disponibilizacao", "po")


def _mapear_colunas(colunas) -> dict:
    """{destino: nome da coluna no arquivo}; None para o que não veio."""
    cmap = {_sem_acento(c): c for c in colunas}
    mapa = {k: next((cmap[_sem_acento(a)] for a in al if _sem_acento(a) in cmap), None)
            for k, al in ALIASES.items()}
    # Lote: se o nome exato não bateu, aceita qualquer coluna que contenha
    # "lote" (ex.: "Lote de Reparo"), para não vir vazio.
    if not mapa.get("lote_prime"):
        achou = next((orig for norm, orig in cmap.items() if "lote" in norm), None)
        if achou:
            mapa["lote_prime"] = achou
    return mapa


def _layout_da_aba(colunas) -> str:
    """"CONTRATO" (RMA + SÉRIE + CATEGORIA), "AVULSO" (RMA + S/N) ou "" quando
    o cabeçalho não serve para importar."""
    cmap = {_sem_acento(c) for c in colunas}
    if not any(_sem_acento(a) in cmap for a in ALIASES["rma"]):
        return ""
    if any(a in cmap for a in COLUNAS_SERIE) and any(
            _sem_acento(a) in cmap for a in ALIASES["categoria"]):
        return "CONTRATO"
    if any(a in cmap for a in COLUNAS_SN):
        return "AVULSO"
    return ""


def _ler_planilha(conteudo: bytes, nome: str, aba: Optional[str] = None):
    """Devolve `(df, aba usada, abas do arquivo)`. Sem `aba`, escolhe a primeira
    cujo cabeçalho case com um dos layouts conhecidos (8.2); nenhuma servindo,
    devolve 400 dizendo o que cada aba tem."""
    import pandas as pd
    sufixo = Path(nome or "").suffix.lower()
    if sufixo == ".csv":
        try:
            try:
                df = pd.read_csv(io.BytesIO(conteudo), dtype=object, sep=None,
                                 engine="python", encoding="utf-8-sig")
            except UnicodeDecodeError:
                df = pd.read_csv(io.BytesIO(conteudo), dtype=object, sep=None,
                                 engine="python", encoding="latin-1")
        except Exception as exc:  # noqa: BLE001
            raise HTTPException(400, f"Arquivo inválido: {exc}")
        return df, "", []

    try:
        xls = pd.ExcelFile(io.BytesIO(conteudo))
        abas = [str(a) for a in xls.sheet_names]
    except Exception as exc:  # noqa: BLE001
        raise HTTPException(400, f"Arquivo inválido: {exc}")

    if not _vazio(aba):
        alvo = _sem_acento(aba)
        escolhida = next((a for a in abas if _sem_acento(a) == alvo), None)
        if escolhida is None:
            raise HTTPException(400, f"Aba {_texto(aba)!r} não existe neste arquivo. "
                                     f"Abas disponíveis: {', '.join(abas)}")
    else:
        # Só o cabeçalho de cada aba: ler 1.700 linhas de cada uma para
        # descobrir qual serve sairia caro.
        cabecalhos: dict[str, list] = {}
        escolhida = None
        for nome_aba in abas:
            try:
                cabecalhos[nome_aba] = list(
                    pd.read_excel(xls, sheet_name=nome_aba, dtype=object, nrows=0).columns)
            except Exception:  # noqa: BLE001
                cabecalhos[nome_aba] = []
            if _layout_da_aba(cabecalhos[nome_aba]):
                escolhida = nome_aba
                break
        if escolhida is None:
            detalhe = "; ".join(
                f"{a}: {', '.join(map(str, cols))[:200] or '(sem cabeçalho)'}"
                for a, cols in cabecalhos.items())
            raise HTTPException(400, "Nenhuma aba tem as colunas necessárias "
                                     "(RMA + SÉRIE + CATEGORIA, ou RMA + S/N). " + detalhe)

    try:
        df = pd.read_excel(xls, sheet_name=escolhida, dtype=object)
    except Exception as exc:  # noqa: BLE001
        raise HTTPException(400, f"Arquivo inválido: {exc}")
    return df, escolhida, abas


def _linha_para_colunas(d: SimpleNamespace, agora, usuario: str,
                       extras: tuple[str, ...] = ()) -> dict:
    """Colunas a gravar. O bloco fixo é o que toda planilha define; `extras`
    são os campos que só entram quando o arquivo trouxe a coluna, para o upsert
    não apagar o que ele nem menciona (ver CAMPOS_OPCIONAIS)."""
    cols = {
        "rma": d.rma, "serie": d.serie, "categoria": d.categoria,
        "familia": d.familia, "empresa": d.empresa, "orcamento": d.orcamento,
        "garantia": d.garantia, "valor_compra": d.valor_compra,
        "valor_compra_fonte": d.valor_compra_fonte, "percentual": d.percentual,
        "avaliacao": d.avaliacao, "status": d.status, "status_original": d.status_original,
        "tipo_manutencao": d.tipo_manutencao, "tipo_original": d.tipo_original,
        "ano": d.ano, "mes_referencia": d.mes_referencia,
        "atualizado_em": agora, "atualizado_por": usuario,
    }
    for campo in extras:
        cols[campo] = getattr(d, campo)
    return cols


@router.post("/importar")
def importar(req: Request, file: UploadFile = File(...),
             substituir: bool = Form(False), aba: str = Form(""),
             dry_run: bool = Form(False)):
    """`substituir=true`: a planilha vira a base. Linhas importadas antes que
    não estão no arquivo são removidas; o que foi digitado no portal fica.
    `aba`: qual aba do xlsx ler; sem ela, a primeira que servir (8.2)."""
    sd = _exigir(req, "admin")
    check_rate_limit(req, "api")
    usuario = _usuario(sd)
    nome = Path(file.filename or "planilha.xlsx").name[:200]
    df, aba_usada, abas_disponiveis = _ler_planilha(file.file.read(), nome, aba)

    colunas = _mapear_colunas(df.columns)
    layout = _layout_da_aba(df.columns)
    # A aba AVULSO não tem CATEGORIA: exigi-la rejeitaria o arquivo inteiro.
    obrigatorias = ("rma", "serie") if layout == "AVULSO" else OBRIGATORIAS
    faltando = [k for k in obrigatorias if not colunas[k]]
    if faltando:
        raise HTTPException(400, "Colunas obrigatórias ausentes: " + ", ".join(faltando)
                            + f". Colunas encontradas: {', '.join(map(str, df.columns))}")
    extras = tuple(c for c in CAMPOS_OPCIONAIS if colunas.get(c))
    tem_categoria = bool(colunas["categoria"])

    detalhes, avisos_detalhes = [], []
    avisos = Counter(status_nao_reconhecido=0, tipo_nao_reconhecido=0,
                     tipo_assumido_contrato=0,
                     empresa_vazia=0, mes_referencia_nulo=0,
                     mes_ambiguo=0, status_do_fornecedor=0)
    novos, alterados, vistos = [], [], set()
    # Todo RMA que aparece no arquivo, mesmo em linha rejeitada: com
    # `substituir`, um erro de formato não pode apagar o que já existe.
    no_arquivo: set[str] = set()
    lidas = 0

    def rejeitar(linha: int, motivo: str) -> None:
        if len(detalhes) < 200:
            detalhes.append({"linha": linha, "motivo": motivo})

    with db.SessionLocal.begin() as s:
        config = db.ler_config(s)
        limiar = config["limiar_percentual"]
        # Um SELECT só: o que já existe, com o que a ordem 3.4 precisa preservar.
        # Um SELECT só, com o que o upsert precisa preservar quando o arquivo
        # não traz a coluna.
        existentes = {
            linha.rma: linha for linha in s.execute(
                select(R.id, R.rma, R.valor_compra, R.valor_compra_fonte, R.categoria,
                       R.familia, R.empresa, R.tipo_manutencao, R.tipo_original)).all()
        }
        agora = db.utcnow()

        for idx, linha in enumerate(df.to_dict("records"), start=2):
            lidas += 1

            def val(k, _l=linha):
                c = colunas.get(k)
                return _l.get(c) if c else None

            rma = _texto_rma(val("rma"))
            if not rma:
                rejeitar(idx, "RMA vazio")
                continue
            no_arquivo.add(rma)
            if rma in vistos:
                rejeitar(idx, f"RMA repetido na planilha: {rma}")
                continue
            serie = _texto(val("serie"), 60).upper()
            if not serie:
                rejeitar(idx, "série vazia")
                continue
            ant = existentes.get(rma)
            if tem_categoria:
                categoria, familia = normalizar_categoria(val("categoria"))
                if not categoria:
                    rejeitar(idx, "categoria vazia")
                    continue
            else:
                # Sem coluna CATEGORIA (aba AVULSO): mantém a que já existe.
                categoria, familia = (ant.categoria, ant.familia) if ant else ("", "OUTRO")
            try:
                orcamento, garantia = interpretar_orcamento(val("orcamento"))
            except ValueError as exc:
                rejeitar(idx, str(exc))
                continue
            vistos.add(rma)

            # O texto do status do fornecedor também diz tipo, mês e empresa (8.2).
            status_txt = _texto(val("status"), 60)
            status, ok = normalizar_status(status_txt)
            if not ok:
                avisos["status_nao_reconhecido"] += 1
                if len(avisos_detalhes) < 200:
                    avisos_detalhes.append({"linha": idx, "motivo": f"status não reconhecido: {status_txt!r}"})
            (status, tipo_status, mes_status, empresa_status,
             garantia_status, zerar) = interpretar_status_fornecedor(status_txt)
            if zerar:
                orcamento = 0.0
            garantia = garantia or garantia_status
            if len(set(meses_por_extenso(status_txt))) > 1:
                avisos["mes_ambiguo"] += 1
                if len(avisos_detalhes) < 200:
                    avisos_detalhes.append({"linha": idx, "motivo": f"mais de um mês no status, vale o primeiro: {status_txt!r}"})

            # Tipo: a coluna manda; sem ela, o texto do status; sem os dois, a
            # aba AVULSO já diz que a manutenção é avulsa.
            tipo_txt = _texto(val("tipo"), 60)
            tipo, ok_tipo = normalizar_tipo(tipo_txt)
            usou_status = bool(mes_status)
            if not ok_tipo and tipo_status:
                tipo, ok_tipo, usou_status = tipo_status, True, True
            if not ok_tipo and layout == "AVULSO":
                tipo, ok_tipo = "AVULSA", True
            if not ok_tipo:
                if ant:  # ninguém falou do tipo: o que já estava gravado vale
                    tipo, tipo_txt = ant.tipo_manutencao, ant.tipo_original or ""
                if tipo_txt and not ant:
                    # Só é "não reconhecido" quando havia texto para reconhecer.
                    avisos["tipo_nao_reconhecido"] += 1
                    if len(avisos_detalhes) < 200:
                        avisos_detalhes.append({"linha": idx, "motivo": f"tipo não reconhecido: {tipo_txt!r}"})
                elif not ant:
                    # Planilha sem coluna de tipo e reparo novo: assume contrato.
                    avisos["tipo_assumido_contrato"] += 1
            if usou_status:
                avisos["status_do_fornecedor"] += 1

            # A empresa do status só vale quando a planilha não trouxe a coluna.
            empresa = (normalizar_empresa(val("empresa")) or (empresa_status or "")
                       or (ant.empresa if ant and not colunas["empresa"] else ""))
            if not empresa:
                avisos["empresa_vazia"] += 1
            # Mês: o do status manda; senão o da coluna de contrato; senão o
            # da disponibilização.
            disponibilizacao = data_celula(val("disponibilizacao"))
            mes = (mes_status or mes_referencia(val("mes_contrato"))
                   or mes_referencia(disponibilizacao))
            if not mes:
                avisos["mes_referencia_nulo"] += 1

            d = SimpleNamespace(
                rma=rma, serie=serie, loja=_inteiro(val("loja")), categoria=categoria,
                familia=familia, empresa=empresa, orcamento=orcamento, garantia=garantia,
                valor_compra=ant.valor_compra if ant else None,
                valor_compra_fonte=(ant.valor_compra_fonte if ant else "") or "",
                percentual=None, avaliacao="", status=status,
                status_original=status_txt or db.STATUS_ROTULOS[status],
                tipo_manutencao=tipo, tipo_original=tipo_txt,
                status_retorno=normalizar_status_retorno(val("status_retorno")),
                ano=_inteiro(val("ano")) or (int(mes[:4]) if mes else date.today().year),
                mes_referencia=mes, ano_devolucao=_inteiro(val("ano_devolucao")),
                lote_prime=_texto(val("lote_prime"), 200), qtde=_inteiro(val("qtde")) or 1,
                origem_equipamento=_texto(val("origem_equipamento"), 10).upper(),
                disponibilizacao=disponibilizacao, po=_texto(val("po"), 40),
            )
            d.valor_compra, d.valor_compra_fonte = _valor_compra_para(d, config, razao=_numero(val("razao")))
            calcular(d, limiar)

            cols = _linha_para_colunas(d, agora, usuario, extras)
            if ant:
                alterados.append({"id": ant.id, **cols})
            else:
                novos.append({**cols, "observacao": "", "ebs_erro": "", "ebs_consultado_em": None,
                              "criado_em": agora, "criado_por": usuario, "origem": "PLANILHA"})

        # Amostra para a tela conferir antes de gravar (o que entraria).
        def _amostra(reg, acao):
            return {
                "acao": acao, "rma": reg.get("rma", ""), "serie": reg.get("serie", ""),
                "categoria": reg.get("categoria", ""), "familia": reg.get("familia", ""),
                "lote_prime": reg.get("lote_prime", ""), "status": reg.get("status", ""),
                "tipo_manutencao": reg.get("tipo_manutencao", ""),
                "mes_referencia": reg.get("mes_referencia") or "",
                "empresa": reg.get("empresa", ""),
                "orcamento": reg.get("orcamento", 0), "loja": reg.get("loja"),
            }
        previa = ([_amostra(r_, "incluir") for r_ in novos[:400]]
                  + [_amostra(r_, "atualizar") for r_ in alterados[:400]])[:500]

        removidas = 0
        if not dry_run:
            for i in range(0, len(novos), 500):
                s.execute(insert(R), novos[i:i + 500])
            for i in range(0, len(alterados), 500):
                s.execute(update(R), alterados[i:i + 500])

            if substituir:
                # Só sai o que veio de planilha: o que foi digitado no portal fica.
                sobrando = [
                    rid for rid, rma in s.execute(
                        select(R.id, R.rma).where(R.origem == "PLANILHA")).all()
                    if rma not in no_arquivo
                ]
                for i in range(0, len(sobrando), 500):
                    s.execute(delete(R).where(R.id.in_(sobrando[i:i + 500])))
                removidas = len(sobrando)

            rejeitadas_log = lidas - len(novos) - len(alterados)
            s.add(db.Importacao(
                arquivo=nome, usuario=usuario, lidas=lidas, incluidas=len(novos),
                atualizadas=len(alterados), rejeitadas=rejeitadas_log,
                detalhes=json.dumps({"substituir": bool(substituir), "removidas": removidas,
                                     "aba": aba_usada, "rejeicoes": detalhes}, ensure_ascii=False),
            ))
        # dry_run: nada é gravado — o bloco só leu, então fecha sem alterações.

    rejeitadas = lidas - len(novos) - len(alterados)
    return {
        "dry_run": bool(dry_run),
        "lidas": lidas, "incluidas": len(novos), "atualizadas": len(alterados),
        "rejeitadas": rejeitadas, "removidas": removidas,
        "substituiu": bool(substituir), "detalhes": detalhes, "previa": previa,
        "aba": aba_usada, "abas_disponiveis": abas_disponiveis, "layout": layout,
        "avisos": dict(avisos), "avisos_detalhes": avisos_detalhes,
    }


# ── 5.8 Exportação ──────────────────────────────────────────────────────
def _xlsx(colunas: list, linhas: list, titulo: str, nome: str) -> StreamingResponse:
    """Padrão de planilha do módulo: cabeçalho AB4807 branco e negrito, congela
    A2, filtro automático e largura pelo conteúdo. `colunas` é
    `[(título, função que lê a linha)]`."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = titulo
    ws.append([c for c, _ in colunas])
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor="AB4807")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center")
    for linha in linhas:
        ws.append([fn(linha) for _, fn in colunas])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for coluna in ws.columns:
        largura = max(len(str(c.value or "")) for c in coluna) + 2
        ws.column_dimensions[coluna[0].column_letter].width = min(45, max(12, largura))

    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{nome}"'},
    )


@router.get("/exportar.xlsx")
def exportar(req: Request, ano: Optional[int] = None, mes: Optional[str] = None,
             familia: Optional[str] = None, categoria: Optional[str] = None,
             status: Optional[str] = None, status_retorno: Optional[str] = None,
             empresa: Optional[str] = None, tipo_manutencao: Optional[str] = None,
             q: Optional[str] = None, min_reparos: Optional[int] = None,
             lote: Optional[str] = None):
    _exigir(req, "export")
    with db.SessionLocal() as s:
        config = db.ler_config(s)
        itens = s.scalars(_ordenado(_filtrar(
            select(R), ano=ano, mes=mes, familia=familia, categoria=categoria, status=status,
            status_retorno=status_retorno, empresa=empresa, tipo_manutencao=tipo_manutencao,
            q=q, min_reparos=min_reparos, lote=lote,
        ))).all()
        # Sem paginação a exportação pode levar a base toda: um agregado só.
        agregado = _agregado_series(s)

    lim = f"{config['limiar_percentual'] * 100:g}".replace(".", ",")
    aval = {"DENTRO": f"Dentro dos {lim}%", "FORA": f"Fora dos {lim}%", "": ""}
    colunas = [
        ("SÉRIE", lambda r: r.serie), ("LOJA", lambda r: r.loja), ("RMA", lambda r: r.rma),
        ("CATEGORIA", lambda r: r.categoria),
        ("ORÇAMENTO", lambda r: "Garantia" if r.garantia else float(r.orcamento or 0)),
        ("60% ORÇAMENTO", lambda r: r.percentual), ("ANO", lambda r: r.ano),
        ("LOTE PRIME", lambda r: r.lote_prime),
        ("STATUS ORÇAMENTO", lambda r: db.STATUS_ROTULOS.get(r.status, r.status)),
        ("MÊS CONTRATO", lambda r: r.mes_referencia or ""),
        ("TIPO DE MANUTENÇÃO", lambda r: db.TIPOS_ROTULOS.get(r.tipo_manutencao, r.tipo_manutencao)),
        ("STATUS DE RETORNO", lambda r: db.RETORNO_ROTULOS.get(r.status_retorno, r.status_retorno)),
        ("ANO DEVOLUÇÃO", lambda r: r.ano_devolucao),
        ("DATA DEVOLUÇÃO", lambda r: r.devolvido_em),
        ("DEVOLUÇÃO CONFIRMADA POR", lambda r: r.devolvido_por or ""),
        ("AVALIAÇÃO ORÇAMENTO", lambda r: aval.get(r.avaliacao or "", r.avaliacao)),
        ("QTDE", lambda r: r.qtde), ("EMPRESA", lambda r: r.empresa),
        ("VALOR COMPRA", lambda r: r.valor_compra), ("FONTE VALOR", lambda r: r.valor_compra_fonte),
        ("PERCENTUAL", lambda r: round(r.percentual * 100, 2) if r.percentual is not None else None),
        ("FAMÍLIA", lambda r: r.familia),
        ("ORIGEM EQUIPAMENTO", lambda r: r.origem_equipamento or ""),
        ("DISPONIBILIZAÇÃO", lambda r: r.disponibilizacao),
        ("PO", lambda r: r.po or ""),
        ("REPAROS DA SÉRIE", lambda r: agregado.get(r.serie, (1, 0.0))[0]),
        ("CUSTO ACUMULADO DA SÉRIE", lambda r: agregado.get(r.serie, (1, 0.0))[1]),
        ("OBSERVAÇÃO", lambda r: r.observacao),
    ]

    return _xlsx(colunas, itens, "Manutenção", f"orcamento_manutencao_{ano or 'todos'}.xlsx")


# ── 5.9 Configuração ────────────────────────────────────────────────────
@router.get("/config")
def config_ler(req: Request):
    _exigir(req, "view")
    with db.SessionLocal() as s:
        return db.ler_config(s)


def _mapa_numerico(mapa: dict, rotulo: str) -> dict:
    """{chave: número > 0}; entradas vazias/zeradas são removidas."""
    saida = {}
    for k, v in (mapa or {}).items():
        chave = _texto(k, 40)
        if not chave or _vazio(v):
            continue
        n = _numero(v)
        if n is None or n < 0:
            raise HTTPException(422, f"{rotulo}: valor inválido para {chave!r}: {_texto(v)!r}")
        if n > 0:
            saida[chave] = round(n, 2)
    return saida


@router.put("/config")
def config_gravar(body: ConfigIn, req: Request):
    sd = _exigir(req, "admin")
    check_rate_limit(req, "api")
    dados: dict = {}
    if body.cota_mensal is not None:
        cotas = _mapa_numerico(body.cota_mensal, "cota_mensal")
        if any(not k.isdigit() or len(k) != 4 for k in cotas):
            raise HTTPException(422, "cota_mensal: as chaves devem ser o ano (AAAA).")
        dados["cota_mensal"] = cotas
    if body.limiar_percentual is not None:
        n = _numero(body.limiar_percentual)
        if n is None or n <= 0:
            raise HTTPException(422, "limiar_percentual deve ser uma fração maior que zero (ex.: 0.60).")
        dados["limiar_percentual"] = round(n / 100 if n > 1 else n, 4)  # aceita 60 no lugar de 0.60
    if body.valor_compra_padrao is not None:
        dados["valor_compra_padrao"] = _mapa_numerico(body.valor_compra_padrao, "valor_compra_padrao")
    with db.SessionLocal.begin() as s:
        db.gravar_config(s, dados, _usuario(sd))
        return db.ler_config(s)


@router.post("/recalcular")
def recalcular(req: Request):
    """Refaz percentual/avaliação de todas as linhas (e o valor PADRAO, se a
    configuração mudou) e aplica a regra 3.5 nas pendentes."""
    sd = _exigir(req, "admin")
    check_rate_limit(req, "api")
    mudaram = 0
    with db.SessionLocal.begin() as s:
        config = db.ler_config(s)
        linhas = s.scalars(select(R)).all()
        for r in linhas:
            antes = (r.valor_compra, r.valor_compra_fonte)
            r.valor_compra, r.valor_compra_fonte = _valor_compra_para(r, config)
            if calcular(r, config["limiar_percentual"]) or antes != (r.valor_compra, r.valor_compra_fonte):
                mudaram += 1
                r.atualizado_por = _usuario(sd)
    return {"total": len(linhas), "mudaram": mudaram}


# ── 5.10 Opções para os selects ─────────────────────────────────────────
@router.get("/opcoes")
def opcoes(req: Request):
    _exigir(req, "view")
    with db.SessionLocal() as s:
        config = db.ler_config(s)
        categorias = {c for (c,) in s.execute(select(R.categoria).distinct()).all() if c}
        categorias |= set((config.get("valor_compra_padrao") or {}).keys())
        empresas = {e for (e,) in s.execute(select(R.empresa).distinct()).all() if e}
        anos = sorted({int(a) for (a,) in s.execute(select(R.ano).distinct()).all() if a}
                      | {date.today().year}, reverse=True)
        # Todos os lotes distintos (o filtro tem autocompletar por conteúdo,
        # então precisa da lista inteira), do mais usado para o menos usado.
        lotes = [l for (l, _) in s.execute(
            select(R.lote_prime, func.count(R.id)).where(R.lote_prime != "")
            .group_by(R.lote_prime).order_by(func.count(R.id).desc())).all()]
    return {
        "categorias": sorted(categorias),
        "familias": list(db.FAMILIAS),
        "status": [{"valor": k, "rotulo": v} for k, v in db.STATUS_ROTULOS.items()],
        "tipos": [{"valor": k, "rotulo": v} for k, v in db.TIPOS_ROTULOS.items()],
        "status_retorno": [{"valor": k, "rotulo": v} for k, v in db.RETORNO_ROTULOS.items()],
        "empresas": sorted(set(db.EMPRESAS) | empresas),
        "fontes_valor": list(db.FONTES_VALOR),
        "anos": anos,
        "lotes": lotes,
        "limiar_percentual": config["limiar_percentual"],
    }


# ── 8.3 Retorno de reparo ───────────────────────────────────────────────
class RetornoIn(BaseModel):
    """Lista de RMAs bipados na tela de retorno. A caixa de texto da tela já
    separa por vírgula, ponto e vírgula, espaço ou linha; aqui só chega a lista."""
    model_config = ConfigDict(extra="ignore")
    rmas: list[Any] = []


RETORNO_MAX = 500


def _rmas_do_corpo(body: RetornoIn) -> list[str]:
    """Limpa (inclusive `\\xa0`), descarta vazios e repetidos preservando a
    ordem digitada, e recusa lote grande demais."""
    vistos, saida = set(), []
    for bruto in (body.rmas or []):
        rma = _texto_rma(bruto)
        if not rma or rma in vistos:
            continue
        vistos.add(rma)
        saida.append(rma)
    if not saida:
        raise HTTPException(422, "Informe ao menos um RMA.")
    if len(saida) > RETORNO_MAX:
        raise HTTPException(422, f"No máximo {RETORNO_MAX} RMAs por vez; vieram {len(saida)}.")
    return saida


def _item_retorno(r) -> dict:
    return {
        "rma": r.rma, "encontrado": True,
        "ja_devolvido": r.status_retorno == "DEVOLVIDO",
        "id": r.id, "serie": r.serie, "categoria": r.categoria, "familia": r.familia,
        "loja": r.loja, "empresa": r.empresa, "orcamento": float(r.orcamento or 0),
        "status": r.status, "status_rotulo": db.STATUS_ROTULOS.get(r.status, r.status),
        "status_retorno": r.status_retorno, "mes_referencia": r.mes_referencia,
        "lote_prime": r.lote_prime or "",
        "devolvido_em": r.devolvido_em.isoformat() if r.devolvido_em else None,
    }


def _itens_retorno(rmas: list[str], achados: dict) -> list[dict]:
    """Na ordem digitada; o que não está na base vem só como não encontrado."""
    return [_item_retorno(achados[rma]) if rma in achados else {"rma": rma, "encontrado": False}
            for rma in rmas]


@router.post("/retorno/consultar")
def retorno_consultar(body: RetornoIn, req: Request):
    """Mostra o que a base sabe dos RMAs bipados. Não altera nada."""
    _exigir(req, "view")
    rmas = _rmas_do_corpo(body)
    with db.SessionLocal() as s:
        # Uma consulta para o lote inteiro, não uma por RMA.
        achados = {r.rma: r for r in s.scalars(select(R).where(R.rma.in_(rmas))).all()}
        itens = _itens_retorno(rmas, achados)
    encontrados = [i for i in itens if i["encontrado"]]
    return {
        "itens": itens, "total": len(itens), "encontrados": len(encontrados),
        "nao_encontrados": len(itens) - len(encontrados),
        "ja_devolvidos": sum(1 for i in encontrados if i["ja_devolvido"]),
    }


@router.post("/retorno/confirmar")
def retorno_confirmar(body: RetornoIn, req: Request):
    """Marca os RMAs como devolvidos. Nada é criado: RMA fora da base só é
    reportado, e o status do orçamento não é tocado."""
    sd = _exigir(req, "edit")
    check_rate_limit(req, "api")
    rmas = _rmas_do_corpo(body)
    usuario = _usuario(sd)
    agora = db.utcnow()
    devolvidos = ja_devolvidos = 0
    with db.SessionLocal.begin() as s:
        achados = {r.rma: r for r in s.scalars(select(R).where(R.rma.in_(rmas))).all()}
        for r in achados.values():
            if r.status_retorno == "DEVOLVIDO":
                ja_devolvidos += 1
                continue
            r.status_retorno = "DEVOLVIDO"
            r.devolvido_em = agora
            r.devolvido_por = usuario
            r.ano_devolucao = date.today().year
            r.atualizado_por = usuario
            devolvidos += 1
        s.flush()
        itens = _itens_retorno(rmas, achados)
    return {
        "devolvidos": devolvidos, "ja_devolvidos": ja_devolvidos,
        "nao_encontrados": [rma for rma in rmas if rma not in achados],
        "itens": itens, "total": len(itens), "encontrados": len(achados),
    }


# ── 8.4 Reincidência ────────────────────────────────────────────────────
def _filtrar_serie(stmt, familia=None, categoria=None, ano=None, q=None):
    """Filtros do resumo por série. Aqui `q` procura só na série."""
    if familia:
        stmt = stmt.where(R.familia == familia.strip().upper())
    if categoria:
        stmt = stmt.where(R.categoria == normalizar_categoria(categoria)[0])
    if ano:
        stmt = stmt.where(R.ano == int(ano))
    if q and q.strip():
        stmt = stmt.where(R.serie.like(f"%{q.strip().upper()}%"))
    return stmt


def _resumo_series(s, min_reparos: int, familia=None, categoria=None, ano=None,
                   q=None, limit: Optional[int] = None, offset: int = 0):
    """`(total, resumo, itens)` do resumo por série. Os filtros valem para o
    que é contado: pedir `familia=SLED` conta os reparos de SLED da série.
    Duas consultas — a agregada e a dos detalhes da página (categoria mais
    recente e lojas) —, nunca uma por série."""
    reparos = func.count(R.id).label("reparos")
    custo = _CUSTO_APROVADO.label("custo_total")
    reprovados = func.coalesce(
        func.sum(case((R.status == "REPROVADO", 1), else_=0)), 0).label("reprovados")
    grupo = (_filtrar_serie(
        select(R.serie.label("serie"), reparos, custo, reprovados,
               func.min(R.mes_referencia).label("primeiro"),
               func.max(R.mes_referencia).label("ultimo")),
        familia, categoria, ano, q)
        .where(R.serie != "").group_by(R.serie)
        .having(func.count(R.id) >= min_reparos))

    sub = grupo.subquery()
    total, soma_reparos, soma_custo = s.execute(
        select(func.count(), func.coalesce(func.sum(sub.c.reparos), 0),
               func.coalesce(func.sum(sub.c.custo_total), 0)).select_from(sub)).one()

    ordenado = grupo.order_by(reparos.desc(), custo.desc(), R.serie)
    if limit is not None:
        ordenado = ordenado.limit(limit).offset(offset)
    linhas = s.execute(ordenado).all()

    series = [l.serie for l in linhas]
    detalhes: dict[str, dict] = {}
    for i in range(0, len(series), 400):
        for (serie, cat, fam, loja, rid, rma, mes, status,
             orcamento, retorno) in s.execute(
            _filtrar_serie(select(R.serie, R.categoria, R.familia, R.loja, R.id,
                                  R.rma, R.mes_referencia, R.status, R.orcamento,
                                  R.status_retorno),
                           familia, categoria, ano, q)
            .where(R.serie.in_(series[i:i + 400]))
            .order_by(R.serie, R.mes_referencia.asc().nullsfirst(), R.id)
        ).all():
            d = detalhes.setdefault(serie, {"categoria": "", "familia": "",
                                            "lojas": [], "atendimentos": []})
            if cat:  # em ordem crescente de mês: sobra a categoria mais recente
                d["categoria"], d["familia"] = cat, fam
            if loja is not None and loja not in d["lojas"]:
                d["lojas"].append(loja)
            # Os RMAs que compõem a reincidência: é por eles que se chega ao
            # atendimento, na tela e na planilha de análise.
            d["atendimentos"].append({
                "id": rid, "rma": rma, "mes": mes, "status": status,
                "status_rotulo": db.STATUS_ROTULOS.get(status, status),
                "orcamento": _dinheiro(orcamento), "status_retorno": retorno,
                "loja": loja,
            })

    itens = []
    for l in linhas:
        d = detalhes.get(l.serie) or {}
        qtde = int(l.reparos)
        total_custo = _dinheiro(l.custo_total)
        itens.append({
            "serie": l.serie, "categoria": d.get("categoria", ""),
            "familia": d.get("familia", ""), "reparos": qtde,
            "custo_total": total_custo,
            "custo_medio": _dinheiro(total_custo / qtde) if qtde else 0.0,
            "primeiro": l.primeiro, "ultimo": l.ultimo,
            "reprovados": int(l.reprovados), "lojas": (d.get("lojas") or [])[:10],
            "atendimentos": (d.get("atendimentos") or [])[:50],
            "rmas": [a["rma"] for a in (d.get("atendimentos") or [])],
            "ultimo_rma": (d.get("atendimentos") or [{}])[-1].get("rma", ""),
        })
    resumo = {"series": int(total), "reparos": int(soma_reparos),
              "custo_total": _dinheiro(soma_custo)}
    return int(total), resumo, itens


@router.get("/reincidencia")
def reincidencia(req: Request, min_reparos: int = 2, familia: Optional[str] = None,
                 categoria: Optional[str] = None, ano: Optional[int] = None,
                 q: Optional[str] = None, limit: int = 100, offset: int = 0):
    """Séries que voltaram para reparo (8.4), das que mais voltaram para as que
    menos voltaram e, no empate, das mais caras para as mais baratas."""
    _exigir(req, "view")
    min_reparos = max(1, int(min_reparos or 2))
    limit = max(1, min(int(limit or 100), 1000))
    offset = max(0, int(offset or 0))
    with db.SessionLocal() as s:
        total, resumo, itens = _resumo_series(
            s, min_reparos, familia, categoria, ano, q, limit, offset)
    return {"total": total, "itens": itens, "resumo": resumo}


@router.get("/reincidencia.xlsx")
def reincidencia_xlsx(req: Request, min_reparos: int = 2, familia: Optional[str] = None,
                      categoria: Optional[str] = None, ano: Optional[int] = None,
                      q: Optional[str] = None):
    _exigir(req, "export")
    min_reparos = max(1, int(min_reparos or 2))
    with db.SessionLocal() as s:
        _, _, itens = _resumo_series(s, min_reparos, familia, categoria, ano, q)
    colunas = [
        ("SÉRIE", lambda i: i["serie"]), ("CATEGORIA", lambda i: i["categoria"]),
        ("FAMÍLIA", lambda i: i["familia"]), ("REPAROS DA SÉRIE", lambda i: i["reparos"]),
        ("CUSTO ACUMULADO DA SÉRIE", lambda i: i["custo_total"]),
        ("CUSTO MÉDIO", lambda i: i["custo_medio"]),
        ("PRIMEIRO MÊS", lambda i: i["primeiro"] or ""),
        ("ÚLTIMO MÊS", lambda i: i["ultimo"] or ""),
        ("REPROVADOS", lambda i: i["reprovados"]),
        ("LOJAS", lambda i: ", ".join(str(x) for x in i["lojas"])),
        ("ÚLTIMO RMA", lambda i: i.get("ultimo_rma", "")),
        ("RMAS DA SÉRIE", lambda i: ", ".join(i.get("rmas") or [])),
    ]
    return _xlsx(colunas, itens, "Reincidência", f"reincidencia_{ano or 'todos'}.xlsx")
