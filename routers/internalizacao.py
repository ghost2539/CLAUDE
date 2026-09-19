"""Internalização (menu Entrada) — API /api/internalizacao.

Recebe os agendamentos cujo recebimento foi confirmado (status RECEBIDO) e,
para cada equipamento fisicamente recebido, registra: item do EBS, descrição,
plaqueta (nº do bem), número de série e a NF. Grava em banco próprio
(`db/internalizacao.py`) e exporta a planilha "Placa Patrimonial" no mesmo
layout do modelo oficial.

Lê o banco dos Agendamentos apenas para LISTAR os recebidos e tirar o
snapshot do processo; nunca escreve nele.

Mantém também dois cadastros que o lançamento consome (`/cadastro/...`):
o estoque de etiquetas de patrimônio, com o local de cada lote, e a lista
dos itens do EBS que são imobilizados.

Permissão pelo módulo "internalizacao": view lê, create/edit lançam e salvam,
export exporta, admin exclui.
"""
from __future__ import annotations

import io
import logging
from datetime import date

from fastapi import APIRouter, HTTPException, Request
from fastapi.responses import FileResponse, StreamingResponse
from pydantic import BaseModel, ConfigDict, field_validator

import db.internalizacao as db
import db.agendamentos_forn as agf_db
from core.security import check_rate_limit, client_ip, get_session, require_permission
from config import get_settings as _get_settings

_log = logging.getLogger("internalizacao")
MODULO = "internalizacao"
_cfg = _get_settings()
AREA_TMP = "lancamentos"
AREA_TMP_NOTAS = "recebimento_forn"

router = APIRouter(prefix="/api/internalizacao", tags=["Internalização"])

_XLSX_MIME = ("application/vnd.openxmlformats-officedocument."
              "spreadsheetml.sheet")


class AtivoIn(BaseModel):
    model_config = ConfigDict(extra="ignore")
    id: int | None = None
    ebs_item: str = ""
    descricao: str = ""
    plaqueta: str = ""
    numero_serie: str = ""

    @field_validator("ebs_item", "descricao", "plaqueta", "numero_serie")
    @classmethod
    def _limpa(cls, v: str) -> str:
        return (v or "").strip()[:200]


class SalvarIn(BaseModel):
    model_config = ConfigDict(extra="ignore")
    ativos: list[AtivoIn] = []
    concluir: bool = False


def _exigir(req: Request, acao: str) -> dict:
    return require_permission(req, MODULO, acao)


def _garantir_processo(s, agendamento_id: int):
    """Retorna o Processo do agendamento, criando-o (com snapshot) se preciso.
    O snapshot vem do banco dos Agendamentos e só é lido para copiar os dados."""
    from sqlalchemy import select
    proc = s.scalar(select(db.Processo).where(
        db.Processo.agendamento_id == agendamento_id))
    if proc:
        return proc
    agf_db.ensure_db()
    with agf_db.SessionLocal() as sa:
        ag = sa.get(agf_db.Agendamento, agendamento_id)
        if not ag:
            raise HTTPException(404, "Agendamento não encontrado.")
        if ag.status != "RECEBIDO":
            raise HTTPException(409, "O recebimento deste agendamento ainda não "
                                     "foi confirmado.")
        dados = ag.to_dict()
    proc = db.Processo(
        agendamento_id=agendamento_id,
        bu=dados["bu"], fornecedor=dados["fornecedor"], nf=dados["nf"],
        estoque_destino=dados["estoque_destino"],
        estoque_destino_rotulo=dados["estoque_destino_rotulo"],
        data_recebimento=(date.fromisoformat(dados["data_recebimento"])
                          if dados["data_recebimento"] else None),
        status="PENDENTE",
    )
    s.add(proc)
    s.flush()
    return proc


def abrir_processo_do_recebimento(agendamento: dict, recebimento: dict,
                                  usuario: str) -> dict:
    """Abre o processo de lançamento a partir da conferência da chegada.

    Uma linha por unidade recebida de item imobilizado, já com o que a
    planilha do CSC Lançamentos vai pedir: item do EBS, descrição, serial,
    PO/linha, NF — e a ETIQUETA, consumida do estoque na mesma transação
    que cria o ativo (reservar numa e gravar noutra é como se perde
    etiqueta quando a segunda falha).

    Idempotente de propósito: a conferência grava no banco dos
    Agendamentos e este processo no banco da Internalização, dois arquivos
    sem transação em comum. Se o segundo passo falhar, o Lançamento chama
    de novo ao abrir a tela — e encontra o processo pela metade ou
    inexistente e completa, em vez de criar um segundo.
    """
    from sqlalchemy import select
    db.ensure_db()
    agendamento_id = int(agendamento["id"])
    unidades = [
        (it, serial)
        for it in recebimento.get("itens", []) if it.get("imobilizado", True)
        for serial in it.get("seriais", [])
    ]
    with db.SessionLocal.begin() as s:
        proc = s.scalar(select(db.Processo).where(
            db.Processo.agendamento_id == agendamento_id))
        if proc and proc.ativos:
            return proc.to_dict()
        if proc is None:
            proc = db.Processo(
                agendamento_id=agendamento_id,
                bu=agendamento.get("bu") or "",
                fornecedor=agendamento.get("fornecedor") or "",
                nf=agendamento.get("nf") or "",
                estoque_destino=agendamento.get("estoque_destino") or "",
                estoque_destino_rotulo=agendamento.get("estoque_destino_rotulo") or "",
                data_recebimento=(date.fromisoformat(agendamento["data_recebimento"])
                                  if agendamento.get("data_recebimento") else date.today()),
                status="PENDENTE",
            )
            s.add(proc)
            s.flush()
        etiquetas = etiquetas_disponiveis(s, len(unidades))
        if len(etiquetas) < len(unidades):
            raise HTTPException(
                409, f"Faltam {len(unidades) - len(etiquetas)} etiqueta(s) no estoque "
                     f"para {len(unidades)} unidade(s). Cadastre em Internalização → "
                     "Cadastro de Etiquetas e confirme de novo.")
        agora = db.utcnow()
        for (it, serial), etq in zip(unidades, etiquetas):
            ativo = db.Ativo(
                ebs_item=(it.get("item_ebs") or "")[:60],
                descricao=(it.get("descricao") or "")[:200],
                plaqueta=etq.codigo, numero_serie=str(serial)[:80],
                po=(it.get("po") or "")[:40], linha=it.get("linha"),
                nf=(it.get("nf") or "")[:40],
                etiqueta_id=etq.id, etiqueta_local=etq.local or "",
                item_id_recebimento=it.get("id"),
                criado_por=usuario,
            )
            proc.ativos.append(ativo)
            s.flush()
            etq.situacao = db.ETIQUETA_CONSUMIDA
            etq.consumida_em = agora
            etq.ativo_id = ativo.id
        proc.atualizado_em = agora
        s.flush()
        dado = proc.to_dict()
    _log.info("processo aberto pelo recebimento agendamento=%s unidades=%s etiquetas=%s por=%s",
              agendamento_id, len(unidades), len(etiquetas), usuario)
    return dado


def _devolver_etiqueta(s, ativo) -> None:
    """A etiqueta de um ativo removido do lançamento volta ao estoque: ela
    não foi colada em nada."""
    if not ativo.etiqueta_id:
        return
    etq = s.get(db.Etiqueta, ativo.etiqueta_id)
    if etq and etq.situacao == db.ETIQUETA_CONSUMIDA and etq.ativo_id == ativo.id:
        etq.situacao = db.ETIQUETA_DISPONIVEL
        etq.consumida_em = None
        etq.ativo_id = None


def _limpar_expirados() -> None:
    try:
        from core import arquivos_temporarios
        arquivos_temporarios.limpar_expirados(area=AREA_TMP)
        arquivos_temporarios.limpar_expirados(area=AREA_TMP_NOTAS)
    except Exception as exc:  # noqa: BLE001
        _log.debug("limpeza dos temporários não rodou: %s", exc)


@router.get("")
@router.get("/")
def listar(req: Request, status: str = "", busca: str = ""):
    """Lista os agendamentos recebidos e o andamento da internalização de cada."""
    _exigir(req, "view")
    _limpar_expirados()
    from sqlalchemy import select
    agf_db.ensure_db()
    db.ensure_db()
    with db.SessionLocal() as s:
        procs = {p.agendamento_id: p.to_dict(com_ativos=False)
                 for p in s.scalars(select(db.Processo)).all()}
    with agf_db.SessionLocal() as sa:
        recs = sa.scalars(
            select(agf_db.Agendamento)
            .where(agf_db.Agendamento.status == "RECEBIDO")
            .order_by(agf_db.Agendamento.id.desc())
        ).all()
        linhas = []
        for ag in recs:
            d = ag.to_dict()
            p = procs.get(ag.id)
            linhas.append({
                "agendamento_id": ag.id,
                "nf": d["nf"], "po": d["po"], "bu": d["bu"],
                "fornecedor": d["fornecedor"],
                "estoque_destino_rotulo": d["estoque_destino_rotulo"],
                "data_recebimento": d["data_recebimento"],
                "equipamentos": d["equipamentos"],
                "status": (p or {}).get("status", "PENDENTE"),
                "status_rotulo": (p or {}).get("status_rotulo", "Pendente"),
                "total_ativos": (p or {}).get("total_ativos", 0),
                "entrega_parcial": bool(d.get("entrega_parcial")),
                "sn_request_number": (p or {}).get("sn_request_number", ""),
                "sn_pendente": bool((p or {}).get("sn_pendente")),
            })
    st = (status or "").strip().upper()
    if st in db.STATUS:
        linhas = [l for l in linhas if l["status"] == st]
    termo = (busca or "").strip().lower()
    if termo:
        linhas = [l for l in linhas
                  if any(termo in str(l.get(c, "")).lower()
                         for c in ("nf", "po", "fornecedor", "bu"))]
    return {"total": len(linhas), "itens": linhas}


@router.get("/{agendamento_id}")
def obter(agendamento_id: int, req: Request):
    """Abre a internalização de um agendamento (cria o processo se preciso) e
    devolve os ativos já lançados + os equipamentos esperados (para semear)."""
    sd = _exigir(req, "view")
    db.ensure_db()
    agf_db.ensure_db()
    with agf_db.SessionLocal() as sa:
        ag = sa.get(agf_db.Agendamento, agendamento_id)
        ag_dado = ag.to_dict() if ag else None
        rec_dado = ag.recebimento.to_dict() if ag and ag.recebimento else None
        esperados = [e.to_dict() for e in ag.equipamentos] if ag else []
    notas = _dados_do_agendamento(agendamento_id)[1] if ag_dado else []
    for n in notas:
        n.pop("pdf_arquivo", None)
        n.pop("xml_arquivo", None)
    if ag_dado and rec_dado:
        abrir_processo_do_recebimento(ag_dado, rec_dado, sd.get("username", ""))
    with db.SessionLocal.begin() as s:
        proc = _garantir_processo(s, agendamento_id)
        dado = proc.to_dict()
    dado["equipamentos_esperados"] = esperados
    dado["recebimento"] = rec_dado
    dado["notas"] = notas
    return dado


@router.put("/{agendamento_id}")
def salvar(agendamento_id: int, body: SalvarIn, req: Request):
    """Substitui a lista de ativos internalizados do agendamento."""
    sd = _exigir(req, "edit")
    db.ensure_db()
    usuario = sd.get("username", "")
    with db.SessionLocal.begin() as s:
        proc = _garantir_processo(s, agendamento_id)
        if proc.lancado_em:
            raise HTTPException(409, "Este lançamento já foi enviado; os ativos não mudam mais aqui.")
        atuais = {a.id: a for a in proc.ativos}
        vistos = set()
        for a in body.ativos:
            if not (a.ebs_item or a.descricao or a.plaqueta or a.numero_serie):
                continue
            existente = atuais.get(a.id) if a.id else None
            if existente is not None:
                existente.ebs_item, existente.descricao = a.ebs_item, a.descricao
                existente.numero_serie = a.numero_serie
                if not existente.etiqueta_id:
                    existente.plaqueta = a.plaqueta
                vistos.add(existente.id)
                continue
            proc.ativos.append(db.Ativo(
                ebs_item=a.ebs_item, descricao=a.descricao,
                plaqueta=a.plaqueta, numero_serie=a.numero_serie,
                criado_por=usuario,
            ))
        for ativo_id, existente in atuais.items():
            if ativo_id not in vistos:
                _devolver_etiqueta(s, existente)
                proc.ativos.remove(existente)
        proc.status = "CONCLUIDA" if body.concluir else "PENDENTE"
        proc.atualizado_em = db.utcnow()
        s.flush()
        dado = proc.to_dict()
    _log.info("internalização salva agendamento=%s ativos=%s status=%s por=%s ip=%s",
              agendamento_id, len(dado["ativos"]), dado["status"], usuario,
              client_ip(req))
    return dado


@router.get("/{agendamento_id}/exportar")
def exportar(agendamento_id: int, req: Request):
    """Gera a planilha "Placa Patrimonial" (mesmo layout do modelo oficial)."""
    _exigir(req, "export")
    db.ensure_db()
    from sqlalchemy import select
    with db.SessionLocal() as s:
        proc = s.scalar(select(db.Processo).where(
            db.Processo.agendamento_id == agendamento_id))
        if not proc:
            raise HTTPException(404, "Nada internalizado para este agendamento.")
        nf = proc.nf or ""
        linhas = [a.to_dict() for a in proc.ativos]
    conteudo = _montar_xlsx(nf, linhas)
    nome = f"Placa_Patrimonial_NF_{nf or agendamento_id}.xlsx"
    return StreamingResponse(
        io.BytesIO(conteudo), media_type=_XLSX_MIME,
        headers={"Content-Disposition": f'attachment; filename="{nome}"'},
    )


@router.delete("/{agendamento_id}", status_code=204)
def excluir(agendamento_id: int, req: Request):
    """Remove o processo de internalização.

    Some o que foi registrado AQUI — patrimônio, plaquetas, confirmações e os
    ativos do processo. O agendamento em si e o que já deu entrada no estoque
    do portal NÃO são tocados: são de outros módulos.

    Não dá para desfazer, então fica no log de acesso com quantos ativos o
    processo tinha. Na tela o botão é só do admin do portal inteiro."""
    sd = _exigir(req, "admin")
    db.ensure_db()
    from sqlalchemy import func, select
    with db.SessionLocal.begin() as s:
        proc = s.scalar(select(db.Processo).where(
            db.Processo.agendamento_id == agendamento_id))
        if not proc:
            raise HTTPException(404, "Processo não encontrado.")
        nf = proc.nf or ""
        ativos = s.scalar(select(func.count(db.Ativo.id)).where(
            db.Ativo.processo_id == proc.id)) or 0
        for a in proc.ativos:
            _devolver_etiqueta(s, a)
        s.delete(proc)
    _registrar_exclusao(sd, req, agendamento_id, nf, ativos)
    return None


def _registrar_exclusao(sd: dict, req: Request, agendamento_id: int,
                        nf: str, ativos: int) -> None:
    """Grava no log de acesso do PORTAL (não no banco deste módulo, que acabou
    de perder a linha). Falha aqui nunca desfaz a exclusão já efetivada."""
    try:
        from db.portal import AccessLog, SessionLocal as PortalSession
        with PortalSession.begin() as s:
            s.add(AccessLog(
                login=sd.get("username", ""),
                auth_source=sd.get("auth_source", "LOCAL"),
                success=True,
                ip=client_ip(req),
                detail=(f"Internalização do agendamento {agendamento_id} "
                        f"(NF {nf}) excluída — {ativos} ativo(s)")[:500],
            ))
    except Exception as exc:  # noqa: BLE001
        _log.warning("Exclusão da internalização %s não foi registrada: %s",
                     agendamento_id, exc)


def _montar_xlsx(nf: str, linhas: list[dict]) -> bytes:
    """Monta a planilha no MESMO desenho do modelo: aba 'Placa Patrimonial',
    faixa vermelha (C00000) com texto branco, cabeçalho NF | Descrição do item
    | Plaqueta | Número de série | Item, dados a partir da linha 7."""
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Placa Patrimonial"
    ws.sheet_view.showGridLines = False

    vermelho = PatternFill("solid", fgColor="C00000")
    branco_bold = Font(name="Calibri", size=12, bold=True, color="FFFFFF")
    corpo = Font(name="Calibri", size=11)
    centro = Alignment(horizontal="center", vertical="center")
    fina = Side(style="thin")
    borda = Border(left=fina, right=fina, top=fina, bottom=fina)

    for col, larg in {"A": 7.57, "B": 123.29, "C": 14.14,
                      "D": 22.14, "E": 8.14, "F": 13.0}.items():
        ws.column_dimensions[col].width = larg

    ws.merge_cells("A5:D5")
    c = ws["A5"]
    c.value = "Placa Patrimonial (N° do bem)"
    c.fill = vermelho
    c.font = branco_bold
    c.alignment = centro
    for col in "ABCD":
        cc = ws[f"{col}5"]
        cc.fill = vermelho
        cc.border = borda
    ws.row_dimensions[5].height = 15.75

    cabec = ["NF", "Descrição do item", "Plaqueta", "Número de série", "Item"]
    for i, titulo in enumerate(cabec):
        cell = ws.cell(row=6, column=i + 1, value=titulo)
        cell.fill = vermelho
        cell.font = branco_bold
        cell.alignment = centro
        cell.border = borda
    ws.row_dimensions[6].height = 15.75

    r = 7
    for it in linhas:
        valores = [
            int(nf) if str(nf).isdigit() else nf,
            it.get("descricao", ""),
            it.get("plaqueta", ""),
            it.get("numero_serie", ""),
            int(it["ebs_item"]) if str(it.get("ebs_item", "")).isdigit()
            else it.get("ebs_item", ""),
        ]
        for i, v in enumerate(valores):
            cell = ws.cell(row=r, column=i + 1, value=v)
            cell.font = corpo
            cell.alignment = centro
            cell.border = borda
        ws.row_dimensions[r].height = 12.75
        r += 1

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()




def _linha_ativo(a, proc) -> dict:
    """O equipamento com o contexto da remessa que a tela precisa mostrar."""
    d = a.to_dict()
    d.update({
        "processo_id": proc.id,
        "agendamento_id": proc.agendamento_id,
        "bu": proc.bu or "",
        "fornecedor": proc.fornecedor or "",
        "nf": proc.nf or "",
        "estoque_destino": proc.estoque_destino or "",
        "estoque_destino_rotulo": proc.estoque_destino_rotulo or "",
        "data_recebimento": (proc.data_recebimento.isoformat()
                             if proc.data_recebimento else ""),
        "tem_ebs": (proc.bu or "") in db.BUS_COM_EBS,
    })
    return d


def _ativos_na_etapa(s, etapa: str) -> list[dict]:
    from sqlalchemy import select
    saida = []
    for proc in s.scalars(select(db.Processo)).all():
        if proc.status != "CONCLUIDA":
            continue
        for a in proc.ativos:
            if (a.etapa or db.ETAPA_PATRIMONIO) == etapa:
                saida.append(_linha_ativo(a, proc))
    saida.sort(key=lambda x: (x.get("data_recebimento") or "", x["id"]))
    return saida


@router.get("/fluxo/patrimonio")
def patrimonio_listar(req: Request):
    """Equipamentos esperando o patrimônio aparecer no EBS."""
    _exigir(req, "view")
    db.ensure_db()
    with db.SessionLocal() as s:
        itens = _ativos_na_etapa(s, db.ETAPA_PATRIMONIO)
    return {"itens": itens, "total": len(itens),
            "aguardando_ebs": sum(1 for x in itens if x["tem_ebs"]),
            "aguardando_confirmacao": sum(1 for x in itens if not x["tem_ebs"])}


@router.post("/fluxo/patrimonio/consultar")
def patrimonio_consultar(req: Request):
    """Procura no EBS o serial de cada equipamento que ainda espera.

    Roda no botão, nunca sozinha: automação que dispara por conta própria
    não tem quem responda por ela quando falha, e o portal não tem agendador.

    Quem não é Renner nem Camicado fica de fora — não há patrimônio no EBS
    para encontrar, e insistir só geraria consulta inútil e ruído no log.
    """
    sd = _exigir(req, "edit")
    check_rate_limit(req, "api")
    db.ensure_db()

    with db.SessionLocal() as s:
        pendentes = [x for x in _ativos_na_etapa(s, db.ETAPA_PATRIMONIO)
                     if x["tem_ebs"] and x["numero_serie"]]
    if not pendentes:
        return {"ok": True, "consultados": 0, "encontrados": 0, "ids": [],
                "aviso": "Nenhum equipamento de Renner ou Camicado esperando "
                         "com número de série preenchido."}

    try:
        from integracoes import ebs_oracle
    except ImportError as exc:
        raise HTTPException(
            503, f"O driver Oracle não está instalado neste servidor: {exc}") from exc

    achados, erros = {}, []
    for item in pendentes:
        try:
            linhas = ebs_oracle.run_named(
                "ativo_por_serial", {"numero_serie": item["numero_serie"]}, max_rows=5)
        except Exception as exc:  # noqa: BLE001
            if type(exc).__name__ == "EbsOracleSemCredencial":
                raise HTTPException(503, str(exc)) from exc
            _log.warning("consulta do serial %s falhou: %s", item["numero_serie"], exc)
            erros.append(item["numero_serie"])
            continue
        if linhas:
            achados[item["id"]] = linhas[0]

    if not achados:
        return {"ok": True, "consultados": len(pendentes), "encontrados": 0,
                "erros": len(erros), "ids": [],
                "aviso": "Nenhum dos seriais consultados está no EBS ainda."}

    from sqlalchemy import select
    agora = db.utcnow()
    with db.SessionLocal.begin() as s:
        for ativo_id, linha in achados.items():
            a = s.get(db.Ativo, ativo_id)
            if not a or a.etapa != db.ETAPA_PATRIMONIO:
                continue
            a.ebs_ativo = str(linha.get("ativo") or "")[:60]
            a.ebs_descricao = str(linha.get("descricao") or "")[:200]
            a.ebs_encontrado_em = agora
            if not a.plaqueta and linha.get("plaqueta"):
                a.plaqueta = str(linha["plaqueta"])[:60]
            a.etapa = db.ETAPA_ENTRADA
    _log.info("patrimônio: %s de %s seriais achados no EBS por=%s ip=%s",
              len(achados), len(pendentes), sd.get("username", ""), client_ip(req))
    return {"ok": True, "consultados": len(pendentes), "encontrados": len(achados),
            "erros": len(erros), "ids": sorted(achados)}


class ConfirmarIn(BaseModel):
    model_config = ConfigDict(extra="ignore")
    ids: list[int] = []


@router.post("/fluxo/patrimonio/confirmar")
def patrimonio_confirmar(body: ConfirmarIn, req: Request):
    """Confirma à mão a entrada do ativo (BU sem patrimônio no EBS)."""
    sd = _exigir(req, "edit")
    db.ensure_db()
    if not body.ids:
        raise HTTPException(422, "Escolha ao menos um equipamento.")
    usuario = sd.get("username", "")
    agora = db.utcnow()
    confirmados, recusados = [], []
    with db.SessionLocal.begin() as s:
        for ativo_id in body.ids[:500]:
            a = s.get(db.Ativo, ativo_id)
            if not a or a.etapa != db.ETAPA_PATRIMONIO:
                recusados.append(ativo_id)
                continue
            if (a.processo.bu or "") in db.BUS_COM_EBS:
                recusados.append(ativo_id)
                continue
            a.confirmado_por = usuario
            a.confirmado_em = agora
            a.etapa = db.ETAPA_ENTRADA
            confirmados.append(ativo_id)
    if not confirmados:
        raise HTTPException(
            422, "Nada a confirmar: os equipamentos escolhidos já saíram desta "
                 "etapa, ou são de BU que tem patrimônio no EBS e precisa da "
                 "consulta em vez da confirmação manual.")
    _log.info("patrimônio confirmado à mão ids=%s por=%s ip=%s",
              confirmados, usuario, client_ip(req))
    return {"ok": True, "confirmados": confirmados, "recusados": recusados}


@router.get("/fluxo/entrada")
def entrada_listar(req: Request):
    """Equipamentos com patrimônio resolvido, esperando o técnico."""
    _exigir(req, "view")
    db.ensure_db()
    with db.SessionLocal() as s:
        itens = _ativos_na_etapa(s, db.ETAPA_ENTRADA)
    return {"itens": itens, "total": len(itens)}


class EntradaIn(BaseModel):
    model_config = ConfigDict(extra="ignore")
    ids: list[int] = []
    espaco_corredor: str = ""

    @field_validator("espaco_corredor")
    @classmethod
    def _espaco(cls, v: str) -> str:
        v = " ".join((v or "").split())[:60]
        if not v:
            raise ValueError("Informe o espaço e corredor onde o equipamento fica.")
        return v


@router.post("/fluxo/entrada")
def entrada_concluir(body: EntradaIn, req: Request):
    """Sobe o equipamento no ServiceNow e o coloca no estoque do portal.

    É a última etapa do fluxo, e a que faz "entrou em estoque" significar
    alguma coisa: além de marcar no ServiceNow, cria o ativo e o ciclo de
    recebimento no banco do portal. Sem isso o equipamento ficaria só
    registrado como concluído aqui, sem aparecer na Consulta nem poder ser
    separado — e a tela mentiria sobre onde ele está.

    O espaço e corredor é obrigatório pelo mesmo motivo que já é no
    Recebimento: sem ele o ServiceNow recebe o ativo sem lugar, e depois
    ninguém acha o equipamento na prateleira.

    A escrita no ServiceNow sai no nome de quem está operando — nunca numa
    conta de serviço.
    """
    sd = _exigir(req, "edit")
    check_rate_limit(req, "api")
    db.ensure_db()
    if not body.ids:
        raise HTTPException(422, "Escolha ao menos um equipamento.")

    usuario = sd.get("username", "")
    with db.SessionLocal() as s:
        alvos, recusados = [], []
        for ativo_id in body.ids[:200]:
            a = s.get(db.Ativo, ativo_id)
            if not a or a.etapa != db.ETAPA_ENTRADA:
                recusados.append(ativo_id)
                continue
            alvos.append(_linha_ativo(a, a.processo))
    if not alvos:
        raise HTTPException(
            422, "Nada a concluir: os equipamentos escolhidos não estão na "
                 "etapa de entrada.")

    from datetime import date as _date
    from sqlalchemy import func, select as _select
    from db.portal import (SessionLocal as PortalSession, Asset, ReceiptCycle,
                           Movement)
    from routers.helpers import upsert_asset

    criados = []
    with PortalSession.begin() as ps:
        for item in alvos:
            ativo = upsert_asset(ps, {
                "empresa": item.get("bu") or "",
                "ativo": item.get("ebs_ativo") or item.get("plaqueta") or "",
                "etiqueta": item.get("plaqueta") or "",
                "numero_serie": item.get("numero_serie") or "",
                "descricao": item.get("ebs_descricao") or item.get("descricao") or "",
                "fonte": "INTERNALIZACAO",
            })
            ps.flush()
            n = (ps.scalar(_select(func.max(ReceiptCycle.cycle_number))
                           .where(ReceiptCycle.asset_id == ativo.id)) or 0) + 1
            hoje = _date.today()
            iso = hoje.isocalendar()
            ciclo = ReceiptCycle(
                asset_id=ativo.id, cycle_number=n, received_date=hoje,
                iso_week=f"{iso.year}-S{iso.week:02d}", status="RECEBIDO",
                origem_entrada="FORNECEDOR", po="", nf=item.get("nf") or "",
                lot_number="", open=False,
                note=f"Internalização do agendamento {item.get('agendamento_id')}",
                created_by=usuario, updated_by=usuario)
            ps.add(ciclo)
            ps.flush()
            ps.add(Movement(asset_id=ativo.id, cycle_id=ciclo.id,
                            new_status="RECEBIDO", origin="INTERNALIZACAO",
                            note=f"Espaço e corredor: {body.espaco_corredor}",
                            username=usuario))
            criados.append({"ativo_id": item["id"], "asset_id": ativo.id})

    itens_sn = [{
        "etiqueta": x.get("plaqueta") or "",
        "numero_serie": x.get("numero_serie") or "",
        "descricao": x.get("ebs_descricao") or x.get("descricao") or "",
        "ativo": x.get("ebs_ativo") or "",
    } for x in alvos]
    resumo_sn = {"ativo": False, "falhas": []}
    try:
        from routers.recebimento import _marcar_no_servicenow
        resumo_sn = _marcar_no_servicenow(itens_sn, req, body.espaco_corredor)
    except Exception as exc:  # noqa: BLE001
        resumo_sn = {"ativo": True, "falhas": [str(exc)]}
        _log.error("ServiceNow: %s equipamento(s) entraram sem marcação: %s",
                   len(itens_sn), exc)

    agora = db.utcnow()
    por_ativo = {c["ativo_id"]: c["asset_id"] for c in criados}
    with db.SessionLocal.begin() as s:
        for ativo_id, asset_id in por_ativo.items():
            a = s.get(db.Ativo, ativo_id)
            if not a or a.etapa != db.ETAPA_ENTRADA:
                continue
            a.espaco_corredor = body.espaco_corredor
            a.entrada_por = usuario
            a.entrada_em = agora
            a.asset_id = asset_id
            a.etapa = db.ETAPA_CONCLUIDO

    _log.info("entrada concluída ids=%s espaco=%r por=%s ip=%s",
              sorted(por_ativo), body.espaco_corredor, usuario, client_ip(req))
    return {"ok": True, "concluidos": sorted(por_ativo), "recusados": recusados,
            "servicenow": resumo_sn,
            "aviso": ("O ServiceNow não confirmou a marcação; o equipamento já "
                      "está no estoque do portal. Repita a marcação depois.")
                     if resumo_sn.get("falhas") else ""}



import re as _re

_ETIQUETA_RE = _re.compile(r"^[A-Za-z0-9][A-Za-z0-9.\-]{0,59}$")
_MAX_ETIQUETAS_POR_VEZ = 5000

def _separar_codigos(texto: str) -> list[str]:
    return [c.strip() for c in _re.split(r"[\s,;]+", texto or "") if c.strip()]

def _faixa(prefixo: str, de: str, ate: str) -> list[str]:
    de, ate = (de or "").strip(), (ate or "").strip()
    if not (de.isdigit() and ate.isdigit()):
        raise HTTPException(422, "Faixa: informe o número inicial e o final, só dígitos.")
    ini, fim = int(de), int(ate)
    if fim < ini:
        raise HTTPException(422, "Faixa: o número final é menor que o inicial.")
    if fim - ini + 1 > _MAX_ETIQUETAS_POR_VEZ:
        raise HTTPException(422, f"Faixa grande demais: máximo de {_MAX_ETIQUETAS_POR_VEZ} etiquetas por vez.")
    largura = len(de)
    return [f"{prefixo.upper()}{n:0{largura}d}" for n in range(ini, fim + 1)]

class EtiquetasIn(BaseModel):
    model_config = ConfigDict(extra="ignore")
    codigos: str = ""
    prefixo: str = ""
    de: str = ""
    ate: str = ""
    local: str = ""

    @field_validator("local")
    @classmethod
    def _local(cls, v: str) -> str:
        v = " ".join((v or "").split())[:120]
        if not v:
            raise ValueError("Informe o local em que as etiquetas estão guardadas.")
        return v

    @field_validator("prefixo")
    @classmethod
    def _prefixo(cls, v: str) -> str:
        v = (v or "").strip()
        if v and not _re.fullmatch(r"[A-Za-z0-9.\-]{1,20}", v):
            raise ValueError("Prefixo da faixa inválido (letras, números, ponto ou hífen).")
        return v

@router.get("/cadastro/etiquetas")
def etiquetas_listar(req: Request, situacao: str = "", busca: str = "", local: str = ""):
    _exigir(req, "view")
    db.ensure_db()
    from sqlalchemy import func, select
    sit = (situacao or "").strip().upper()
    termo = (busca or "").strip().lower()
    loc = (local or "").strip().lower()
    with db.SessionLocal() as s:
        totais = {k: 0 for k in db.ETIQUETA_SITUACOES}
        for k, n in s.execute(select(db.Etiqueta.situacao, func.count(db.Etiqueta.id))
                              .group_by(db.Etiqueta.situacao)).all():
            totais[k or db.ETIQUETA_DISPONIVEL] = int(n or 0)
        q = select(db.Etiqueta).order_by(db.Etiqueta.id.desc())
        if sit in db.ETIQUETA_SITUACOES:
            q = q.where(db.Etiqueta.situacao == sit)
        linhas = [e.to_dict() for e in s.scalars(q.limit(_MAX_ETIQUETAS_POR_VEZ)).all()]
        locais = sorted({(l or "") for (l,) in s.execute(
            select(db.Etiqueta.local).distinct()).all() if l})
    if termo:
        linhas = [l for l in linhas if termo in l["codigo"].lower()]
    if loc:
        linhas = [l for l in linhas if l["local"].lower() == loc]
    return {"total": len(linhas), "totais": totais, "locais": locais, "itens": linhas}

@router.post("/cadastro/etiquetas", status_code=201)
def etiquetas_cadastrar(body: EtiquetasIn, req: Request):
    sd = _exigir(req, "create")
    db.ensure_db()
    codigos = _separar_codigos(body.codigos)
    if body.de or body.ate:
        codigos += _faixa(body.prefixo, body.de, body.ate)
    if not codigos:
        raise HTTPException(422, "Nenhuma etiqueta informada: cole a lista ou preencha a faixa.")
    if len(codigos) > _MAX_ETIQUETAS_POR_VEZ:
        raise HTTPException(422, f"Máximo de {_MAX_ETIQUETAS_POR_VEZ} etiquetas por vez.")

    invalidas = [c for c in codigos if not _ETIQUETA_RE.match(c)]
    validas: list[str] = []
    vistos: set[str] = set()
    for c in codigos:
        chave = c.upper()
        if c in invalidas or chave in vistos:
            continue
        vistos.add(chave)
        validas.append(chave)

    from sqlalchemy import func, select
    usuario = sd.get("username", "")
    with db.SessionLocal.begin() as s:
        existentes = {
            (cod or "").upper() for (cod,) in s.execute(
                select(db.Etiqueta.codigo).where(
                    func.upper(db.Etiqueta.codigo).in_([v.upper() for v in validas]))).all()
        } if validas else set()
        repetidas = [v for v in validas if v.upper() in existentes]
        novas = [v for v in validas if v.upper() not in existentes]
        for cod in novas:
            s.add(db.Etiqueta(codigo=cod, local=body.local, criado_por=usuario))
    _log.info("etiquetas cadastradas: %s novas, %s repetidas, %s inválidas, local=%r por=%s ip=%s",
              len(novas), len(repetidas), len(invalidas), body.local, usuario, client_ip(req))
    return {"criadas": len(novas), "repetidas": repetidas[:200], "invalidas": invalidas[:200],
            "local": body.local}

class EtiquetaEditIn(BaseModel):
    model_config = ConfigDict(extra="ignore")
    local: str = ""

    @field_validator("local")
    @classmethod
    def _local(cls, v: str) -> str:
        v = " ".join((v or "").split())[:120]
        if not v:
            raise ValueError("Informe o local.")
        return v

@router.patch("/cadastro/etiquetas/{etiqueta_id}")
def etiqueta_editar(etiqueta_id: int, body: EtiquetaEditIn, req: Request):
    _exigir(req, "edit")
    db.ensure_db()
    with db.SessionLocal.begin() as s:
        e = s.get(db.Etiqueta, etiqueta_id)
        if not e:
            raise HTTPException(404, "Etiqueta não encontrada.")
        if e.situacao != db.ETIQUETA_DISPONIVEL:
            raise HTTPException(409, "Só etiqueta disponível muda de local.")
        e.local = body.local
        return e.to_dict()

class EtiquetaMoverIn(BaseModel):
    model_config = ConfigDict(extra="ignore")
    ids: list[int] = []
    local: str = ""

    @field_validator("local")
    @classmethod
    def _local(cls, v: str) -> str:
        v = " ".join((v or "").split())[:120]
        if not v:
            raise ValueError("Informe o local.")
        return v

@router.post("/cadastro/etiquetas/mover")
def etiquetas_mover(body: EtiquetaMoverIn, req: Request):
    _exigir(req, "edit")
    db.ensure_db()
    if not body.ids:
        raise HTTPException(422, "Escolha ao menos uma etiqueta.")
    movidas = 0
    with db.SessionLocal.begin() as s:
        for eid in body.ids[:_MAX_ETIQUETAS_POR_VEZ]:
            e = s.get(db.Etiqueta, eid)
            if e and e.situacao == db.ETIQUETA_DISPONIVEL:
                e.local = body.local
                movidas += 1
    return {"movidas": movidas, "local": body.local}

class EtiquetaCancelarIn(BaseModel):
    model_config = ConfigDict(extra="ignore")
    motivo: str = ""

    @field_validator("motivo")
    @classmethod
    def _motivo(cls, v: str) -> str:
        v = " ".join((v or "").split())[:200]
        if not v:
            raise ValueError("Diga por que a etiqueta sai do estoque (perdida, danificada…).")
        return v

@router.post("/cadastro/etiquetas/{etiqueta_id}/cancelar")
def etiqueta_cancelar(etiqueta_id: int, body: EtiquetaCancelarIn, req: Request):
    sd = _exigir(req, "edit")
    db.ensure_db()
    with db.SessionLocal.begin() as s:
        e = s.get(db.Etiqueta, etiqueta_id)
        if not e:
            raise HTTPException(404, "Etiqueta não encontrada.")
        if e.situacao != db.ETIQUETA_DISPONIVEL:
            raise HTTPException(409, "Só etiqueta disponível pode ser cancelada.")
        e.situacao = db.ETIQUETA_CANCELADA
        e.cancelada_por = sd.get("username", "")
        e.cancelada_motivo = body.motivo
        dado = e.to_dict()
    _log.info("etiqueta %s cancelada (%s) por=%s ip=%s",
              dado["codigo"], body.motivo, sd.get("username", ""), client_ip(req))
    return dado

@router.delete("/cadastro/etiquetas/{etiqueta_id}", status_code=204)
def etiqueta_excluir(etiqueta_id: int, req: Request):
    _exigir(req, "admin")
    db.ensure_db()
    with db.SessionLocal.begin() as s:
        e = s.get(db.Etiqueta, etiqueta_id)
        if not e:
            raise HTTPException(404, "Etiqueta não encontrada.")
        if e.situacao != db.ETIQUETA_DISPONIVEL:
            raise HTTPException(409, "Etiqueta já usada ou cancelada não se apaga: o histórico é dela.")
        s.delete(e)
    return None

def etiquetas_disponiveis(s, quantidade: int) -> list:
    from sqlalchemy import select
    if quantidade <= 0:
        return []
    return list(s.scalars(
        select(db.Etiqueta)
        .where(db.Etiqueta.situacao == db.ETIQUETA_DISPONIVEL)
        .order_by(db.Etiqueta.id.asc())
        .limit(quantidade)).all())

class ItemImobilizadoIn(BaseModel):
    model_config = ConfigDict(extra="ignore")
    item_ebs: str = ""
    descricao: str = ""

    @field_validator("item_ebs")
    @classmethod
    def _item(cls, v: str) -> str:
        v = db.normalizar_item_ebs(v)
        if not v or not _re.fullmatch(r"[A-Za-z0-9.\-/]{1,60}", v):
            raise ValueError("Código do item do EBS inválido.")
        return v

    @field_validator("descricao")
    @classmethod
    def _desc(cls, v: str) -> str:
        return " ".join((v or "").split())[:200]

class ItensImobilizadosIn(BaseModel):
    model_config = ConfigDict(extra="ignore")
    itens: list[ItemImobilizadoIn] = []
    texto: str = ""

def _itens_do_texto(texto: str) -> list[ItemImobilizadoIn]:
    saida = []
    for linha in (texto or "").splitlines():
        linha = linha.strip()
        if not linha:
            continue
        partes = _re.split(r"[\s;\t]+", linha, maxsplit=1)
        codigo = partes[0].strip(" ;,")
        desc = partes[1].strip() if len(partes) > 1 else ""
        try:
            saida.append(ItemImobilizadoIn(item_ebs=codigo, descricao=desc))
        except ValueError:
            raise HTTPException(422, f"Linha inválida no texto colado: {linha[:80]!r}")
    return saida

@router.get("/cadastro/itens-imobilizados")
def imobilizados_listar(req: Request, busca: str = ""):
    _exigir(req, "view")
    db.ensure_db()
    from sqlalchemy import select
    termo = (busca or "").strip().lower()
    with db.SessionLocal() as s:
        linhas = [i.to_dict() for i in s.scalars(
            select(db.ItemImobilizado).order_by(db.ItemImobilizado.item_ebs)).all()]
    if termo:
        linhas = [l for l in linhas
                  if termo in l["item_ebs"].lower() or termo in l["descricao"].lower()]
    return {"total": len(linhas), "itens": linhas}

@router.post("/cadastro/itens-imobilizados", status_code=201)
def imobilizados_cadastrar(body: ItensImobilizadosIn, req: Request):
    sd = _exigir(req, "create")
    db.ensure_db()
    itens = list(body.itens) + _itens_do_texto(body.texto)
    if not itens:
        raise HTTPException(422, "Nenhum item informado.")
    if len(itens) > 2000:
        raise HTTPException(422, "Máximo de 2000 itens por vez.")
    from sqlalchemy import select
    usuario = sd.get("username", "")
    criados, atualizados = 0, 0
    with db.SessionLocal.begin() as s:
        for it in itens:
            atual = s.scalar(select(db.ItemImobilizado)
                             .where(db.ItemImobilizado.item_ebs == it.item_ebs))
            if atual:
                if it.descricao and it.descricao != atual.descricao:
                    atual.descricao = it.descricao
                    atualizados += 1
                continue
            s.add(db.ItemImobilizado(item_ebs=it.item_ebs, descricao=it.descricao,
                                     criado_por=usuario))
            criados += 1
    _log.info("itens imobilizados: %s criados, %s atualizados por=%s ip=%s",
              criados, atualizados, usuario, client_ip(req))
    return {"criados": criados, "atualizados": atualizados}

@router.patch("/cadastro/itens-imobilizados/{item_id}")
def imobilizado_editar(item_id: int, body: ItemImobilizadoIn, req: Request):
    _exigir(req, "edit")
    db.ensure_db()
    from sqlalchemy import select
    with db.SessionLocal.begin() as s:
        i = s.get(db.ItemImobilizado, item_id)
        if not i:
            raise HTTPException(404, "Item não encontrado.")
        outro = s.scalar(select(db.ItemImobilizado).where(
            db.ItemImobilizado.item_ebs == body.item_ebs, db.ItemImobilizado.id != item_id))
        if outro:
            raise HTTPException(409, f"O item {body.item_ebs} já está na lista.")
        i.item_ebs, i.descricao = body.item_ebs, body.descricao
        return i.to_dict()

@router.delete("/cadastro/itens-imobilizados/{item_id}", status_code=204)
def imobilizado_excluir(item_id: int, req: Request):
    sd = _exigir(req, "admin")
    db.ensure_db()
    with db.SessionLocal.begin() as s:
        i = s.get(db.ItemImobilizado, item_id)
        if not i:
            raise HTTPException(404, "Item não encontrado.")
        codigo = i.item_ebs
        s.delete(i)
    _log.info("item imobilizado %s removido por=%s ip=%s", codigo, sd.get("username", ""), client_ip(req))
    return None

def codigos_imobilizados(s) -> set[str]:
    from sqlalchemy import select
    return {db.normalizar_item_ebs(c) for (c,) in s.execute(
        select(db.ItemImobilizado.item_ebs)).all() if c}


XLSX_MIME = _XLSX_MIME


def _dados_do_agendamento(agendamento_id: int) -> tuple[dict, list[dict], dict]:
    agf_db.ensure_db()
    with agf_db.SessionLocal() as sa:
        ag = sa.get(agf_db.Agendamento, agendamento_id)
        if not ag:
            raise HTTPException(404, "Agendamento não encontrado.")
        dado = ag.to_dict()
        notas = [n.to_dict() | {"pdf_arquivo": n.pdf_arquivo or "", "xml_arquivo": n.xml_arquivo or ""} for n in ag.notas]
    po_da_nf: dict[str, list[str]] = {}
    for p in dado.get("pedidos", []):
        if p.get("nf"):
            po_da_nf.setdefault(p["nf"], []).append(p["po"])
    for nf in po_da_nf:
        if not any(n["nf"] == nf for n in notas):
            notas.append({"nf": nf, "chave": "", "vencimento": "", "origem": "", "origem_rotulo": "",
                          "tem_xml": False, "tem_pdf": False, "itens": [], "emitente": "", "cstat": "",
                          "erro": "", "pdf_arquivo": "", "xml_arquivo": ""})
    for n in notas:
        n["po"] = " / ".join(po_da_nf.get(n["nf"], [])) or dado.get("po", "")
    return dado, notas, po_da_nf


def _planilha(proc, notas: list[dict]) -> tuple[str, bytes]:
    from core import planilha_cadastro_ativos as plan
    linhas = [{"item": a.ebs_item or "", "descricao": a.descricao or "", "plaqueta": a.plaqueta or "",
               "numero_serie": a.numero_serie or "", "nf": a.nf or proc.nf or ""} for a in proc.ativos]
    nfs = sorted({l["nf"] for l in linhas if l["nf"]}) or [proc.nf or str(proc.agendamento_id)]
    return plan.nome_arquivo(nfs), plan.montar(linhas)


def _sessao_sn(req: Request):
    sd = get_session(req)
    cookies = sd.get("sn_cookies")
    if not cookies:
        raise HTTPException(409, "Sessão ServiceNow não ativa. Saia e entre de novo no portal (Logon AD).")
    from integracoes import http as http_saida
    from integracoes.sn_catalogo import CatalogoServiceNow
    from routers import servicenow as sn
    s = http_saida.sessao("servicenow", proxy=sn.SN_PROXY or None)
    s.cookies.update(cookies)
    return CatalogoServiceNow(s, sn.SERVICENOW_BASE)


def _abrir_chamado(req: Request, agendamento: dict, notas: list[dict], planilha_nome: str,
                   planilha_bytes: bytes) -> dict:
    from integracoes.sn_catalogo import montar_variaveis_lancamento_nf
    from core import arquivos_temporarios
    cat = _sessao_sn(req)
    eu = cat.usuario_atual()
    impacto = cat.procurar_usuario(_cfg.SN_CATALOGO_LANCAMENTO_NF_IMPACTO) or {}
    item = cat.descrever_item(_cfg.SN_CATALOGO_LANCAMENTO_NF_ITEM)
    dados = {
        "requested_for_sys_id": eu.get("sys_id", ""),
        "impact_employee_sys_id": impacto.get("sys_id", ""),
        "telefone": _cfg.SN_CATALOGO_LANCAMENTO_NF_TELEFONE,
        "bu": agendamento.get("bu", ""),
        "fornecedor": agendamento.get("fornecedor", ""),
        "notas": [{"po": n["po"], "nf": n["nf"], "vencimento": n.get("vencimento", "")} for n in notas],
    }
    variaveis, faltantes = montar_variaveis_lancamento_nf(item, dados)
    pedido = cat.enviar_pedido(item["sys_id"], variaveis)
    ritm = cat.ritm_do_pedido(pedido["request_sys_id"]) or {}
    tabela, alvo = ("sc_req_item", ritm["sys_id"]) if ritm.get("sys_id") else ("sc_request", pedido["request_sys_id"])
    anexos, sem_pdf = [], []
    cat.anexar(tabela, alvo, planilha_nome, planilha_bytes, XLSX_MIME)
    anexos.append(planilha_nome)
    for n in notas:
        caminho = arquivos_temporarios.caminho(AREA_TMP_NOTAS, agendamento["id"], n.get("pdf_arquivo") or "") if n.get("pdf_arquivo") else None
        if not caminho:
            sem_pdf.append(n["nf"])
            continue
        cat.anexar(tabela, alvo, caminho.name, caminho.read_bytes(), "application/pdf")
        anexos.append(caminho.name)
    return {"request_sys_id": pedido["request_sys_id"], "request_number": pedido["request_number"],
            "ritm_sys_id": ritm.get("sys_id", ""), "ritm_number": ritm.get("number", ""),
            "faltantes": faltantes, "anexos": anexos, "sem_pdf": sem_pdf,
            "aviso": (f"NF sem PDF anexado: {', '.join(sem_pdf)}. " if sem_pdf else "")
                     + (f"Variáveis que o item não tem: {', '.join(faltantes)}." if faltantes else "")}


def _lancar(agendamento_id: int, req: Request, reenvio: bool) -> dict:
    sd = _exigir(req, "edit")
    check_rate_limit(req, "api")
    db.ensure_db()
    from sqlalchemy import select
    from core import arquivos_temporarios
    usuario = sd.get("username", "")
    agendamento, notas, _ = _dados_do_agendamento(agendamento_id)
    with db.SessionLocal.begin() as s:
        proc = s.scalar(select(db.Processo).where(db.Processo.agendamento_id == agendamento_id))
        if not proc or not proc.ativos:
            raise HTTPException(422, "Nada lançado para este agendamento.")
        if reenvio and not proc.lancado_em:
            raise HTTPException(409, "Este lançamento ainda não foi enviado; use o OK.")
        if reenvio and proc.sn_request_sys_id:
            raise HTTPException(409, f"O chamado {proc.sn_request_number} já foi aberto.")
        if not reenvio and proc.lancado_em:
            raise HTTPException(409, "Este lançamento já foi enviado.")
        faltando = [a.numero_serie or a.plaqueta or f"linha {i + 1}" for i, a in enumerate(proc.ativos)
                    if not (a.plaqueta and a.numero_serie)]
        if faltando:
            raise HTTPException(422, "Toda linha precisa de plaqueta e número de série: " + ", ".join(faltando[:10]))
        nome, conteudo = _planilha(proc, notas)
        caminho = arquivos_temporarios.gravar(AREA_TMP, agendamento_id, nome, conteudo)
        agora = db.utcnow()
        proc.planilha_arquivo, proc.planilha_em = caminho.name, agora
        if not reenvio:
            proc.lancado_por, proc.lancado_em = usuario, agora
            proc.status = "CONCLUIDA"
            proc.atualizado_em = agora
        s.flush()
    chamado, erro = None, ""
    try:
        chamado = _abrir_chamado(req, agendamento, notas, nome, conteudo)
    except HTTPException as exc:
        erro = str(exc.detail)
    except Exception as exc:  # noqa: BLE001
        erro = str(exc)[:300]
        _log.error("chamado do lançamento %s falhou: %s", agendamento_id, exc, exc_info=True)
    with db.SessionLocal.begin() as s:
        proc = s.scalar(select(db.Processo).where(db.Processo.agendamento_id == agendamento_id))
        if chamado:
            proc.sn_request_number = chamado["request_number"][:40]
            proc.sn_request_sys_id = chamado["request_sys_id"][:64]
            proc.sn_ritm_number = chamado["ritm_number"][:40]
            proc.sn_ritm_sys_id = chamado["ritm_sys_id"][:64]
            proc.sn_enviado_em = db.utcnow()
            proc.sn_erro = ""
        else:
            proc.sn_erro = erro[:300]
        dado = proc.to_dict()
    _log.info("lançamento %s agendamento=%s planilha=%s chamado=%s erro=%r por=%s ip=%s",
              "reenviado" if reenvio else "enviado", agendamento_id, nome,
              (chamado or {}).get("request_number", ""), erro, usuario, client_ip(req))
    dado["chamado"] = chamado
    dado["aviso"] = (chamado or {}).get("aviso", "") or (
        f"Planilha gerada; o chamado no ServiceNow não foi aberto: {erro}" if erro else "")
    return dado


@router.post("/{agendamento_id}/lancar")
def lancar(agendamento_id: int, req: Request):
    return _lancar(agendamento_id, req, reenvio=False)


@router.post("/{agendamento_id}/lancar/reenviar")
def lancar_reenviar(agendamento_id: int, req: Request):
    return _lancar(agendamento_id, req, reenvio=True)


@router.get("/{agendamento_id}/planilha")
def baixar_planilha(agendamento_id: int, req: Request):
    _exigir(req, "view")
    db.ensure_db()
    from sqlalchemy import select
    from core import arquivos_temporarios
    with db.SessionLocal() as s:
        proc = s.scalar(select(db.Processo).where(db.Processo.agendamento_id == agendamento_id))
        nome = proc.planilha_arquivo if proc else ""
    caminho = arquivos_temporarios.caminho(AREA_TMP, agendamento_id, nome) if nome else None
    if not caminho:
        raise HTTPException(404, "Planilha não encontrada (pode ter passado dos cinco dias).")
    return FileResponse(str(caminho), filename=caminho.name, media_type=XLSX_MIME)


@router.get("/{agendamento_id}/nota/{nf}/pdf")
def baixar_pdf_nota(agendamento_id: int, nf: str, req: Request):
    _exigir(req, "view")
    from core import arquivos_temporarios
    _, notas, _ = _dados_do_agendamento(agendamento_id)
    nota = next((n for n in notas if n["nf"] == nf), None)
    caminho = arquivos_temporarios.caminho(AREA_TMP_NOTAS, agendamento_id, nota["pdf_arquivo"]) if nota and nota.get("pdf_arquivo") else None
    if not caminho:
        raise HTTPException(404, "PDF da nota não encontrado (pode ter passado dos cinco dias).")
    return FileResponse(str(caminho), filename=caminho.name, media_type="application/pdf")


class NotaVencimentoIn(BaseModel):
    model_config = ConfigDict(extra="ignore")
    vencimento: str = ""

    @field_validator("vencimento")
    @classmethod
    def _venc(cls, v: str) -> str:
        v = (v or "").strip()
        if v:
            date.fromisoformat(v)
        return v


@router.patch("/{agendamento_id}/nota/{nf}")
def nota_vencimento(agendamento_id: int, nf: str, body: NotaVencimentoIn, req: Request):
    sd = _exigir(req, "edit")
    from routers import agendamentos_forn as agf
    agf_db.ensure_db()
    with agf_db.SessionLocal.begin() as sa:
        ag = sa.get(agf_db.Agendamento, agendamento_id)
        if not ag:
            raise HTTPException(404, "Agendamento não encontrado.")
        nota = agf.registrar_nota(sa, ag, nf, sd.get("username", ""),
                                  vencimento=date.fromisoformat(body.vencimento) if body.vencimento else None)
        return nota.to_dict()


@router.get("/catalogo-sn/conferir")
def conferir_catalogo(req: Request):
    _exigir(req, "edit")
    check_rate_limit(req, "api")
    from integracoes.sn_catalogo import montar_variaveis_lancamento_nf, ErroServiceNow
    try:
        cat = _sessao_sn(req)
        eu = cat.usuario_atual()
        item = cat.descrever_item(_cfg.SN_CATALOGO_LANCAMENTO_NF_ITEM)
        impacto = cat.procurar_usuario(_cfg.SN_CATALOGO_LANCAMENTO_NF_IMPACTO) or {}
    except ErroServiceNow as exc:
        raise HTTPException(502, str(exc)) from exc
    exemplo = {"requested_for_sys_id": eu.get("sys_id", ""), "impact_employee_sys_id": impacto.get("sys_id", ""),
               "telefone": _cfg.SN_CATALOGO_LANCAMENTO_NF_TELEFONE, "bu": "Renner", "fornecedor": "EXEMPLO",
               "notas": [{"po": "0000000", "nf": "0", "vencimento": "2026-01-01"}]}
    try:
        variaveis, faltantes = montar_variaveis_lancamento_nf(item, exemplo)
        erro = ""
    except (KeyError, ValueError) as exc:
        variaveis, faltantes, erro = {}, [], str(exc)
    return {"usuario": eu, "impacto": impacto, "item": item, "exemplo": variaveis,
            "faltantes": faltantes, "erro": erro}


@router.get("/nfe/certificados")
def certificados_nfe(req: Request):
    _exigir(req, "view")
    saida = []
    for bu in ("Renner", "Camicado", "Youcom"):
        linha = {"bu": bu, "configurado": False, "cnpj": "", "valido_ate": "", "vencido": False, "erro": ""}
        try:
            from integracoes import nfe_sefaz
            cert = nfe_sefaz.certificado_da_bu(bu)
            if cert:
                linha.update({"configurado": True, "cnpj": cert.cnpj, "valido_ate": cert.valido_ate.isoformat(),
                              "vencido": cert.vencido})
        except Exception as exc:  # noqa: BLE001
            linha["erro"] = str(exc)[:200]
        saida.append(linha)
    return {"certificados": saida}
